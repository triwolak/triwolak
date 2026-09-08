# -*- coding: utf-8 -*-
"""Testy modułu :mod:`flexit2xlsx.xlsxwrite` — zapis XLSX.

Każdy test formatu jest uruchamiany DWA RAZY: raz dla wbudowanego zapisu
(``stdlib``) i raz dla zapisu przez ``openpyxl``.  Realizuje to klasa-mieszanka
:class:`_BackendCase`, po której dziedziczą dwie klasy testowe różniące się
wyłącznie polem ``BACKEND`` (ustawianym w zmiennej środowiskowej
``FLEXIT_XLSX_BACKEND``).

Weryfikacja idzie dwiema drogami:

* **odczyt przez openpyxl** (``load_workbook``) — wartości, typy komórek
  (``cell.data_type``), formaty liczbowe, pogrubienie nagłówka, zamrożenie,
  autofiltr, szerokości kolumn;
* **surowy XML z archiwum ZIP** — bo tylko tak da się udowodnić, że tekst
  ``"=A1+1"`` NIE trafił do pliku jako element ``<f>`` (formuła) oraz że pakiet
  OOXML jest spójny (``[Content_Types].xml``, relacje, kolejność elementów).

Uruchamiane przez ``python3 -m unittest`` oraz ``python3 -m pytest``.
"""

import datetime as dt
import io
import os
import pathlib
import re
import shutil
import sys
import tempfile
import time
import tracemalloc
import unittest
import xml.etree.ElementTree as ET
import zipfile
from contextlib import redirect_stderr
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from flexit2xlsx import xlsxwrite  # noqa: E402
from flexit2xlsx.values import MAX_CELL_CHARS  # noqa: E402
from flexit2xlsx.xlsxwrite import (  # noqa: E402
    ENV_BACKEND,
    Sheet,
    backend_name,
    safe_sheet_name,
    write_workbook,
)

try:  # openpyxl służy WYŁĄCZNIE do weryfikacji wyników w testach
    import openpyxl
    from openpyxl.styles.numbers import is_date_format
except ImportError:  # pragma: no cover - zależy od środowiska
    openpyxl = None
    is_date_format = None

POLSKIE = "Zażółć gęślą jaźń ŁÓDŹ ĄĘĆŃÓŚŹŻ"

#: Kolejność elementów potomnych ``<worksheet>`` wymuszona przez schemat OOXML.
_WORKSHEET_ORDER = [
    "sheetPr",
    "dimension",
    "sheetViews",
    "sheetFormatPr",
    "cols",
    "sheetData",
    "sheetCalcPr",
    "sheetProtection",
    "protectedRanges",
    "scenarios",
    "autoFilter",
    "sortState",
    "dataConsolidate",
    "customSheetViews",
    "mergeCells",
    "phoneticPr",
    "conditionalFormatting",
    "dataValidations",
    "hyperlinks",
    "printOptions",
    "pageMargins",
    "pageSetup",
    "headerFooter",
]


# --------------------------------------------------------------------------- #
# Pomocnicze narzędzia do zaglądania do archiwum
# --------------------------------------------------------------------------- #


def read_part(path, name):
    """Zwraca zawartość jednej części pakietu XLSX jako tekst."""
    with zipfile.ZipFile(path) as archive:
        return archive.read(name).decode("utf-8")


def sheet_xml(path, index=1):
    """Zwraca surowy XML arkusza o podanym numerze (1-based)."""
    return read_part(path, "xl/worksheets/sheet%d.xml" % index)


def cell_xml(xml, ref):
    """Wycina z XML-a arkusza element ``<c>`` o podanym adresie (np. ``A2``)."""
    match = re.search(
        r'<c r="%s"(?:\s[^>]*)?(?:/>|>.*?</c>)' % re.escape(ref), xml, re.DOTALL
    )
    return match.group(0) if match else ""


def resolve_target(base, target):
    """Rozwiązuje ``Target`` relacji OPC do nazwy części w archiwum.

    Cel zaczynający się od ``/`` jest liczony od korzenia pakietu (tak zapisuje
    openpyxl), pozostałe — względem katalogu części, do której należy plik
    ``.rels`` (tak zapisuje Excel i wbudowany zapis).
    """
    if target.startswith("/"):
        return target[1:]
    return os.path.normpath(os.path.join(base, target)).replace(os.sep, "/")


def make_rows(count, columns):
    """Generator wierszy testowych (mieszanka tekstu, liczb całkowitych i zmiennoprzecinkowych)."""
    for row_index in range(count):
        yield [
            "tekst %d-%d" % (row_index, col_index)
            if col_index % 3 == 0
            else (row_index * col_index if col_index % 3 == 1 else row_index / 7.0)
            for col_index in range(columns)
        ]


# --------------------------------------------------------------------------- #
# Testy niezależne od ścieżki zapisu
# --------------------------------------------------------------------------- #


class TestSafeSheetName(unittest.TestCase):
    """Oczyszczanie i unikalność nazw arkuszy."""

    def test_zwykla_nazwa_bez_zmian(self):
        used = set()
        self.assertEqual(safe_sheet_name("Wszystkie aukcje", used), "Wszystkie aukcje")
        self.assertIn("Wszystkie aukcje", used)

    def test_znaki_zakazane_zamieniane_na_podkreslenie(self):
        cases = [
            ("a/b", "a_b"),
            ("a\\b", "a_b"),
            ("a:b", "a_b"),
            ("a*b", "a_b"),
            ("a?b", "a_b"),
            ("[Dane]", "_Dane_"),
        ]
        for raw, expected in cases:
            with self.subTest(raw=raw):
                self.assertEqual(safe_sheet_name(raw, set()), expected)

    def test_dlugosc_przycieta_do_31_znakow(self):
        raw = "Bardzo długa nazwa arkusza aukcyjnego z portalu flexitauctions"
        name = safe_sheet_name(raw, set())
        self.assertLessEqual(len(name), 31)
        self.assertTrue(raw.startswith(name.rstrip()))

    def test_polskie_znaki_zachowane(self):
        self.assertEqual(safe_sheet_name("Zażółć gęślą", set()), "Zażółć gęślą")

    def test_biale_znaki_i_znaki_sterujace_zwijane(self):
        self.assertEqual(safe_sheet_name("  a\t\tb\n c  ", set()), "a b c")
        self.assertEqual(safe_sheet_name("x\x00y", set()), "x y")

    def test_pusta_nazwa_dostaje_wartosc_domyslna(self):
        for raw in ("", "   ", None, "'''", "\x00\x01"):
            with self.subTest(raw=raw):
                self.assertEqual(safe_sheet_name(raw, set()), "Arkusz")

    def test_apostrofy_na_brzegach_usuwane(self):
        self.assertEqual(safe_sheet_name("'Dane'", set()), "Dane")

    def test_nazwa_zarezerwowana(self):
        self.assertEqual(safe_sheet_name("History", set()), "History_")
        self.assertEqual(safe_sheet_name("history", set()), "history_")

    def test_duplikaty_dostaja_przyrostek(self):
        used = set()
        self.assertEqual(safe_sheet_name("Dane", used), "Dane")
        self.assertEqual(safe_sheet_name("Dane", used), "Dane (2)")
        self.assertEqual(safe_sheet_name("Dane", used), "Dane (3)")

    def test_unikalnosc_bez_wzgledu_na_wielkosc_liter(self):
        used = set()
        safe_sheet_name("Dane", used)
        self.assertEqual(safe_sheet_name("DANE", used), "DANE (2)")

    def test_duplikat_dlugiej_nazwy_nadal_miesci_sie_w_limicie(self):
        used = set()
        raw = "A" * 40
        first = safe_sheet_name(raw, used)
        second = safe_sheet_name(raw, used)
        self.assertEqual(len(first), 31)
        self.assertLessEqual(len(second), 31)
        self.assertNotEqual(first, second)
        self.assertTrue(second.endswith(" (2)"))

    def test_duplikaty_po_przycieciu_roznych_nazw(self):
        """Dwie różne, długie nazwy skracają się do tego samego — muszą się rozjechać."""
        used = set()
        first = safe_sheet_name("Aukcja flexit 2026 pozycje sprzętu A", used)
        second = safe_sheet_name("Aukcja flexit 2026 pozycje sprzętu B", used)
        self.assertNotEqual(first, second)
        self.assertLessEqual(len(second), 31)


class TestBackendSelection(unittest.TestCase):
    """Wybór ścieżki zapisu przez zmienną środowiskową."""

    def setUp(self):
        self._previous = os.environ.get(ENV_BACKEND)
        self.addCleanup(self._restore)

    def _restore(self):
        if self._previous is None:
            os.environ.pop(ENV_BACKEND, None)
        else:
            os.environ[ENV_BACKEND] = self._previous

    def _auto_expected(self):
        return "openpyxl" if xlsxwrite._openpyxl is not None else "stdlib"

    def test_brak_zmiennej_to_tryb_auto(self):
        os.environ.pop(ENV_BACKEND, None)
        self.assertEqual(backend_name(), self._auto_expected())

    def test_wymuszenie_stdlib(self):
        os.environ[ENV_BACKEND] = "stdlib"
        self.assertEqual(backend_name(), "stdlib")

    def test_wartosc_jest_normalizowana(self):
        os.environ[ENV_BACKEND] = "  STDLIB \n"
        self.assertEqual(backend_name(), "stdlib")

    def test_pusta_wartosc_to_auto(self):
        os.environ[ENV_BACKEND] = ""
        self.assertEqual(backend_name(), self._auto_expected())

    def test_nieznana_wartosc_ostrzega_i_wraca_do_auto(self):
        os.environ[ENV_BACKEND] = "excel95"
        stderr = io.StringIO()
        with redirect_stderr(stderr):
            self.assertEqual(backend_name(), self._auto_expected())
        self.assertIn("excel95", stderr.getvalue())

    def test_zadanie_openpyxl_bez_biblioteki_schodzi_na_stdlib(self):
        os.environ[ENV_BACKEND] = "openpyxl"
        stderr = io.StringIO()
        with mock.patch.object(xlsxwrite, "_openpyxl", None), redirect_stderr(stderr):
            self.assertEqual(backend_name(), "stdlib")
        self.assertIn("openpyxl", stderr.getvalue())

    @unittest.skipIf(openpyxl is None, "openpyxl niedostępny")
    def test_zadanie_openpyxl_gdy_jest_biblioteka(self):
        os.environ[ENV_BACKEND] = "openpyxl"
        self.assertEqual(backend_name(), "openpyxl")


class TestSheetDataclass(unittest.TestCase):
    """Kontrakt dataclassy :class:`Sheet`."""

    def test_pola_i_wartosc_domyslna_wierszy(self):
        sheet = Sheet(name="A", columns=["x"])
        self.assertEqual(sheet.name, "A")
        self.assertEqual(sheet.columns, ["x"])
        self.assertEqual(list(sheet.rows), [])

    def test_wiersze_moga_byc_generatorem(self):
        sheet = Sheet(name="A", columns=["x"], rows=(row for row in ([1], [2])))
        self.assertEqual(list(sheet.rows), [[1], [2]])


# --------------------------------------------------------------------------- #
# Testy wspólne dla obu ścieżek zapisu
# --------------------------------------------------------------------------- #


class _BackendCase:
    """Zestaw testów uruchamiany osobno dla każdej ścieżki zapisu.

    Klasa NIE dziedziczy po ``unittest.TestCase`` — robią to dopiero podklasy,
    dzięki czemu ten zestaw nie uruchamia się „sam z siebie” bez ustawionego
    backendu.
    """

    BACKEND = "stdlib"

    # -- infrastruktura ----------------------------------------------------- #

    def setUp(self):
        if openpyxl is None:  # pragma: no cover - zależy od środowiska
            self.skipTest("openpyxl jest potrzebny do weryfikacji wyników")
        self._previous = os.environ.get(ENV_BACKEND)
        os.environ[ENV_BACKEND] = self.BACKEND
        self.addCleanup(self._restore_env)
        self.tmpdir = tempfile.mkdtemp(prefix="flexit2xlsx-test-")
        self.addCleanup(shutil.rmtree, self.tmpdir, ignore_errors=True)
        self.assertEqual(backend_name(), self.BACKEND)

    def _restore_env(self):
        if self._previous is None:
            os.environ.pop(ENV_BACKEND, None)
        else:
            os.environ[ENV_BACKEND] = self._previous

    def out(self, name="wynik.xlsx"):
        """Ścieżka pliku w katalogu tymczasowym testu."""
        return os.path.join(self.tmpdir, name)

    def write(self, sheets, name="wynik.xlsx", **options):
        """Zapisuje skoroszyt i zwraca ścieżkę pliku."""
        path = self.out(name)
        write_workbook(path, sheets, **options)
        return path

    def assert_backend_used(self, path):
        """Sprawdza, że plik naprawdę powstał w oczekiwanej ścieżce zapisu.

        Rozróżnienie po ``xl/theme/theme1.xml``: openpyxl zawsze dokłada motyw,
        wbudowany zapis nigdy go nie tworzy (nie odwołuje się do motywu wcale).
        """
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
        if self.BACKEND == "openpyxl":
            self.assertIn("xl/theme/theme1.xml", names)
        else:
            self.assertNotIn("xl/theme/theme1.xml", names)

    # -- podstawy ----------------------------------------------------------- #

    def test_podstawowy_zapis_i_odczyt(self):
        path = self.write(
            [Sheet(name="Dane", columns=["Aukcja", "Model", "Cena"],
                   rows=[["A1", "ThinkPad T480", 1299.5], ["A2", "Latitude 7490", 999]])]
        )
        self.assert_backend_used(path)
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual(worksheet.title, "Dane")
        self.assertEqual([cell.value for cell in worksheet[1]], ["Aukcja", "Model", "Cena"])
        self.assertEqual([cell.value for cell in worksheet[2]], ["A1", "ThinkPad T480", 1299.5])
        self.assertEqual([cell.value for cell in worksheet[3]], ["A2", "Latitude 7490", 999])
        self.assertEqual(worksheet.max_row, 3)
        self.assertEqual(worksheet.max_column, 3)

    def test_naglowek_pogrubiony_zamrozony_z_autofiltrem_i_szerokoscia(self):
        path = self.write(
            [Sheet(name="Dane", columns=["Krótka", "Bardzo długa nazwa kolumny"],
                   rows=[["x", "y"]])]
        )
        worksheet = openpyxl.load_workbook(path).active
        self.assertTrue(worksheet["A1"].font.bold)
        self.assertTrue(worksheet["B1"].font.bold)
        self.assertFalse(bool(worksheet["A2"].font.bold))
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(worksheet.auto_filter.ref, "A1:B2")
        wide = worksheet.column_dimensions["B"].width
        narrow = worksheet.column_dimensions["A"].width
        self.assertGreater(wide, narrow)
        self.assertLessEqual(wide, xlsxwrite.MAX_COL_WIDTH)
        self.assertGreaterEqual(narrow, xlsxwrite.MIN_COL_WIDTH)

    def test_wylaczone_opcje_formatowania(self):
        path = self.write(
            [Sheet(name="Dane", columns=["a", "b"], rows=[["x", "y"]])],
            freeze_header=False,
            autofilter=False,
            auto_width=False,
        )
        worksheet = openpyxl.load_workbook(path).active
        self.assertIsNone(worksheet.freeze_panes)
        self.assertIsNone(worksheet.auto_filter.ref)
        self.assertNotIn("A", worksheet.column_dimensions)

    # -- ochrona przed formułami ------------------------------------------- #

    def test_tekst_wygladajacy_na_formule_jest_tekstem(self):
        groźne = [
            "=A1+1",
            "=cmd|' /C calc'!A0",
            "+SUM(A1:A9)",
            "-2+3",
            "@SUM(1)",
            "\tuciekinier",
            "\rpowrot",
        ]
        path = self.write([Sheet(name="Formuły", columns=["wartość"],
                                 rows=[[text] for text in groźne])])
        worksheet = openpyxl.load_workbook(path).active
        xml = sheet_xml(path, 1)

        # 1. Nigdzie w arkuszu nie ma elementu formuły.
        self.assertNotIn("<f>", xml)
        self.assertNotIn("<f ", xml)

        for index, text in enumerate(groźne, start=2):
            with self.subTest(text=text):
                cell = worksheet.cell(row=index, column=1)
                # 2. openpyxl widzi typ „napis”, a nie „formuła” ('f').
                self.assertEqual(cell.data_type, "s")
                self.assertEqual(cell.value, text)
                # 3. Surowy XML komórki: napis osadzony, bez <f>.
                raw = cell_xml(xml, "A%d" % index)
                self.assertIn('t="inlineStr"', raw)
                self.assertNotIn("<f", raw)

    def test_naglowek_wygladajacy_na_formule_tez_jest_tekstem(self):
        path = self.write([Sheet(name="H", columns=["=SUMA", "-ujemna"], rows=[[1, 2]])])
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual(worksheet["A1"].data_type, "s")
        self.assertEqual(worksheet["A1"].value, "=SUMA")
        self.assertEqual(worksheet["B1"].value, "-ujemna")
        self.assertNotIn("<f", sheet_xml(path, 1))

    def test_apostrof_nie_jest_doklejany(self):
        path = self.write([Sheet(name="A", columns=["k"], rows=[["=A1+1"]])])
        worksheet = openpyxl.load_workbook(path).active
        self.assertFalse(worksheet["A2"].value.startswith("'"))

    # -- typy komórek ------------------------------------------------------- #

    def test_typy_liczbowe_logiczne_i_puste(self):
        rows = [[42, 3.5, True, None], [-7, -0.25, False, None], [0, 0.0, True, None]]
        path = self.write([Sheet(name="T", columns=["int", "float", "bool", "pusto"], rows=rows)])
        worksheet = openpyxl.load_workbook(path).active

        self.assertEqual(worksheet["A2"].value, 42)
        self.assertEqual(worksheet["A2"].data_type, "n")
        self.assertEqual(worksheet["B2"].value, 3.5)
        self.assertEqual(worksheet["B2"].data_type, "n")
        self.assertIs(worksheet["C2"].value, True)
        self.assertEqual(worksheet["C2"].data_type, "b")
        self.assertIs(worksheet["C3"].value, False)
        self.assertIsNone(worksheet["D2"].value)

        self.assertEqual(worksheet["A3"].value, -7)
        self.assertEqual(worksheet["B3"].value, -0.25)
        self.assertEqual(worksheet["A4"].value, 0)

    def test_bool_zapisany_jako_typ_natywny(self):
        """W pliku ma być ``t="b"``, a nie napis — inaczej PRAWDA/FAŁSZ zależałoby od wersji Excela."""
        path = self.write([Sheet(name="B", columns=["flaga"], rows=[[True], [False]]) ])
        xml = sheet_xml(path, 1)
        self.assertIn('t="b"', cell_xml(xml, "A2"))
        self.assertIn("<v>1</v>", cell_xml(xml, "A2"))
        self.assertIn("<v>0</v>", cell_xml(xml, "A3"))

    def test_daty_maja_format_liczbowy(self):
        rows = [
            [dt.date(2026, 2, 26), dt.datetime(2026, 2, 26, 14, 30, 15), dt.time(9, 5, 0)],
        ]
        path = self.write([Sheet(name="D", columns=["data", "dataczas", "czas"], rows=rows)])
        worksheet = openpyxl.load_workbook(path).active

        data = worksheet["A2"]
        self.assertEqual(data.data_type, "d")
        self.assertTrue(is_date_format(data.number_format), data.number_format)
        self.assertEqual(getattr(data.value, "year"), 2026)
        self.assertEqual(getattr(data.value, "month"), 2)
        self.assertEqual(getattr(data.value, "day"), 26)

        stamp = worksheet["B2"]
        self.assertEqual(stamp.data_type, "d")
        self.assertTrue(is_date_format(stamp.number_format), stamp.number_format)
        self.assertEqual(stamp.value, dt.datetime(2026, 2, 26, 14, 30, 15))

        clock = worksheet["C2"]
        self.assertTrue(is_date_format(clock.number_format), clock.number_format)
        self.assertEqual(clock.value, dt.time(9, 5, 0))

        # W surowym XML data to liczba (numer seryjny), nie tekst.
        self.assertNotIn("inlineStr", cell_xml(sheet_xml(path, 1), "A2"))

    def test_data_z_offsetem_strefy_sprowadzona_do_utc(self):
        aware = dt.datetime(2026, 2, 26, 14, 0, tzinfo=dt.timezone(dt.timedelta(hours=2)))
        path = self.write([Sheet(name="D", columns=["kiedy"], rows=[[aware]])])
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual(worksheet["A2"].value, dt.datetime(2026, 2, 26, 12, 0))

    def test_data_sprzed_1900_zapisana_jako_tekst(self):
        path = self.write([Sheet(name="D", columns=["kiedy"], rows=[[dt.date(1899, 5, 1)]])])
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual(worksheet["A2"].value, "1899-05-01")
        self.assertEqual(worksheet["A2"].data_type, "s")

    def test_nan_i_nieskonczonosc_jako_tekst(self):
        path = self.write([Sheet(name="N", columns=["x"], rows=[[float("nan")], [float("inf")]])])
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual(worksheet["A2"].data_type, "s")
        self.assertEqual(worksheet["A3"].data_type, "s")

    # -- tekst: znaki sterujące, długość, polskie znaki --------------------- #

    def test_bardzo_dlugi_tekst_przyciety_do_limitu_excela(self):
        long_text = "ą" * (MAX_CELL_CHARS + 5000)
        path = self.write([Sheet(name="L", columns=["tekst"], rows=[[long_text]])])
        worksheet = openpyxl.load_workbook(path).active
        value = worksheet["A2"].value
        self.assertEqual(len(value), MAX_CELL_CHARS)
        self.assertEqual(set(value), {"ą"})

    def test_znaki_sterujace_usuniete_a_dozwolone_zachowane(self):
        raw = "A\x00B\x07C\x1fD\tE\nF\rG"
        path = self.write([Sheet(name="C", columns=["tekst"], rows=[[raw]])])
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual(worksheet["A2"].value, "ABCD\tE\nF\rG")

    def test_polskie_znaki_w_danych_i_naglowkach(self):
        path = self.write([Sheet(name="Pł", columns=["Zażółć", "gęślą"], rows=[[POLSKIE, "ŁÓDŹ"]])])
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual(worksheet["A1"].value, "Zażółć")
        self.assertEqual(worksheet["B1"].value, "gęślą")
        self.assertEqual(worksheet["A2"].value, POLSKIE)
        self.assertEqual(worksheet["B2"].value, "ŁÓDŹ")

    def test_biale_znaki_na_brzegach_zachowane(self):
        path = self.write([Sheet(name="W", columns=["k"], rows=[["  odstęp  "], ["   "]])])
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual(worksheet["A2"].value, "  odstęp  ")
        self.assertEqual(worksheet["A3"].value, "   ")

    def test_znaki_wymagajace_escapowania_xml(self):
        raw = '<tag attr="x"> & \'apostrof\''
        path = self.write([Sheet(name="X", columns=["<kol&umna>"], rows=[[raw]])])
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual(worksheet["A1"].value, "<kol&umna>")
        self.assertEqual(worksheet["A2"].value, raw)

    def test_wartosci_nietypowych_typow_ida_jako_tekst(self):
        class Dziwny:
            def __str__(self):
                return "obiekt-dziwny"

        path = self.write([Sheet(name="O", columns=["k"], rows=[[Dziwny()], [b"bajty"]])])
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual(worksheet["A2"].value, "obiekt-dziwny")
        self.assertEqual(worksheet["A3"].value, "bajty")

    # -- kształt danych ----------------------------------------------------- #

    def test_generator_jako_zrodlo_wierszy(self):
        source = (["wiersz %d" % index, index] for index in range(5))
        path = self.write([Sheet(name="G", columns=["tekst", "nr"], rows=source)])
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual(worksheet.max_row, 6)
        self.assertEqual(worksheet["A2"].value, "wiersz 0")
        self.assertEqual(worksheet["B6"].value, 4)
        # Generator został wyczerpany — nic nie zostało w pamięci „na później”.
        self.assertEqual(list(source), [])

    def test_wiersze_krotsze_i_dluzsze_niz_naglowek(self):
        rows = [["a"], ["a", "b", "c", "nadmiar"], []]
        path = self.write([Sheet(name="R", columns=["k1", "k2", "k3"], rows=rows)])
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual(worksheet.max_column, 3)
        self.assertEqual([cell.value for cell in worksheet[2]], ["a", None, None])
        self.assertEqual([cell.value for cell in worksheet[3]], ["a", "b", "c"])
        self.assertEqual([cell.value for cell in worksheet[4]], [None, None, None])

    def test_wiersz_jako_krotka_i_dowolna_sekwencja(self):
        rows = [("a", 1), range(2)]
        path = self.write([Sheet(name="S", columns=["k1", "k2"], rows=rows)])
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual([cell.value for cell in worksheet[2]], ["a", 1])
        self.assertEqual([cell.value for cell in worksheet[3]], [0, 1])

    def test_arkusz_bez_wierszy_ma_sam_naglowek(self):
        path = self.write([Sheet(name="Puste", columns=["k1", "k2"], rows=[])])
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual(worksheet.max_row, 1)
        self.assertEqual([cell.value for cell in worksheet[1]], ["k1", "k2"])

    def test_pusta_lista_arkuszy_daje_poprawny_plik(self):
        path = self.write([])
        workbook = openpyxl.load_workbook(path)
        self.assertEqual(workbook.sheetnames, ["Arkusz"])

    # -- nazwy arkuszy w gotowym pliku -------------------------------------- #

    def test_nazwy_arkuszy_sa_czyszczone_i_unikalne(self):
        sheets = [
            Sheet(name="Aukcja: flexit/2026 [pakiet]", columns=["k"], rows=[["a"]]),
            Sheet(name="Bardzo długa nazwa arkusza, która nie mieści się w limicie",
                  columns=["k"], rows=[["b"]]),
            Sheet(name="Dane", columns=["k"], rows=[["c"]]),
            Sheet(name="Dane", columns=["k"], rows=[["d"]]),
            Sheet(name="dane", columns=["k"], rows=[["e"]]),
        ]
        path = self.write(sheets)
        names = openpyxl.load_workbook(path).sheetnames
        self.assertEqual(len(names), 5)
        self.assertEqual(len(set(name.lower() for name in names)), 5)
        for name in names:
            with self.subTest(name=name):
                self.assertLessEqual(len(name), 31)
                self.assertFalse(set(name) & set("[]:*?/\\"))
        self.assertEqual(names[0], "Aukcja_ flexit_2026 _pakiet_")
        self.assertEqual(names[2], "Dane")
        self.assertEqual(names[3], "Dane (2)")
        self.assertEqual(names[4], "dane (3)")

    def test_wiele_arkuszy_ma_wlasne_dane(self):
        sheets = [
            Sheet(name="Pierwszy", columns=["a"], rows=[["1"]]),
            Sheet(name="Drugi", columns=["b"], rows=[["2"], ["3"]]),
        ]
        path = self.write(sheets)
        workbook = openpyxl.load_workbook(path)
        self.assertEqual(workbook.sheetnames, ["Pierwszy", "Drugi"])
        self.assertEqual(workbook["Pierwszy"]["A2"].value, "1")
        self.assertEqual(workbook["Drugi"].max_row, 3)

    # -- limity Excela ------------------------------------------------------ #

    def test_podzial_po_przekroczeniu_limitu_wierszy(self):
        """Limit podmieniany przez stałą modułu ``MAX_ROWS_PER_SHEET`` (patrz docstring modułu)."""
        stderr = io.StringIO()
        rows = ([index, "w%d" % index] for index in range(13))
        with mock.patch.object(xlsxwrite, "MAX_ROWS_PER_SHEET", 6), redirect_stderr(stderr):
            path = self.write([Sheet(name="Dużo", columns=["nr", "tekst"], rows=rows)])

        workbook = openpyxl.load_workbook(path)
        self.assertEqual(workbook.sheetnames, ["Dużo", "Dużo (2)", "Dużo (3)"])
        # 6 wierszy arkusza = nagłówek + 5 wierszy danych.
        self.assertEqual([sheet.max_row for sheet in workbook.worksheets], [6, 6, 4])

        odczytane = []
        for sheet in workbook.worksheets:
            self.assertEqual([cell.value for cell in sheet[1]], ["nr", "tekst"])
            for row in sheet.iter_rows(min_row=2, values_only=True):
                odczytane.append(row)
        self.assertEqual(odczytane, [(index, "w%d" % index) for index in range(13)])

        message = stderr.getvalue()
        self.assertIn("Dużo", message)
        self.assertIn("wiersz", message)

    def test_podzial_po_przekroczeniu_limitu_kolumn(self):
        """Limit podmieniany przez stałą modułu ``MAX_COLS_PER_SHEET``."""
        stderr = io.StringIO()
        columns = ["k%d" % index for index in range(7)]
        rows = ([["r%dc%d" % (r, c) for c in range(7)] for r in range(3)])
        with mock.patch.object(xlsxwrite, "MAX_COLS_PER_SHEET", 3), redirect_stderr(stderr):
            path = self.write([Sheet(name="Szerokie", columns=columns, rows=iter(rows))])

        workbook = openpyxl.load_workbook(path)
        self.assertEqual(len(workbook.sheetnames), 3)
        self.assertEqual(
            [[cell.value for cell in sheet[1]] for sheet in workbook.worksheets],
            [["k0", "k1", "k2"], ["k3", "k4", "k5"], ["k6"]],
        )
        # Wiersze muszą być te same w każdym bloku kolumn (źródło było iteratorem!).
        scalone = []
        for row_index in range(2, 5):
            scalone.append(
                [cell.value
                 for sheet in workbook.worksheets
                 for cell in sheet[row_index]]
            )
        self.assertEqual(scalone, rows)
        self.assertIn("kolumn", stderr.getvalue())

    def test_brak_ostrzezenia_gdy_limity_nieprzekroczone(self):
        stderr = io.StringIO()
        with redirect_stderr(stderr):
            self.write([Sheet(name="Małe", columns=["a"], rows=[["x"]])])
        self.assertEqual(stderr.getvalue(), "")

    # -- poprawność pakietu OOXML ------------------------------------------ #

    def test_struktura_pakietu_ooxml(self):
        sheets = [
            Sheet(name="Pierwszy", columns=["a", "b"], rows=[["x", 1]]),
            Sheet(name="Drugi", columns=["c"], rows=[["y"]]),
        ]
        path = self.write(sheets)
        with zipfile.ZipFile(path) as archive:
            self.assertIsNone(archive.testzip(), "archiwum ZIP jest uszkodzone")
            names = set(archive.namelist())
            for required in (
                "[Content_Types].xml",
                "_rels/.rels",
                "xl/workbook.xml",
                "xl/_rels/workbook.xml.rels",
                "xl/styles.xml",
                "xl/worksheets/sheet1.xml",
                "xl/worksheets/sheet2.xml",
            ):
                self.assertIn(required, names)

            # 1. Każda część XML jest poprawnie zbudowana.
            for name in names:
                if name.endswith(".xml") or name.endswith(".rels"):
                    with self.subTest(part=name):
                        ET.fromstring(archive.read(name))

            # 2. [Content_Types].xml opisuje wszystkie części niebędące relacjami.
            types = ET.fromstring(archive.read("[Content_Types].xml"))
            ns = "{http://schemas.openxmlformats.org/package/2006/content-types}"
            defaults = {node.get("Extension").lower() for node in types.findall(ns + "Default")}
            overrides = {node.get("PartName") for node in types.findall(ns + "Override")}
            self.assertIn("rels", defaults)
            self.assertIn("xml", defaults)
            for name in names:
                if name == "[Content_Types].xml" or name.endswith(".rels"):
                    continue
                extension = name.rsplit(".", 1)[-1].lower()
                with self.subTest(part=name):
                    self.assertTrue(
                        ("/" + name) in overrides or extension in defaults,
                        "brak typu MIME dla części %s" % name,
                    )

            # 3. Relacje wskazują na istniejące części.
            rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
            root_rels = ET.fromstring(archive.read("_rels/.rels"))
            targets = [resolve_target("", node.get("Target")) for node in root_rels]
            self.assertIn("xl/workbook.xml", targets)
            for target in targets:
                self.assertIn(target, names)

            book_rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            relations = {
                node.get("Id"): resolve_target("xl", node.get("Target"))
                for node in book_rels.findall(rel_ns + "Relationship")
            }
            self.assertIn("xl/styles.xml", relations.values())
            for target in relations.values():
                self.assertIn(target, names)

            # 4. Każdy arkusz z workbook.xml ma swoją relację i plik.
            main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            rel = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
            book = ET.fromstring(archive.read("xl/workbook.xml"))
            declared = book.find(main + "sheets").findall(main + "sheet")
            self.assertEqual([node.get("name") for node in declared], ["Pierwszy", "Drugi"])
            for node in declared:
                self.assertIn(relations[node.get(rel + "id")], names)

    def test_kolejnosc_elementow_w_arkuszu_zgodna_ze_schematem(self):
        path = self.write([Sheet(name="K", columns=["a"], rows=[["x"]])])
        main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        root = ET.fromstring(sheet_xml(path, 1))
        tags = [child.tag.replace(main, "") for child in root]
        positions = []
        for tag in tags:
            self.assertIn(tag, _WORKSHEET_ORDER, "nieznany element arkusza: %s" % tag)
            positions.append(_WORKSHEET_ORDER.index(tag))
        self.assertEqual(positions, sorted(positions), "zła kolejność elementów: %s" % tags)
        self.assertIn("sheetData", tags)

    def test_deklaracja_kodowania_i_utf8(self):
        path = self.write([Sheet(name="U", columns=[POLSKIE], rows=[[POLSKIE]])])
        with zipfile.ZipFile(path) as archive:
            raw = archive.read("xl/worksheets/sheet1.xml")
        self.assertTrue(raw.startswith(b'<?xml version="1.0" encoding="UTF-8"')
                        or raw.startswith(b"<worksheet"))
        # Polskie znaki muszą przetrwać zapis w UTF-8 i dać się sparsować.
        text = raw.decode("utf-8")
        self.assertTrue(POLSKIE in text or "&#" in text)
        self.assertEqual(
            openpyxl.load_workbook(path).active["A2"].value, POLSKIE
        )
        ET.fromstring(raw)

    # -- zapis pliku -------------------------------------------------------- #

    def test_zapis_jest_powtarzalny_bajt_w_bajt(self):
        first = self.write([Sheet(name="D", columns=["a"], rows=[["x"]])], name="a.xlsx")
        time.sleep(0.01)
        second = self.write([Sheet(name="D", columns=["a"], rows=[["x"]])], name="b.xlsx")
        with open(first, "rb") as handle_a, open(second, "rb") as handle_b:
            self.assertEqual(handle_a.read(), handle_b.read())

    def test_nadpisanie_istniejacego_pliku(self):
        path = self.write([Sheet(name="D", columns=["a"], rows=[["stare"]])])
        self.write([Sheet(name="D", columns=["a"], rows=[["nowe"]])])
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual(worksheet["A2"].value, "nowe")

    def test_blad_w_trakcie_nie_niszczy_pliku_ani_nie_zostawia_smieci(self):
        path = self.write([Sheet(name="D", columns=["a"], rows=[["dobre"]])])

        def wybuchowy():
            yield ["ok"]
            raise RuntimeError("awaria źródła danych")

        with self.assertRaises(RuntimeError):
            write_workbook(path, [Sheet(name="D", columns=["a"], rows=wybuchowy())])

        # Stary plik nadal jest poprawny…
        worksheet = openpyxl.load_workbook(path).active
        self.assertEqual(worksheet["A2"].value, "dobre")
        # …a po nieudanym zapisie nie zostały pliki tymczasowe.
        leftovers = [name for name in os.listdir(self.tmpdir) if name.startswith(".flexit2xlsx-")]
        self.assertEqual(leftovers, [])

    def test_sciezka_jako_obiekt_pathlike(self):
        target = pathlib.Path(self.tmpdir) / "przez-pathlib.xlsx"
        write_workbook(target, [Sheet(name="D", columns=["a"], rows=[["x"]])])
        self.assertTrue(target.exists())
        self.assertEqual(openpyxl.load_workbook(target).active["A2"].value, "x")

    # -- wydajność i pamięć ------------------------------------------------- #

    def test_szybki_wariant_duzego_zapisu(self):
        """Skrócony wariant docelowego przypadku 200 000 x 30 (tu 20 000 x 30)."""
        rows_count, cols_count = 20000, 30
        columns = ["kol%d" % index for index in range(cols_count)]
        started = time.perf_counter()
        path = self.write(
            [Sheet(name="Duży", columns=columns, rows=make_rows(rows_count, cols_count))],
            name="duzy.xlsx",
        )
        elapsed = time.perf_counter() - started
        self.assertLess(elapsed, 120.0, "zapis 20 000 x 30 trwał %.1f s" % elapsed)

        workbook = openpyxl.load_workbook(path, read_only=True)
        try:
            worksheet = workbook.worksheets[0]
            rows = worksheet.iter_rows(values_only=True)
            self.assertEqual(next(rows)[:2], ("kol0", "kol1"))
            self.assertEqual(next(rows)[:3], ("tekst 0-0", 0, 0.0))
            counted = 2 + sum(1 for _ in rows)
        finally:
            workbook.close()
        self.assertEqual(counted, rows_count + 1)

    def test_zuzycie_pamieci_nie_rosnie_z_liczba_wierszy(self):
        """Dowód strumieniowości: szczyt alokacji jest niemal taki sam dla 1 000 i 5 000 wierszy."""
        peaks = {}
        for count in (1000, 5000):
            tracemalloc.start()
            try:
                self.write(
                    [Sheet(name="P", columns=["k%d" % i for i in range(30)],
                           rows=make_rows(count, 30))],
                    name="pamiec-%d.xlsx" % count,
                )
                peaks[count] = tracemalloc.get_traced_memory()[1]
            finally:
                tracemalloc.stop()
        self.assertLess(
            peaks[5000],
            peaks[1000] * 1.8,
            "szczyt pamięci rośnie z liczbą wierszy: %r" % (peaks,),
        )


@unittest.skipIf(openpyxl is None, "openpyxl niedostępny")
class TestStdlibBackend(_BackendCase, unittest.TestCase):
    """Wbudowany zapis OOXML — wyłącznie biblioteka standardowa."""

    BACKEND = "stdlib"

    def test_brak_sharedstrings_napisy_sa_osadzone(self):
        path = self.write([Sheet(name="S", columns=["k"], rows=[["tekst"]])])
        with zipfile.ZipFile(path) as archive:
            self.assertNotIn("xl/sharedStrings.xml", archive.namelist())
        self.assertIn('t="inlineStr"', sheet_xml(path, 1))

    def test_znaczniki_czasu_w_archiwum_sa_stale(self):
        path = self.write([Sheet(name="S", columns=["k"], rows=[["x"]])])
        with zipfile.ZipFile(path) as archive:
            stamps = {info.date_time for info in archive.infolist()}
        self.assertEqual(stamps, {(1980, 1, 1, 0, 0, 0)})


@unittest.skipIf(openpyxl is None, "openpyxl niedostępny")
class TestOpenpyxlBackend(_BackendCase, unittest.TestCase):
    """Zapis przez openpyxl w trybie strumieniowym (``write_only``)."""

    BACKEND = "openpyxl"


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
