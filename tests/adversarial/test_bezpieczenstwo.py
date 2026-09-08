# -*- coding: utf-8 -*-
"""Testy ADWERSARYJNE: przegląd bezpieczeństwa całego pakietu ``flexit2xlsx``.

Zakres ataku: XXE i bomby encyjne, path traversal i zip-slip przy zapisie,
wstrzyknięcie formuły do Excela, SSRF przez ``--base-url``, wstrzyknięcie
nagłówków HTTP, wyciek ciasteczka sesyjnego przy przekierowaniu, nadpisywanie
plików poza katalogiem docelowym oraz zużycie pamięci jako DoS.

Wszystkie testy (``test_ok_*``) potwierdzają ochrony, które MAJĄ działać,
i chronią przed regresją.  Historycznie moduł zawierał też testy ``test_blad_*``
utrwalające zaobserwowane wady — po ich naprawieniu każdy został przepisany na
asercję stanu poprawnego.

Uruchamianie (w katalogu repozytorium):

    python3 -m unittest discover -s tests/adversarial -t . -v
    python3 tests/adversarial/test_bezpieczenstwo.py
"""

from __future__ import annotations

import contextlib
import http.server
import io
import os
import socketserver
import sys
import tempfile
import threading
import tracemalloc
import unittest
import urllib.request
import zipfile

_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

from flexit2xlsx import cli, scrape, values, xlsxwrite  # noqa: E402
from flexit2xlsx import xmlflatten as xf  # noqa: E402

try:  # openpyxl i lxml służą WYŁĄCZNIE do weryfikacji wyników
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

try:
    from lxml import etree as _lxml_etree
except ImportError:  # pragma: no cover
    _lxml_etree = None


# --------------------------------------------------------------------------- #
# Narzędzia pomocnicze
# --------------------------------------------------------------------------- #


class _Serwer:
    """Malutki serwer HTTP na 127.0.0.1 z losowym portem (do testów sieciowych).

    ``routes`` to mapa ``ścieżka -> (kod, nagłówki, treść)``.  Każde żądanie jest
    zapisywane w ``self.trafienia`` razem z otrzymanymi nagłówkami, dzięki czemu
    test może sprawdzić, czy wyciekło ciasteczko.
    """

    def __init__(self, routes):
        self.routes = routes
        self.trafienia = []
        serwer = self

        class Handler(http.server.BaseHTTPRequestHandler):
            protocol_version = "HTTP/1.0"

            def log_message(self, *args):  # cisza w testach
                pass

            def do_GET(self):  # noqa: N802 - nazwa narzucona przez stdlib
                sciezka = self.path.split("?", 1)[0]
                serwer.trafienia.append((sciezka, dict(self.headers)))
                trasa = serwer.routes.get(sciezka)
                if trasa is None:
                    self.send_response(404)
                    self.send_header("Content-Length", "0")
                    self.end_headers()
                    return
                kod, naglowki, tresc = trasa
                if callable(tresc):
                    tresc = tresc()
                self.send_response(kod)
                for klucz, wartosc in naglowki.items():
                    self.send_header(klucz, wartosc)
                self.send_header("Content-Length", str(len(tresc)))
                self.end_headers()
                if tresc:
                    self.wfile.write(tresc)

        self._srv = socketserver.TCPServer(("127.0.0.1", 0), Handler)
        self.port = self._srv.server_address[1]
        self.base = "http://127.0.0.1:%d/" % self.port
        self.base_localhost = "http://localhost:%d/" % self.port
        self._watek = threading.Thread(target=self._srv.serve_forever, daemon=True)
        self._watek.start()

    def zamknij(self):
        self._srv.shutdown()
        self._srv.server_close()


@contextlib.contextmanager
def _backend(nazwa):
    """Wymusza ścieżkę zapisu XLSX (``stdlib``/``openpyxl``) na czas bloku."""
    poprzedni = os.environ.get(xlsxwrite.ENV_BACKEND)
    os.environ[xlsxwrite.ENV_BACKEND] = nazwa
    try:
        yield
    finally:
        if poprzedni is None:
            os.environ.pop(xlsxwrite.ENV_BACKEND, None)
        else:
            os.environ[xlsxwrite.ENV_BACKEND] = poprzedni


def _xml_glebokie(glebokosc, liscie):
    """Poprawny XML: ``glebokosc`` zagnieżdżeń, a na dnie ``liscie`` różnych pól.

    Bez encji, bez DTD, bez rekurencji — czysty, w pełni legalny dokument.
    """
    czesci = ["<?xml version='1.0'?>"]
    czesci += ["<n%d>" % i for i in range(glebokosc)]
    czesci += ["<f%d>v</f%d>" % (j, j) for j in range(liscie)]
    czesci += ["</n%d>" % i for i in reversed(range(glebokosc))]
    return "".join(czesci).encode("utf-8")


class _BazaZKatalogiem(unittest.TestCase):
    """Wspólny katalog tymczasowy dla testów operujących na plikach."""

    def setUp(self):
        self.katalog = tempfile.mkdtemp(prefix="flexit-bezp-")
        self.addCleanup(self._sprzataj)

    def _sprzataj(self):
        import shutil

        shutil.rmtree(self.katalog, ignore_errors=True)

    def sciezka(self, nazwa):
        return os.path.join(self.katalog, nazwa)

    def zapisz(self, nazwa, tresc):
        pelna = self.sciezka(nazwa)
        os.makedirs(os.path.dirname(pelna), exist_ok=True)
        tryb = "wb" if isinstance(tresc, bytes) else "w"
        with open(pelna, tryb) as uchwyt:
            uchwyt.write(tresc)
        return pelna


# --------------------------------------------------------------------------- #
# 1. XXE i encje zewnętrzne
# --------------------------------------------------------------------------- #


class TestXXE(_BazaZKatalogiem):
    """Encje zewnętrzne i odwołania do zasobów spoza dokumentu."""

    def test_ok_encja_systemowa_nie_czyta_pliku(self):
        """``<!ENTITY x SYSTEM "file://...">`` musi zostać zablokowana."""
        tajne = self.zapisz("tajne.txt", "HASLO-DO-BANKU-9911")
        doc = (
            "<?xml version='1.0'?><!DOCTYPE r [<!ENTITY x SYSTEM 'file://%s'>]>"
            "<r><a>&x;</a></r>" % tajne
        ).encode("utf-8")
        with self.assertRaises(xf.XmlParseError) as ctx:
            xf.parse_bytes(doc, "xxe.xml")
        self.assertIn("XXE", str(ctx.exception))
        self.assertNotIn("HASLO-DO-BANKU", str(ctx.exception))

    def test_ok_encja_publiczna_zablokowana(self):
        doc = (
            b"<?xml version='1.0'?><!DOCTYPE r ["
            b"<!ENTITY x PUBLIC '-//x//EN' 'file:///etc/passwd'>]><r><a>&x;</a></r>"
        )
        with self.assertRaises(xf.XmlParseError):
            xf.parse_bytes(doc, "xxe2.xml")

    def test_ok_encja_parametryczna_zablokowana(self):
        doc = (
            b"<?xml version='1.0'?><!DOCTYPE r ["
            b"<!ENTITY % p SYSTEM 'file:///etc/passwd'> %p;]><r><a>b</a></r>"
        )
        with self.assertRaises(xf.XmlParseError):
            xf.parse_bytes(doc, "xxe3.xml")

    def test_ok_zewnetrzny_dtd_nie_jest_pobierany(self):
        """DTD wskazane przez SYSTEM na http:// nie może wywołać żądania sieciowego."""
        srv = _Serwer({"/zly.dtd": (200, {"Content-Type": "application/xml-dtd"},
                                    b"<!ENTITY x 'WSTRZYKNIETE'>")})
        self.addCleanup(srv.zamknij)
        doc = (
            "<?xml version='1.0'?><!DOCTYPE r SYSTEM '%szly.dtd'><r><a>ok</a></r>"
            % srv.base
        ).encode("utf-8")
        doc_wynik = xf.parse_bytes(doc, "dtd.xml")
        self.assertEqual(srv.trafienia, [], "parser NIE może pobierać zewnętrznego DTD")
        self.assertEqual(doc_wynik.records[0].get("a"), "ok")

    def test_ok_xxe_przez_cli_nie_wycieka_do_xlsx(self):
        """Pełna ścieżka ``build``: plik z XXE jest pomijany, a nie wczytywany."""
        tajne = self.zapisz("sekret.txt", "TOKEN-AWS-ABCDEF")
        wejscie = os.path.join(self.katalog, "we")
        os.makedirs(wejscie, exist_ok=True)
        with open(os.path.join(wejscie, "atak.xml"), "w") as uchwyt:
            uchwyt.write(
                "<?xml version='1.0'?><!DOCTYPE r [<!ENTITY x SYSTEM 'file://%s'>]>"
                "<r><item><a>&x;</a></item><item><a>2</a></item></r>" % tajne
            )
        wyjscie = self.sciezka("wynik.xlsx")
        bufor = io.StringIO()
        with contextlib.redirect_stdout(bufor), contextlib.redirect_stderr(bufor):
            kod = cli.main(["build", "--in", wejscie, "--out", wyjscie, "-q"])
        self.assertEqual(kod, cli.EXIT_NO_DATA)
        self.assertFalse(os.path.exists(wyjscie))
        self.assertNotIn("TOKEN-AWS", bufor.getvalue())


# --------------------------------------------------------------------------- #
# 2. Bomby encyjne
# --------------------------------------------------------------------------- #


class TestBombyEncyjne(unittest.TestCase):
    """Rozwijanie encji nie może wysadzić pamięci."""

    def test_ok_billion_laughs(self):
        doc = (
            b"<?xml version='1.0'?><!DOCTYPE r ["
            b"<!ENTITY a 'aaaaaaaaaa'>"
            b"<!ENTITY b '&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;'>"
            b"<!ENTITY c '&b;&b;&b;&b;&b;&b;&b;&b;&b;&b;'>"
            b"<!ENTITY d '&c;&c;&c;&c;&c;&c;&c;&c;&c;&c;'>"
            b"]><r><a>&d;</a></r>"
        )
        with self.assertRaises(xf.XmlParseError) as ctx:
            xf.parse_bytes(doc, "bomba.xml")
        self.assertIn("bomb", str(ctx.exception).lower())

    def test_ok_bomba_kwadratowa(self):
        """Jedna duża encja użyta tysiące razy — limit amplifikacji tekstu."""
        doc = (
            "<?xml version='1.0'?><!DOCTYPE r [<!ENTITY a '%s'>]><r>%s</r>"
            % ("X" * 1000, "<a>&a;</a>" * 30000)
        ).encode("utf-8")
        with self.assertRaises(xf.XmlParseError) as ctx:
            xf.parse_bytes(doc, "kwadrat.xml")
        self.assertIn("limit", str(ctx.exception).lower())

    def test_ok_zbyt_wiele_deklaracji_encji(self):
        """Limit liczby deklaracji odsiewa bombę, ale przepuszcza zwykłe pliki.

        Eksporty z prawdziwych systemów potrafią deklarować kilkaset encji
        (np. podzbiór encji HTML), więc próg musi być hojny — odrzucamy dopiero
        powyżej :data:`xf._MAX_ENTITY_DECLS`.
        """
        def dokument(ile):
            deklaracje = "".join("<!ENTITY e%d 'x'>" % i for i in range(ile))
            return ("<?xml version='1.0'?><!DOCTYPE r [%s]><r><a>1</a></r>"
                    % deklaracje).encode("utf-8")

        doc = xf.parse_bytes(dokument(300), "duzo.xml")
        self.assertEqual(doc.records[0]["a"], "1")

        with self.assertRaises(xf.XmlParseError) as ctx:
            xf.parse_bytes(dokument(xf._MAX_ENTITY_DECLS + 100), "bomba.xml")
        self.assertIn("deklaracji encji", str(ctx.exception))

# --------------------------------------------------------------------------- #
# 3. Awaryjna naprawa encji nie może ruszać sekcji CDATA
# --------------------------------------------------------------------------- #


class TestNaprawaEncjiPsujeCDATA(unittest.TestCase):
    """Awaryjna naprawa encji NIE MOŻE ruszać sekcji CDATA.

    Gdy w dokumencie wystąpi choć JEDNA nieznana encja (np. HTML-owe ``&nbsp;``
    z eksportu portalu), moduł ponawia parsowanie po podmianie encji — ale
    wyłącznie POZA sekcjami ``<![CDATA[...]]>``, komentarzami i instrukcjami
    przetwarzania, gdzie ``&nazwa;`` jest zwykłym tekstem, a nie encją.
    Dodatkowo o każdej naprawie użytkownik jest informowany na stderr.
    """

    DOC_Z_NIEZNANA_ENCJA = (
        "<?xml version='1.0' encoding='UTF-8'?>"
        "<lot><opis>Cena&nbsp;netto</opis><items>"
        "<item><sku>A1</sku>"
        "<note><![CDATA[Firma &raquo; model &kod_producenta; koniec]]></note></item>"
        "<item><sku>A2</sku><note><![CDATA[R&D 100]]></note></item>"
        "</items></lot>"
    )

    def _bez_nieznanej_encji(self):
        return self.DOC_Z_NIEZNANA_ENCJA.replace("Cena&nbsp;netto", "Cena netto")

    def test_ok_cdata_nietkniete_gdy_plik_jest_poprawny(self):
        """Kontrola: bez nieznanej encji CDATA zostaje dosłownie takie, jakie było."""
        doc = xf.parse_bytes(self._bez_nieznanej_encji().encode("utf-8"), "ok.xml")
        self.assertEqual(doc.records[0]["note"],
                         "Firma &raquo; model &kod_producenta; koniec")

    def test_ok_cdata_nietkniete_mimo_naprawy_encji(self):
        """Naprawa ``&nbsp;`` poza CDATA nie zmienia treści WEWNĄTRZ CDATA."""
        with contextlib.redirect_stderr(io.StringIO()):
            doc = xf.parse_bytes(self.DOC_Z_NIEZNANA_ENCJA.encode("utf-8"), "psuje.xml")
        self.assertEqual(doc.records[0]["note"],
                         "Firma &raquo; model &kod_producenta; koniec")
        self.assertEqual(doc.records[1]["note"], "R&D 100")
        self.assertEqual(doc.context["lot/opis"], "Cena netto")   # &nbsp; POZA CDATA

    def test_ok_naprawa_encji_jest_zgloszona_uzytkownikowi(self):
        """Żaden fragment nie znika po cichu — stderr dostaje ostrzeżenie."""
        blad = io.StringIO()
        with contextlib.redirect_stderr(blad):
            doc = xf.parse_bytes(self.DOC_Z_NIEZNANA_ENCJA.encode("utf-8"), "psuje.xml")
        self.assertIn("kod_producenta", doc.records[0]["note"])
        komunikat = blad.getvalue()
        self.assertIn("encje", komunikat)
        self.assertIn("CDATA", komunikat)

    def test_ok_tresc_cdata_dociera_nietknieta_do_gotowego_xlsx(self):
        """Ta sama treść przechodzi przez cały potok aż do arkusza."""
        if openpyxl is None:  # pragma: no cover
            self.skipTest("openpyxl niedostępny")
        katalog = tempfile.mkdtemp(prefix="flexit-cdata-")
        self.addCleanup(lambda: __import__("shutil").rmtree(katalog, ignore_errors=True))
        wejscie = os.path.join(katalog, "we")
        os.makedirs(wejscie)
        with open(os.path.join(wejscie, "lot.xml"), "w", encoding="utf-8") as uchwyt:
            uchwyt.write(self.DOC_Z_NIEZNANA_ENCJA)
        wyjscie = os.path.join(katalog, "w.xlsx")
        bufor = io.StringIO()
        with contextlib.redirect_stdout(bufor), contextlib.redirect_stderr(bufor):
            kod = cli.main(["build", "--in", wejscie, "--out", wyjscie, "-q"])
        self.assertEqual(kod, cli.EXIT_OK)
        skoroszyt = openpyxl.load_workbook(wyjscie)
        try:
            arkusz = skoroszyt[cli.SHEET_ALL]
            naglowki = [komorka.value for komorka in arkusz[1]]
            kolumna = naglowki.index("note") + 1
            wartosc = arkusz.cell(row=2, column=kolumna).value
        finally:
            skoroszyt.close()
        self.assertEqual(wartosc, "Firma &raquo; model &kod_producenta; koniec")

# --------------------------------------------------------------------------- #
# 4. Zużycie pamięci jako DoS
# --------------------------------------------------------------------------- #


class TestBombaPamieciowa(unittest.TestCase):
    """Mały, poprawny XML nie może zjadać setek MB pamięci.

    Dawniej ``xmlflatten._scan`` budował dla KAŻDEJ ścieżki zbiór wszystkich
    względnych ścieżek pól potomnych, sklejając je łańcuchowo w górę drzewa —
    koszt był iloczynem głębokości i szerokości, więc plik 23 kB zajmował
    kilkadziesiąt MB, a plik 191 kB przekraczał 900 MB.  Dziś zużycie pamięci
    jest proporcjonalne do rozmiaru pliku.
    """

    def _zmierz(self, dane):
        """Zwraca szczytowe zużycie pamięci przy parsowaniu ``dane``."""
        tracemalloc.start()
        try:
            with contextlib.redirect_stderr(io.StringIO()):
                xf.parse_bytes(dane, "bomba.xml")
            _, szczyt = tracemalloc.get_traced_memory()
        finally:
            tracemalloc.stop()
        return szczyt

    def test_ok_maly_plik_nie_zjada_dziesiatek_MB(self):
        dane = _xml_glebokie(120, 1500)
        self.assertLess(len(dane), 30 * 1024, "plik wejściowy jest naprawdę mały")
        szczyt = self._zmierz(dane)
        krotnosc = szczyt / len(dane)
        self.assertLess(
            szczyt, 20 * 1024 * 1024,
            "plik %d B zajął %.1f MB (%.0f-krotność rozmiaru)"
            % (len(dane), szczyt / 1e6, krotnosc),
        )
        self.assertLess(
            krotnosc, 500,
            "zużycie pamięci ma być proporcjonalne do rozmiaru pliku "
            "(rząd kilkudziesięciu krotności), a nie tysiące krotności",
        )

    def test_ok_zuzycie_nie_rosnie_kwadratowo_z_szerokoscia(self):
        """Czterokrotnie szerszy plik nie może zająć wielokrotnie więcej pamięci."""
        maly = _xml_glebokie(120, 1500)
        duzy = _xml_glebokie(120, 6000)
        szczyt_maly = self._zmierz(maly)
        szczyt_duzy = self._zmierz(duzy)
        wzrost_danych = len(duzy) / len(maly)
        wzrost_pamieci = szczyt_duzy / max(szczyt_maly, 1)
        self.assertLess(
            wzrost_pamieci, wzrost_danych * 3,
            "pamięć urosła %.1f-krotnie przy %.1f-krotnym wzroście danych"
            % (wzrost_pamieci, wzrost_danych),
        )

    def test_ok_glebokosc_ponad_limit_jest_obcinana_z_ostrzezeniem(self):
        """Zbyt głębokie zagnieżdżenie nie kasuje całego pliku — tylko nadmiar.

        Odrzucenie CAŁEGO dokumentu oznaczałoby utratę wszystkich pozycji przez
        jedno patologicznie zagnieżdżone pole; zamiast tego obcinamy elementy
        poniżej :data:`xf._MAX_DEPTH` i głośno o tym mówimy na stderr.
        """
        dane = _xml_glebokie(400, 2)
        blad = io.StringIO()
        with contextlib.redirect_stderr(blad):
            doc = xf.parse_bytes(dane, "gleboki.xml")
        self.assertIn("głębiej niż %d" % xf._MAX_DEPTH, blad.getvalue())
        self.assertIn("obcięto", blad.getvalue())
        self.assertTrue(doc.records, "płytsze pola muszą zostać wczytane")

# --------------------------------------------------------------------------- #
# 5. Nazwy plików: path traversal, symlinki, nadpisywanie poza katalogiem
# --------------------------------------------------------------------------- #


def _fetch_atrapa(_opener, _url):
    return b"<?xml version='1.0'?><r><a>1</a></r>", "application/xml"


class _OpenerAtrapa:
    """Minimalny obiekt opener wystarczający dla ``download_xml``."""

    def __init__(self):
        self.stats = {"requests": 0, "bytes": 0, "probes": 0}
        self._ctype_cache = {}


class TestNazwyPlikow(_BazaZKatalogiem):
    """``safe_filename`` i ``download_xml`` nie mogą wyjść poza katalog docelowy."""

    ZLOSLIWE = [
        "../../../../etc/passwd",
        "..%2f..%2fetc%2fshadow",
        "....//....//evil.xml",
        "/etc/cron.d/backdoor",
        "C:\\Windows\\System32\\evil.xml",
        "..\\..\\evil.xml",
        "%2e%2e%2f%2e%2e%2fevil.xml",
        "a\x00/../../evil.xml",
        "\u202egnp.xml",
        ".bashrc",
        "  ..  ",
    ]

    def test_ok_safe_filename_zawsze_daje_sama_nazwe(self):
        for zly in self.ZLOSLIWE:
            with self.subTest(nazwa=zly):
                wynik = scrape.safe_filename(zly)
                self.assertNotIn("/", wynik)
                self.assertNotIn("\\", wynik)
                self.assertNotIn("..", wynik)
                self.assertFalse(wynik.startswith("."))
                self.assertTrue(wynik.endswith(".xml"))
                self.assertEqual(wynik, os.path.basename(wynik))

    def test_ok_download_xml_nie_wychodzi_poza_katalog(self):
        cel = self.sciezka("pobrane")
        os.makedirs(cel, exist_ok=True)
        korzen = os.path.realpath(cel)
        for zly in self.ZLOSLIWE:
            with self.subTest(nazwa=zly):
                ref = scrape.XmlRef(url="http://przyklad.test/x",
                                    filename=zly, auction_id="A1", auction_title="t")
                sciezka = scrape.download_xml(_OpenerAtrapa(), ref, cel,
                                              overwrite=True, fetch_fn=_fetch_atrapa)
                self.assertEqual(os.path.dirname(os.path.realpath(sciezka)), korzen)

    def test_ok_symlink_w_katalogu_docelowym_jest_odrzucany(self):
        """Portal nie może użyć istniejącego dowiązania do nadpisania cudzego pliku."""
        cel = self.sciezka("pobrane2")
        os.makedirs(cel, exist_ok=True)
        ofiara = self.zapisz("wazny.txt", "ORYGINALNA-TRESC")
        os.symlink(ofiara, os.path.join(cel, "podstep.xml"))
        ref = scrape.XmlRef(url="http://przyklad.test/x", filename="podstep.xml",
                            auction_id="A1", auction_title="t")
        with self.assertRaises(scrape.ScrapeError):
            scrape.download_xml(_OpenerAtrapa(), ref, cel, overwrite=True,
                                fetch_fn=_fetch_atrapa)
        with open(ofiara) as uchwyt:
            self.assertEqual(uchwyt.read(), "ORYGINALNA-TRESC")

    def test_ok_nazwa_z_portalu_nie_moze_udawac_pliku_czesciowego(self):
        """Nazwa ``x.xml.part`` nie może kolidować z plikiem roboczym pobierania."""
        self.assertTrue(scrape.safe_filename("x.xml.part").endswith(".xml"))
        self.assertNotEqual(scrape.safe_filename("x.xml.part"), "x.xml.part")


class TestStrukturaZip(_BazaZKatalogiem):
    """Wygenerowany ``.xlsx`` to archiwum ZIP — nie może zawierać zip-slipu."""

    def test_ok_brak_sciezek_wychodzacych_z_archiwum(self):
        arkusze = [
            xlsxwrite.Sheet(name="../../zly", columns=["a"], rows=[["1"]]),
            xlsxwrite.Sheet(name="/abs", columns=["b"], rows=[["2"]]),
        ]
        for nazwa_backendu in ("stdlib", "openpyxl"):
            if nazwa_backendu == "openpyxl" and openpyxl is None:  # pragma: no cover
                continue
            with self.subTest(backend=nazwa_backendu), _backend(nazwa_backendu):
                cel = self.sciezka("zip_%s.xlsx" % nazwa_backendu)
                xlsxwrite.write_workbook(cel, arkusze)
                with zipfile.ZipFile(cel) as archiwum:
                    for wpis in archiwum.namelist():
                        self.assertFalse(wpis.startswith("/"), wpis)
                        self.assertFalse(wpis.startswith("\\"), wpis)
                        self.assertNotIn("..", wpis)
                        self.assertEqual(
                            os.path.normpath(os.path.join("/root", wpis)).startswith(
                                "/root/"), True, wpis)


# --------------------------------------------------------------------------- #
# 6. SSRF, zakres witryny, wyciek ciasteczka
# --------------------------------------------------------------------------- #


class TestZakresWitrynyISSRF(unittest.TestCase):
    """``--base-url`` i linki ze strony nie mogą wyprowadzić poza http(s)/witrynę."""

    ZLE_SCHEMATY = [
        "file:///etc/passwd",
        "ftp://127.0.0.1/x",
        "gopher://127.0.0.1:11211/_x",
        "javascript:alert(1)",
        "data:text/html;base64,PHNjcmlwdD4=",
        "jav\tascript:alert(1)",
    ]

    def test_ok_absolutize_odrzuca_niebezpieczne_schematy(self):
        for zly in self.ZLE_SCHEMATY:
            with self.subTest(url=zly):
                self.assertIsNone(scrape.absolutize("https://portal.test/", zly))

    def test_ok_opener_nie_pobiera_spoza_http(self):
        opener = scrape.make_opener(delay=0, retries=0)
        for zly in ("file:///etc/passwd", "ftp://127.0.0.1/", "gopher://127.0.0.1/"):
            with self.subTest(url=zly):
                self.assertFalse(opener.allows(zly))

    def test_ok_cli_odmawia_pobierania_z_file(self):
        katalog = tempfile.mkdtemp(prefix="flexit-ssrf-")
        self.addCleanup(lambda: __import__("shutil").rmtree(katalog, ignore_errors=True))
        bufor = io.StringIO()
        with contextlib.redirect_stdout(bufor), contextlib.redirect_stderr(bufor):
            kod = cli.main(["download", "--base-url", "file:///etc/",
                            "--out", os.path.join(katalog, "xml"),
                            "--delay", "0", "--retries", "0"])
        self.assertEqual(kod, cli.EXIT_NETWORK)
        self.assertIn("zły schemat", bufor.getvalue())

    def test_ok_link_na_obca_domene_jest_odfiltrowany(self):
        self.assertFalse(scrape.is_same_site("https://flexitauctions.com/",
                                             "https://evil.example/x.xml"))
        self.assertTrue(scrape.is_same_site("https://flexitauctions.com/",
                                            "https://media.flexitauctions.com/x.xml"))

    def test_ok_domena_publicznego_sufiksu_nie_jest_ta_sama_witryna(self):
        """Dwuczłonowy sufiks publiczny nie może zlewać obcych domen w jedną.

        Gdyby ``is_same_site`` brało po prostu dwie ostatnie etykiety hosta, to
        dla portalu pod ``sklep.example.co.uk`` KAŻDA domena w ``.co.uk``
        wyglądałaby jak ta sama witryna — i dostawałaby ciasteczko sesyjne.
        """
        pary = [
            ("https://sklep.example.co.uk/", "https://evil.co.uk/kradnij"),
            ("https://aukcje.com.pl/", "https://zlodziej.com.pl/kradnij"),
            ("https://przetargi.gov.pl/", "https://falszywy.gov.pl/kradnij"),
            ("https://klient.github.io/", "https://atakujacy.github.io/kradnij"),
        ]
        for baza, obcy in pary:
            with self.subTest(baza=baza):
                self.assertFalse(
                    scrape.is_same_site(baza, obcy),
                    "%s NIE jest tą samą witryną co %s" % (obcy, baza),
                )

    def test_ok_poddomena_tej_samej_witryny_jest_dozwolona(self):
        """Kontrola: prawdziwa poddomena portalu nadal jest „tą samą witryną''."""
        self.assertTrue(scrape.is_same_site("https://sklep.example.co.uk/",
                                            "https://cdn.sklep.example.co.uk/x.xml"))
        self.assertTrue(scrape.is_same_site("https://aukcje.com.pl/",
                                            "https://aukcje.com.pl/plik.xml"))

    def test_ok_inny_port_to_inna_witryna(self):
        """Inny port to inne źródło (origin) — ciasteczko tam nie trafia."""
        self.assertFalse(scrape.is_same_site("http://intranet:8080/",
                                             "http://intranet:9200/"))
        self.assertTrue(scrape.is_same_site("http://intranet:8080/",
                                            "http://intranet:8080/dane.xml"))

class TestWyciekCiasteczka(unittest.TestCase):
    """Ciasteczko sesyjne z ``--cookie`` a przekierowania i logi."""

    def test_ok_przekierowanie_na_obca_domene_zablokowane(self):
        srv = _Serwer({"/": (302, {"Location": "http://evil.example.org/kradnij"}, b"")})
        self.addCleanup(srv.zamknij)
        opener = scrape.make_opener(cookie="SESSIONID=TAJNE", delay=0, retries=0)
        opener.bind_site(srv.base)
        with self.assertRaises(scrape.SiteBlockedError):
            scrape.fetch(opener, srv.base)

    def test_ok_przekierowanie_na_file_zablokowane(self):
        srv = _Serwer({"/": (302, {"Location": "file:///etc/passwd"}, b"")})
        self.addCleanup(srv.zamknij)
        opener = scrape.make_opener(cookie="SESSIONID=TAJNE", delay=0, retries=0)
        opener.bind_site(srv.base)
        with self.assertRaises(scrape.ScrapeError):
            scrape.fetch(opener, srv.base)

    def test_ok_przekierowanie_308_tez_pilnowane(self):
        self.assertTrue(hasattr(urllib.request.HTTPRedirectHandler, "http_error_308"))
        srv = _Serwer({"/": (308, {"Location": "http://evil.example.org/kradnij"}, b"")})
        self.addCleanup(srv.zamknij)
        opener = scrape.make_opener(cookie="SESSIONID=TAJNE", delay=0, retries=0)
        opener.bind_site(srv.base)
        with self.assertRaises(scrape.SiteBlockedError):
            scrape.fetch(opener, srv.base)

    def test_ok_ciasteczko_nie_wycieka_na_inny_port(self):
        """Przekierowanie na inny port TEGO SAMEGO hosta jest zablokowane.

        Serwer „portalu'' przekierowuje na ``http://localhost:<inny port>/``.
        Inny port to inne źródło, więc strażnik przekierowań musi je odrzucić,
        zanim urllib przeniesie nagłówek ``Cookie`` pod nowy adres.
        """
        ofiara = _Serwer({"/kradnij": (200, {"Content-Type": "application/xml"},
                                       b"<?xml version='1.0'?><r/>")})
        self.addCleanup(ofiara.zamknij)
        portal = _Serwer({"/": (302, {"Location": "http://localhost:%d/kradnij"
                                                  % ofiara.port}, b"")})
        self.addCleanup(portal.zamknij)

        opener = scrape.make_opener(cookie="SESSIONID=TAJNY-TOKEN-UZYTKOWNIKA",
                                    delay=0, retries=0)
        opener.bind_site(portal.base_localhost)
        with self.assertRaises(scrape.ScrapeError) as pulapka:
            scrape.fetch(opener, portal.base_localhost)
        self.assertIn("poza witrynę", str(pulapka.exception))

        naglowki_ofiary = [nagl for sciezka, nagl in ofiara.trafienia
                           if sciezka == "/kradnij"]
        self.assertEqual(naglowki_ofiary, [],
                         "ofiara nie mogła dostać ŻADNEGO żądania")

    def test_ok_ciasteczko_nie_trafia_do_logow_ani_do_arkusza(self):
        """Pełny przebieg ``all`` z ``--cookie --diagnose``: brak sekretu w wyjściu."""
        tajne = "SESSIONID=SUPER-TAJNY-TOKEN-XYZ"
        strony = {
            "/": (200, {"Content-Type": "text/html"},
                  b"<html><a href='/auction/flexit-auctions-01-01-2026-1103/'>A</a></html>"),
            "/auction/flexit-auctions-01-01-2026-1103/":
                (200, {"Content-Type": "text/html"},
                 b"<html><a href='/lot/laptopy-abcd1'>Lot</a></html>"),
            "/lot/laptopy-abcd1": (200, {"Content-Type": "text/html"},
                                   b"<html><a href='/api/batch.xml' download='batch.xml'>"
                                   b"Download Batch Details</a></html>"),
            "/api/batch.xml": (200, {"Content-Type": "application/xml"},
                               b"<?xml version='1.0'?><batch>"
                               b"<item><sku>1</sku></item><item><sku>2</sku></item>"
                               b"</batch>"),
        }
        srv = _Serwer(strony)
        self.addCleanup(srv.zamknij)
        katalog = tempfile.mkdtemp(prefix="flexit-cookie-")
        self.addCleanup(lambda: __import__("shutil").rmtree(katalog, ignore_errors=True))
        wyjscie = os.path.join(katalog, "wynik.xlsx")
        bufor = io.StringIO()
        with contextlib.redirect_stdout(bufor), contextlib.redirect_stderr(bufor):
            kod = cli.main(["all", "--base-url", srv.base, "--cookie", tajne,
                            "--in", os.path.join(katalog, "xml"), "--out", wyjscie,
                            "--overwrite", "--delay", "0", "--retries", "0",
                            "--diagnose"])
        self.assertEqual(kod, cli.EXIT_OK, bufor.getvalue())
        self.assertNotIn("TAJNY", bufor.getvalue())
        with open(wyjscie, "rb") as uchwyt:
            self.assertNotIn(b"TAJNY", uchwyt.read())
        # ciasteczko FAKTYCZNIE poszło do portalu (inaczej test nic nie dowodzi)
        self.assertTrue(any(nagl.get("Cookie") == tajne for _s, nagl in srv.trafienia))


# --------------------------------------------------------------------------- #
# 7. Wstrzyknięcie nagłówków HTTP
# --------------------------------------------------------------------------- #


class TestNaglowkiHttp(unittest.TestCase):
    """``--user-agent`` i ``--cookie`` a wstrzyknięcie CRLF."""

    def _serwer_html(self):
        srv = _Serwer({"/": (200, {"Content-Type": "text/html"},
                             b"<html><a href='/auction/x-1103/'>a</a></html>")})
        self.addCleanup(srv.zamknij)
        return srv

    def test_ok_crlf_w_user_agent_jest_odrzucony_z_komunikatem(self):
        """Próba wstrzyknięcia nagłówka przez ``--user-agent`` kończy się od razu.

        Wartość jest sprawdzana w ``make_opener``, więc żadne żądanie nie
        wychodzi, a użytkownik dostaje czytelny komunikat po polsku zamiast
        surowego ``ValueError`` z ``http.client``.
        """
        srv = self._serwer_html()
        with self.assertRaises(scrape.ScrapeError) as pulapka:
            scrape.make_opener(user_agent="UA\r\nX-Wstrzykniete: 1",
                               delay=0, retries=0)
        self.assertIn("końca linii", str(pulapka.exception))
        self.assertEqual(srv.trafienia, [], "żadne żądanie nie mogło wyjść")

    def test_ok_ciasteczko_z_koncowym_enterem_jest_przycinane(self):
        """USTALENIE (ważny): ciasteczko wklejone z przeglądarki bywa z ``\\n``.

        Nadmiarowy Enter na końcu to najczęstsza pomyłka przy kopiowaniu
        z DevTools — przycinamy go po cichu i pobieranie idzie dalej
        (żadnego ``ValueError`` z ``http.client``, żadnego śladu stosu).
        """
        srv = self._serwer_html()
        katalog = tempfile.mkdtemp(prefix="flexit-crlf-")
        self.addCleanup(lambda: __import__("shutil").rmtree(katalog, ignore_errors=True))
        bufor = io.StringIO()
        with contextlib.redirect_stdout(bufor), contextlib.redirect_stderr(bufor):
            kod = cli.main(["download", "--base-url", srv.base,
                            "--out", os.path.join(katalog, "xml"),
                            "--cookie", "SESSIONID=abc\n",
                            "--delay", "0", "--retries", "0", "-q"])
        self.assertNotEqual(kod, cli.EXIT_ERROR)
        self.assertNotIn("Traceback", bufor.getvalue())
        self.assertTrue(srv.trafienia, "żądanie musiało wyjść mimo nadmiarowego Entera")
        for _sciezka, naglowki in srv.trafienia:
            self.assertEqual(naglowki.get("Cookie"), "SESSIONID=abc")

    def test_ok_ciasteczko_z_wstrzyknietym_naglowkiem_daje_czytelny_blad(self):
        """``\\n`` W ŚRODKU wartości to próba wstrzyknięcia — CLI odmawia po polsku."""
        srv = self._serwer_html()
        katalog = tempfile.mkdtemp(prefix="flexit-crlf2-")
        self.addCleanup(lambda: __import__("shutil").rmtree(katalog, ignore_errors=True))
        bufor = io.StringIO()
        with contextlib.redirect_stdout(bufor), contextlib.redirect_stderr(bufor):
            kod = cli.main(["download", "--base-url", srv.base,
                            "--out", os.path.join(katalog, "xml"),
                            "--cookie", "SESSIONID=abc\nX-Wstrzykniete: 1",
                            "--delay", "0", "--retries", "0", "-q"])
        self.assertEqual(kod, cli.EXIT_ERROR)
        self.assertIn("--cookie", bufor.getvalue())
        self.assertIn("końca linii", bufor.getvalue())
        self.assertNotIn("Traceback", bufor.getvalue())
        self.assertFalse(any("X-Wstrzykniete" in nagl
                             for _s, nagl in srv.trafienia))

# --------------------------------------------------------------------------- #
# 8. Wstrzyknięcie formuły do Excela
# --------------------------------------------------------------------------- #


class TestWstrzykniecieFormuly(_BazaZKatalogiem):
    """Wartości i nagłówki z XML nie mogą stać się formułami w arkuszu."""

    NIEBEZPIECZNE = ["=1+1", "+1+1", "-1+1", "@SUM(A1)",
                     "=cmd|'/C calc'!A0", '=HYPERLINK("http://zly.test")',
                     "\t=1+1", "\r=x"]

    def test_ok_wartosci_zapisane_jako_tekst_w_obu_backendach(self):
        if openpyxl is None:  # pragma: no cover
            self.skipTest("openpyxl niedostępny")
        wiersze = [[wartosc] for wartosc in self.NIEBEZPIECZNE]
        for nazwa_backendu in ("stdlib", "openpyxl"):
            with self.subTest(backend=nazwa_backendu), _backend(nazwa_backendu):
                cel = self.sciezka("formula_%s.xlsx" % nazwa_backendu)
                xlsxwrite.write_workbook(
                    cel, [xlsxwrite.Sheet(name="S", columns=["v"],
                                          rows=[list(w) for w in wiersze])])
                skoroszyt = openpyxl.load_workbook(cel)
                arkusz = skoroszyt.active
                for numer in range(2, 2 + len(wiersze)):
                    komorka = arkusz.cell(row=numer, column=1)
                    self.assertNotEqual(komorka.data_type, "f",
                                        "komórka %s stała się formułą" % komorka.coordinate)
                skoroszyt.close()

    def test_ok_naglowek_kolumny_z_xml_nie_jest_formula(self):
        """Nagłówek potrafi pochodzić z TREŚCI XML (heurystyka klucz–wartość)."""
        if openpyxl is None:  # pragma: no cover
            self.skipTest("openpyxl niedostępny")
        wejscie = os.path.join(self.katalog, "we_formula")
        os.makedirs(wejscie, exist_ok=True)
        with open(os.path.join(wejscie, "lot.xml"), "w", encoding="utf-8") as uchwyt:
            uchwyt.write(
                "<?xml version='1.0'?><lot><items>"
                "<item><sku>1</sku>"
                "<spec name='=cmd|&apos;/C calc&apos;!A0' value='x'/></item>"
                "<item><sku>2</sku>"
                "<spec name='=HYPERLINK(\"http://zly.test\")' value='y'/></item>"
                "</items></lot>"
            )
        for nazwa_backendu in ("stdlib", "openpyxl"):
            with self.subTest(backend=nazwa_backendu), _backend(nazwa_backendu):
                cel = self.sciezka("naglowek_%s.xlsx" % nazwa_backendu)
                bufor = io.StringIO()
                with contextlib.redirect_stdout(bufor), contextlib.redirect_stderr(bufor):
                    kod = cli.main(["build", "--in", wejscie, "--out", cel,
                                    "--overwrite", "-q"])
                self.assertEqual(kod, cli.EXIT_OK, bufor.getvalue())
                skoroszyt = openpyxl.load_workbook(cel)
                arkusz = skoroszyt[cli.SHEET_ALL]
                formuly = [k.coordinate for k in arkusz[1] if k.data_type == "f"]
                naglowki = [str(k.value) for k in arkusz[1] if k.value]
                skoroszyt.close()
                self.assertEqual(formuly, [])
                self.assertTrue(any(n.startswith("=") for n in naglowki),
                                "test musi faktycznie zawierać nagłówek z '='")


# --------------------------------------------------------------------------- #
# 9. Nagłówki kolumn przechodzą tę samą sanityzację co wartości
# --------------------------------------------------------------------------- #


class TestNaglowkiKolumnBezSanityzacji(_BazaZKatalogiem):
    """``xlsxwrite`` sanityzuje ``Sheet.columns`` dokładnie tak jak komórki.

    Nazwa kolumny pochodzi z nazwy tagu/atrybutu w XML-u portalu, więc może
    zawierać cokolwiek.  Obie ścieżki zapisu muszą ją przepuścić przez
    ``values.sanitize_cell``: usunąć znaki zakazane w XLSX i przyciąć do
    ``MAX_CELL_CHARS``.  Wynik obu backendów ma być identyczny.
    """

    def test_ok_znak_sterujacy_w_naglowku_nie_psuje_pliku_stdlib(self):
        """Znak ``\\x0b`` znika z nagłówka, a ``sheet1.xml`` jest poprawnym XML-em."""
        if _lxml_etree is None:  # pragma: no cover
            self.skipTest("lxml niedostępny")
        with _backend("stdlib"):
            cel = self.sciezka("ctrl_stdlib.xlsx")
            xlsxwrite.write_workbook(
                cel, [xlsxwrite.Sheet(name="S", columns=["kolumna\x0bz_pionowa_tab"],
                                      rows=[["1"]])])
        with zipfile.ZipFile(cel) as archiwum:
            surowy = archiwum.read("xl/worksheets/sheet1.xml")
        self.assertNotIn(b"\x0b", surowy)
        _lxml_etree.fromstring(surowy)          # musi się sparsować

    def test_ok_znak_sterujacy_w_naglowku_nie_wywala_openpyxl(self):
        """Ścieżka openpyxl nie rzuca ``IllegalCharacterError`` — nagłówek jest odkażony."""
        if openpyxl is None:  # pragma: no cover
            self.skipTest("openpyxl niedostępny")
        with _backend("openpyxl"):
            cel = self.sciezka("ctrl_openpyxl.xlsx")
            xlsxwrite.write_workbook(
                cel, [xlsxwrite.Sheet(name="S", columns=["kolumna\x0bz_tab"],
                                      rows=[["1"]])])
        skoroszyt = openpyxl.load_workbook(cel)
        try:
            self.assertEqual(skoroszyt["S"].cell(row=1, column=1).value, "kolumnaz_tab")
        finally:
            skoroszyt.close()

    def test_ok_dlugi_naglowek_z_xml_jest_przyciety_tak_samo_w_obu_backendach(self):
        """Nagłówek 40 000 znaków skraca się do limitu Excela w OBU ścieżkach."""
        if openpyxl is None:  # pragma: no cover
            self.skipTest("openpyxl niedostępny")
        wejscie = os.path.join(self.katalog, "we_dlugi")
        os.makedirs(wejscie, exist_ok=True)
        dlugi = "D" * 40000
        with open(os.path.join(wejscie, "lot.xml"), "w", encoding="utf-8") as uchwyt:
            uchwyt.write(
                "<?xml version='1.0'?><lot><items>"
                "<item><sku>1</sku><spec name='%s' value='x'/></item>"
                "<item><sku>2</sku><spec name='inna' value='y'/></item>"
                "</items></lot>" % dlugi
            )
        dlugosci = {}
        for nazwa_backendu in ("stdlib", "openpyxl"):
            with _backend(nazwa_backendu):
                cel = self.sciezka("dlugi_%s.xlsx" % nazwa_backendu)
                bufor = io.StringIO()
                with contextlib.redirect_stdout(bufor), contextlib.redirect_stderr(bufor):
                    kod = cli.main(["build", "--in", wejscie, "--out", cel,
                                    "--overwrite", "-q"])
                self.assertEqual(kod, cli.EXIT_OK, bufor.getvalue())
                skoroszyt = openpyxl.load_workbook(cel)
                try:
                    arkusz = skoroszyt[cli.SHEET_ALL]
                    dlugosci[nazwa_backendu] = max(
                        len(str(k.value)) for k in arkusz[1] if k.value)
                finally:
                    skoroszyt.close()
        self.assertEqual(dlugosci["stdlib"], values.MAX_CELL_CHARS)
        self.assertEqual(dlugosci["openpyxl"], values.MAX_CELL_CHARS)

    def test_ok_wartosci_komorek_sa_przycinane(self):
        """Kontrola: dane (w odróżnieniu od nagłówków) są sanityzowane poprawnie."""
        self.assertEqual(len(values.sanitize_cell("X" * 40000)), values.MAX_CELL_CHARS)
        self.assertEqual(values.sanitize_cell("a\x0bb"), "ab")


# --------------------------------------------------------------------------- #
# 10. Kody błędów Excela („#N/A'') są zwykłym TEKSTEM w obu backendach
# --------------------------------------------------------------------------- #


class TestKomorkiBledu(_BazaZKatalogiem):
    """Ten sam XML musi dać ten sam TYP komórki niezależnie od backendu.

    Ciągi z listy kodów błędów Excela (``#N/A``, ``#REF!``, ``#DIV/0!``,
    ``#VALUE!``, ``#NAME?``, ``#NULL!``, ``#NUM!``) openpyxl domyślnie zamienia
    na komórki typu ``e`` (błąd): przestają być tekstem, propagują błąd do
    formuł i inaczej zachowują się w filtrach.  Wartość wzięta z XML-a jest
    tekstem, więc obie ścieżki zapisują ją jako ``s``.
    """

    KODY = ["#N/A", "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NULL!", "#NUM!"]

    def test_ok_kody_bledow_maja_ten_sam_typ_w_obu_backendach(self):
        if openpyxl is None:  # pragma: no cover
            self.skipTest("openpyxl niedostępny")
        typy = {}
        wartosci = {}
        for nazwa_backendu in ("stdlib", "openpyxl"):
            with _backend(nazwa_backendu):
                cel = self.sciezka("bledy_%s.xlsx" % nazwa_backendu)
                xlsxwrite.write_workbook(
                    cel, [xlsxwrite.Sheet(name="S", columns=["stan"],
                                          rows=[[kod] for kod in self.KODY])])
                skoroszyt = openpyxl.load_workbook(cel)
                try:
                    arkusz = skoroszyt.active
                    typy[nazwa_backendu] = [
                        arkusz.cell(row=i, column=1).data_type
                        for i in range(2, 2 + len(self.KODY))]
                    wartosci[nazwa_backendu] = [
                        arkusz.cell(row=i, column=1).value
                        for i in range(2, 2 + len(self.KODY))]
                finally:
                    skoroszyt.close()
        self.assertEqual(typy["stdlib"], ["s"] * len(self.KODY))
        self.assertEqual(typy["openpyxl"], ["s"] * len(self.KODY))
        self.assertEqual(wartosci["stdlib"], self.KODY)
        self.assertEqual(wartosci["openpyxl"], self.KODY)

    def test_ok_pelny_potok_cli_nie_daje_komorki_bledu(self):
        """Realistyczny scenariusz: pole ``#N/A`` w eksporcie „Batch Details''."""
        if openpyxl is None:  # pragma: no cover
            self.skipTest("openpyxl niedostępny")
        wejscie = os.path.join(self.katalog, "we_bledy")
        os.makedirs(wejscie, exist_ok=True)
        with open(os.path.join(wejscie, "batch.xml"), "w", encoding="utf-8") as uchwyt:
            uchwyt.write(
                "<?xml version='1.0'?><lot id='L-1'><items>"
                "<item><sku>A1</sku><bateria>#N/A</bateria></item>"
                "<item><sku>A2</sku><bateria>85%</bateria></item>"
                "</items></lot>"
            )
        wyniki = {}
        for nazwa_backendu in ("stdlib", "openpyxl"):
            with _backend(nazwa_backendu):
                cel = self.sciezka("potok_%s.xlsx" % nazwa_backendu)
                bufor = io.StringIO()
                with contextlib.redirect_stdout(bufor), contextlib.redirect_stderr(bufor):
                    kod = cli.main(["build", "--in", wejscie, "--out", cel,
                                    "--overwrite", "-q"])
                self.assertEqual(kod, cli.EXIT_OK, bufor.getvalue())
                skoroszyt = openpyxl.load_workbook(cel)
                try:
                    arkusz = skoroszyt[cli.SHEET_ALL]
                    naglowki = [k.value for k in arkusz[1]]
                    kolumna = naglowki.index("bateria") + 1
                    komorka = arkusz.cell(row=2, column=kolumna)
                    wyniki[nazwa_backendu] = (komorka.value, komorka.data_type)
                finally:
                    skoroszyt.close()
        self.assertEqual(wyniki["stdlib"], ("#N/A", "s"))
        self.assertEqual(wyniki["openpyxl"], ("#N/A", "s"),
                         "identyczny wynik niezależnie od zainstalowanych bibliotek")

# --------------------------------------------------------------------------- #
# 11. Sekwencje sterujące terminala z portalu są odsiewane
# --------------------------------------------------------------------------- #


class TestEscapeTerminala(unittest.TestCase):
    """Treść z portalu nie może sterować terminalem użytkownika.

    Złośliwa (albo zhakowana) strona mogłaby wyczyścić ekran, zmienić tytuł
    okna terminala i ukryć ostrzeżenia o pominiętych plikach.  Dlatego znaki
    ``\\x1b`` (i pozostałe znaki sterujące) są usuwane z identyfikatorów,
    tytułów i z raportu ``--diagnose``, zanim cokolwiek trafi na konsolę.
    """

    HTML = (
        "<html><head><title>Aukcje\x1b]0;PRZEJETY-TYTUL\x07</title></head><body>"
        "<a href='/auction/flexit-\x1b[2J\x1b[31mKRWAWY-1103/'>Lot \x1b[5mMIGA\x1b[0m</a>"
        "</body></html>"
    ).encode("utf-8")

    def test_ok_esc_z_portalu_nie_dociera_do_komunikatow_cli(self):
        srv = _Serwer({"/": (200, {"Content-Type": "text/html; charset=utf-8"},
                             self.HTML)})
        self.addCleanup(srv.zamknij)
        opener = scrape.make_opener(delay=0, retries=0)
        aukcje = scrape.discover_auctions(opener, srv.base)
        self.assertEqual(len(aukcje), 1)
        self.assertNotIn("\x1b", aukcje[0].id,
                         "identyfikator bez znaków sterujących terminala")
        self.assertNotIn("\x1b", aukcje[0].title,
                         "tytuł bez znaków sterujących terminala")

        bufor = io.StringIO()
        reporter = cli.Reporter(out=bufor, err=bufor)
        reporter.info("[1/1] %s — %s" % (aukcje[0].id, aukcje[0].title))
        self.assertNotIn("\x1b", bufor.getvalue(),
                         "CLI nie wypisuje sekwencji ESC z obcej strony")

    def test_ok_esc_nie_przechodzi_do_raportu_diagnostycznego(self):
        raport = scrape.describe_page(self.HTML, "http://portal.test/",
                                      content_type="text/html")
        self.assertNotIn("\x1b", raport,
                         "--diagnose oczyszcza treść pobraną z portalu")
        self.assertNotIn("\x07", raport)

    def test_ok_esc_nie_trafia_do_nazwy_pliku(self):
        """Kontrola: w nazwach plików znaki sterujące są usuwane poprawnie."""
        self.assertEqual(scrape.safe_filename("\x1b[2Jbatch\x07.xml"), "2Jbatch.xml")


# --------------------------------------------------------------------------- #
# 12. Nadpisywanie plików i katalog wyjściowy
# --------------------------------------------------------------------------- #


class TestZapisWyniku(_BazaZKatalogiem):
    """``build --out`` nie może po cichu nadpisać cudzych danych."""

    def test_ok_brak_nadpisania_bez_overwrite(self):
        wejscie = os.path.join(self.katalog, "we")
        os.makedirs(wejscie, exist_ok=True)
        with open(os.path.join(wejscie, "a.xml"), "w") as uchwyt:
            uchwyt.write("<r><item><a>1</a></item><item><a>2</a></item></r>")
        cel = self.zapisz("wazny.xlsx", b"NIE-RUSZAJ-MNIE")
        bufor = io.StringIO()
        with contextlib.redirect_stdout(bufor), contextlib.redirect_stderr(bufor):
            kod = cli.main(["build", "--in", wejscie, "--out", cel, "-q"])
        self.assertEqual(kod, cli.EXIT_ERROR)
        with open(cel, "rb") as uchwyt:
            self.assertEqual(uchwyt.read(), b"NIE-RUSZAJ-MNIE")

    def test_ok_nieudany_zapis_nie_zostawia_smieci(self):
        """Wyjątek w trakcie zapisu musi usunąć plik tymczasowy i nie ruszyć celu."""
        katalog = self.sciezka("wyjscie")
        os.makedirs(katalog, exist_ok=True)

        def zly_generator():
            yield ["ok"]
            raise RuntimeError("awaria w połowie zapisu")

        cel = os.path.join(katalog, "wynik.xlsx")
        for nazwa_backendu in ("stdlib", "openpyxl"):
            if nazwa_backendu == "openpyxl" and openpyxl is None:  # pragma: no cover
                continue
            with self.subTest(backend=nazwa_backendu), _backend(nazwa_backendu):
                with self.assertRaises(RuntimeError):
                    xlsxwrite.write_workbook(
                        cel, [xlsxwrite.Sheet(name="S", columns=["a"],
                                              rows=zly_generator())])
                self.assertFalse(os.path.exists(cel))
                pozostale = [n for n in os.listdir(katalog)
                             if n.startswith(".flexit2xlsx-")]
                self.assertEqual(pozostale, [], "plik tymczasowy nie został usunięty")

    def test_ok_download_nie_pisze_poza_wskazany_katalog(self):
        """Pełny przebieg ``download`` z portalem podsuwającym złośliwe nazwy."""
        strony = {
            "/": (200, {"Content-Type": "text/html"},
                  b"<html><a href='/auction/flexit-auctions-01-01-2026-1103/'>A</a></html>"),
            "/auction/flexit-auctions-01-01-2026-1103/":
                (200, {"Content-Type": "text/html"},
                 b"<html>"
                 b"<a href='/api/batch.xml?file=../../../../tmp/przejete.xml' "
                 b"download='../../../../tmp/przejete.xml'>Batch</a>"
                 b"</html>"),
            "/api/batch.xml": (200, {"Content-Type": "application/xml"},
                               b"<?xml version='1.0'?><batch><item><a>1</a></item>"
                               b"<item><a>2</a></item></batch>"),
        }
        srv = _Serwer(strony)
        self.addCleanup(srv.zamknij)
        cel = self.sciezka("pobrane_cli")
        bufor = io.StringIO()
        with contextlib.redirect_stdout(bufor), contextlib.redirect_stderr(bufor):
            kod = cli.main(["download", "--base-url", srv.base, "--out", cel,
                            "--delay", "0", "--retries", "0", "-q"])
        self.assertIn(kod, (cli.EXIT_OK, cli.EXIT_PARTIAL), bufor.getvalue())
        for nazwa in os.listdir(cel):
            self.assertNotIn("..", nazwa)
            self.assertEqual(nazwa, os.path.basename(nazwa))
        self.assertFalse(os.path.exists("/tmp/przejete.xml"))


if __name__ == "__main__":  # pragma: no cover
    unittest.main(verbosity=2)
