# -*- coding: utf-8 -*-
"""Testy ADWERSARYJNE: brudny XML i cały potok XML -> XLSX.

Wszystkie testy (``test_ok_*``) potwierdzają zachowanie WYMAGANE i chronią
przed regresją: XXE, bomby encyjne, uszkodzone pliki, kodowania, limity Excela,
fuzzing strukturalny i bajtowy.

Historycznie moduł zawierał też testy ``test_blad_*`` utrwalające zaobserwowane
WADY (asercja opisywała stan faktyczny, komunikat — oczekiwany).  Po naprawieniu
tych wad każdy taki test został przepisany na asercję stanu poprawnego.

Uruchamianie (w katalogu repozytorium):

    python3 -m unittest discover -s tests/adversarial -t . -v
    python3 tests/adversarial/test_brudny_xml.py
"""

from __future__ import annotations

import contextlib
import io
import itertools
import os
import random
import shutil
import sys
import tempfile
import unittest
import zipfile

_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

from flexit2xlsx import cli, values, xlsxwrite  # noqa: E402
from flexit2xlsx import xmlflatten as xf  # noqa: E402

try:  # openpyxl służy WYŁĄCZNIE do weryfikacji wyników (wolno go użyć w testach)
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


# --------------------------------------------------------------------------- #
# Narzędzia pomocnicze
# --------------------------------------------------------------------------- #

def parse(text, source="test.xml", **kw):
    """Parsuje XML podany jako ``str`` (kodowanie UTF-8) albo ``bytes``."""
    data = text.encode("utf-8") if isinstance(text, str) else text
    return xf.parse_bytes(data, source, **kw)


class BazaCLI(unittest.TestCase):
    """Wspólna baza: katalog tymczasowy + uruchamianie CLI w tym samym procesie."""

    def setUp(self):
        self.katalog = tempfile.mkdtemp(prefix="flexit-adw-")
        self.wejscie = os.path.join(self.katalog, "xml")
        os.makedirs(self.wejscie)
        self.wynik = os.path.join(self.katalog, "wynik.xlsx")
        self.addCleanup(shutil.rmtree, self.katalog, True)

    def zapisz(self, nazwa, tresc):
        """Zapisuje plik wejściowy; ``tresc`` może być ``str`` albo ``bytes``."""
        sciezka = os.path.join(self.wejscie, nazwa)
        tryb = "wb" if isinstance(tresc, (bytes, bytearray)) else "w"
        with open(sciezka, tryb, **({} if tryb == "wb" else {"encoding": "utf-8"})) as fh:
            fh.write(tresc)
        return sciezka

    def uruchom(self, *dodatkowe, backend=None, out=None):
        """Uruchamia ``build`` i zwraca ``(kod, stdout, stderr)``."""
        cel = out or self.wynik
        argv = ["build", "--in", self.wejscie, "--out", cel, "--overwrite"]
        argv.extend(dodatkowe)
        poprzedni = os.environ.get(xlsxwrite.ENV_BACKEND)
        if backend is not None:
            os.environ[xlsxwrite.ENV_BACKEND] = backend
        wyjscie, blad = io.StringIO(), io.StringIO()
        try:
            with contextlib.redirect_stdout(wyjscie), contextlib.redirect_stderr(blad):
                kod = cli.main(argv)
        finally:
            if backend is not None:
                if poprzedni is None:
                    os.environ.pop(xlsxwrite.ENV_BACKEND, None)
                else:
                    os.environ[xlsxwrite.ENV_BACKEND] = poprzedni
        return kod, wyjscie.getvalue(), blad.getvalue()

    def arkusze(self, sciezka=None):
        """Wczytuje wynik przez openpyxl: ``{nazwa arkusza: [wiersze]}``."""
        if openpyxl is None:  # pragma: no cover
            self.skipTest("openpyxl niedostępny")
        skoroszyt = openpyxl.load_workbook(sciezka or self.wynik)
        return {
            nazwa: [list(w) for w in skoroszyt[nazwa].iter_rows(values_only=True)]
            for nazwa in skoroszyt.sheetnames
        }


# --------------------------------------------------------------------------- #
# 1. Wykrywanie elementu powtarzalnego
# --------------------------------------------------------------------------- #

class TestWykrywaniaRekordu(BazaCLI):
    """Heurystyka ``detect_record_path`` na realistycznych plikach „Batch Details”."""

    PAKIET_3 = """<?xml version="1.0" encoding="UTF-8"?>
<batch lotHash="aaaaa" auction="flexit-auctions-18-06-2026-1103">
  <lot><title>3x ThinkPad</title><items>
    <item nr="1"><model>T480</model><serial>TP1</serial><grade>B</grade></item>
    <item nr="2"><model>T490</model><serial>TP2</serial><grade>A</grade></item>
    <item nr="3"><model>T14</model><serial>TP3</serial><grade>C</grade></item>
  </items></lot>
</batch>"""

    PAKIET_1_ZE_ZDJECIAMI = """<?xml version="1.0" encoding="UTF-8"?>
<batch lotHash="bbbbb" auction="flexit-auctions-18-06-2026-1103">
  <lot><title>1x Dell OptiPlex</title>
    <photos><photo>https://media/1.jpg</photo><photo>https://media/2.jpg</photo>
            <photo>https://media/3.jpg</photo><photo>https://media/4.jpg</photo></photos>
    <items>
      <item nr="1"><model>OptiPlex 7050</model><serial>DL9</serial><grade>A</grade></item>
    </items></lot>
</batch>"""

    def test_ok_pozycja_pakietu_wygrywa_ze_zdjeciami(self):
        """Lot z 1 sztuką i 4 zdjęciami daje 1 wiersz, nie 4 wiersze-widma.

        ``<photo>`` powtarza się częściej niż ``<item>``, ale nie jest pozycją
        pakietu — heurystyka musi preferować nazwę z listy rekordowej i element
        o większej liczbie pól-liści.
        """
        doc = parse(self.PAKIET_1_ZE_ZDJECIAMI, "batch_bbbbb.xml")
        self.assertEqual(doc.record_path, "batch/lot/items/item")
        self.assertEqual(len(doc.records), 1)
        self.assertEqual(doc.records[0]["model"], "OptiPlex 7050")
        # Zdjęcia to metadane lotu — trafiają do kontekstu, sklejone separatorem.
        self.assertIn("batch/lot/photos/photo", doc.context)
        self.assertEqual(doc.context["batch/lot/photos/photo"].count("|"), 3)

    def test_ok_zgodne_kolumny_miedzy_plikami_tej_samej_aukcji(self):
        """Ta sama informacja trafia do JEDNEJ rodziny kolumn w całym arkuszu."""
        self.zapisz("batch_aaaaa.xml", self.PAKIET_3)
        self.zapisz("batch_bbbbb.xml", self.PAKIET_1_ZE_ZDJECIAMI)
        kod, out, _ = self.uruchom()
        self.assertEqual(kod, 0, out)
        arkusz = self.arkusze()["Wszystkie aukcje"]
        naglowki = arkusz[0]
        self.assertIn("model", naglowki)
        self.assertNotIn("batch/lot/items/item/model", naglowki,
                         "jedna kolumna 'model' dla obu plików")
        self.assertEqual(len(arkusz) - 1, 4, "3 sztuki + 1 sztuka")
        wiersze_bbbbb = [w for w in arkusz[1:] if w[1] == "batch_bbbbb.xml"]
        self.assertEqual(len(wiersze_bbbbb), 1)
        kol_model = naglowki.index("model")
        self.assertEqual(wiersze_bbbbb[0][kol_model], "OptiPlex 7050")

    def test_ok_pakiet_wielosztukowy_wykrywany_poprawnie(self):
        doc = parse(self.PAKIET_3, "batch_aaaaa.xml")
        self.assertEqual(doc.record_path, "batch/lot/items/item")
        self.assertEqual(len(doc.records), 3)

    def test_ok_ujednolicanie_dziala_gdy_brak_powtorzen(self):
        """Plik z 1 sztuką BEZ zdjęć jest ratowany przez cli.unify_record_paths."""
        jeden = self.PAKIET_1_ZE_ZDJECIAMI.replace(
            """<photos><photo>https://media/1.jpg</photo><photo>https://media/2.jpg</photo>
            <photo>https://media/3.jpg</photo><photo>https://media/4.jpg</photo></photos>""", "")
        self.zapisz("batch_aaaaa.xml", self.PAKIET_3)
        self.zapisz("batch_ccccc.xml", jeden)
        kod, out, _ = self.uruchom()
        self.assertEqual(kod, 0, out)
        arkusz = self.arkusze()["Wszystkie aukcje"]
        self.assertEqual(len(arkusz) - 1, 4)
        self.assertNotIn("batch/lot/items/item/model", arkusz[0])

    def test_ok_pojedynczy_element_i_pusty_kontener(self):
        doc = parse("<pakiet/>", "pusty.xml")
        self.assertIsNone(doc.record_path)
        self.assertEqual(len(doc.records), 1)


# --------------------------------------------------------------------------- #
# 2. Pary klucz–wartość
# --------------------------------------------------------------------------- #

class TestParKluczWartosc(BazaCLI):
    """``<spec name="X">v</spec>`` – zamiana na kolumnę i jej skutki uboczne."""

    def test_ok_nazwa_cechy_nie_zlewa_sie_z_prawdziwym_tagiem(self):
        """Cecha o nazwie kolidującej z tagiem dostaje WŁASNĄ kolumnę."""
        doc = parse(
            "<aukcja id='A-1'>"
            "<item><model>ThinkPad X230</model><spec name='model'>i5-3320M</spec>"
            "<sn>111</sn></item>"
            "<item><model>ThinkPad T440</model><spec name='model'>i5-4300U</spec>"
            "<sn>222</sn></item></aukcja>")
        self.assertEqual(doc.records[0]["model"], "ThinkPad X230")
        self.assertEqual(doc.records[0]["spec/model"], "i5-3320M")
        self.assertEqual(doc.records[1]["model"], "ThinkPad T440")
        self.assertEqual(doc.records[1]["spec/model"], "i5-4300U")
        self.assertNotIn("spec@name", doc.columns)

    def test_ok_rekord_z_jedna_cecha_ma_te_same_kolumny_co_z_dwiema(self):
        """Kolumna 'Kolor' jest wypełniona dla KAŻDEJ pozycji, która kolor ma."""
        doc = parse(
            "<aukcja id='A-1'>"
            "<pozycja><sn>1</sn><opcja nazwa='Kolor'>czarny</opcja>"
            "<opcja nazwa='Rozmiar'>M</opcja></pozycja>"
            "<pozycja><sn>2</sn><opcja nazwa='Kolor'>bialy</opcja></pozycja>"
            "</aukcja>")
        self.assertEqual(doc.records[0].get("Kolor"), "czarny")
        self.assertEqual(doc.records[1].get("Kolor"), "bialy")
        self.assertEqual(doc.records[0].get("Rozmiar"), "M")
        self.assertIsNone(doc.records[1].get("Rozmiar"))
        # Żadnych zbędnych kolumn technicznych po parze klucz-wartość:
        self.assertNotIn("opcja", doc.columns)
        self.assertNotIn("opcja@nazwa", doc.columns)

    def test_ok_tekst_elementu_zostaje_gdy_wartosc_jest_w_atrybucie(self):
        """Treść ``<spec ...>tu tekst</spec>`` ląduje w osobnej kolumnie."""
        doc = parse(
            "<aukcja><item>"
            "<spec name='RAM' value='16 GB'>po rozbudowie u sprzedawcy</spec>"
            "<sn>1</sn></item>"
            "<item><spec name='RAM' value='8 GB'>fabryczne</spec><sn>2</sn></item>"
            "</aukcja>")
        self.assertEqual(doc.records[0].get("RAM"), "16 GB")
        self.assertEqual(doc.records[0].get("RAM (tekst)"),
                         "po rozbudowie u sprzedawcy")
        self.assertEqual(doc.records[1].get("RAM"), "8 GB")
        self.assertEqual(doc.records[1].get("RAM (tekst)"), "fabryczne")

    def test_ok_atrybuty_elementu_klucza_nie_gina(self):
        """Atrybut na elemencie pełniącym rolę klucza trafia do własnej kolumny."""
        doc = parse(
            "<aukcja><item>"
            "<cecha><nazwa jezyk='pl' id='C1'>Kolor</nazwa><wartosc>czarny</wartosc></cecha>"
            "<cecha><nazwa jezyk='pl' id='C2'>Rozmiar</nazwa><wartosc>M</wartosc></cecha>"
            "</item><item><sn>2</sn></item></aukcja>")
        self.assertEqual(doc.records[0].get("Kolor"), "czarny")
        self.assertEqual(doc.records[0].get("Rozmiar"), "M")
        wszystko = " ".join(str(v) for v in doc.records[0].values())
        self.assertIn("C1", wszystko)
        self.assertIn("C2", wszystko)

    def test_ok_zwykla_para_nazwa_tresc(self):
        doc = parse(
            "<aukcja><item><spec name='RAM'>16 GB</spec><spec name='Dysk'>512 GB</spec></item>"
            "<item><spec name='RAM'>8 GB</spec><spec name='Dysk'>256 GB</spec></item></aukcja>")
        self.assertEqual(doc.records[0], {"RAM": "16 GB", "Dysk": "512 GB"})
        self.assertEqual(doc.records[1], {"RAM": "8 GB", "Dysk": "256 GB"})


# --------------------------------------------------------------------------- #
# 3. Przestrzenie nazw
# --------------------------------------------------------------------------- #

class TestPrzestrzeniNazw(BazaCLI):

    XML = ("<b xmlns:dc='http://purl.org/dc/elements/1.1/' xmlns:f='http://flexit/'>"
           "<i><dc:title>Tytul katalogowy</dc:title>"
           "<f:title>Tytul sprzedazy</f:title><sn>1</sn></i>"
           "<i><dc:title>A</dc:title><f:title>B</f:title><sn>2</sn></i></b>")

    def test_ok_obcinanie_prefiksow_nie_zlewa_dwoch_pol(self):
        """``dc:title`` i ``f:title`` dostają OSOBNE kolumny mimo ``strip_ns``.

        Przy kolizji nazw po obcięciu prefiksu druga (i kolejna) przestrzeń
        zachowuje prefiks, a użytkownik dostaje ostrzeżenie na stderr.
        """
        blad = io.StringIO()
        with contextlib.redirect_stderr(blad):
            doc = parse(self.XML)
        self.assertEqual(doc.columns, ["title", "f:title", "sn"])
        self.assertEqual(doc.records[0]["title"], "Tytul katalogowy")
        self.assertEqual(doc.records[0]["f:title"], "Tytul sprzedazy")
        self.assertIn("kolidowały", blad.getvalue())
        self.assertIn("--keep-ns", blad.getvalue())

    def test_ok_keep_ns_rozdziela_pola(self):
        doc = parse(self.XML, strip_ns=False)
        self.assertEqual(doc.records[0]["dc:title"], "Tytul katalogowy")
        self.assertEqual(doc.records[0]["f:title"], "Tytul sprzedazy")

    def test_ok_niezadeklarowany_prefiks_nie_wywraca_parsera(self):
        doc = parse("<a><i><ns:x>1</ns:x></i><i><ns:x>2</ns:x></i></a>")
        self.assertEqual([r.get("x") for r in doc.records], ["1", "2"])


# --------------------------------------------------------------------------- #
# 4. Kodowania, BOM
# --------------------------------------------------------------------------- #

class TestKodowan(BazaCLI):

    SZABLON = "<?xml version='1.0' encoding='%s'?><a><i><n>zażółć gęślą</n></i><i><n>b</n></i></a>"

    def test_ok_utf8_z_bledna_deklaracja_iso_jest_ratowane(self):
        """Bajty UTF-8 z (błędną) deklaracją ISO-8859-2 czytamy jako UTF-8.

        Deklaracja w prologu bywa nieprawdziwa (eksport doklejony do szablonu),
        a poprawny UTF-8 rozpoznaje się jednoznacznie — więc ufamy zawartości,
        nie deklaracji, i mówimy o tym na stderr.
        """
        dane = (self.SZABLON % "ISO-8859-2").encode("utf-8")
        blad = io.StringIO()
        with contextlib.redirect_stderr(blad):
            doc = parse(dane)
        tekst = doc.records[0]["n"]
        self.assertEqual(tekst, "zażółć gęślą")
        self.assertNotIn("\x82", tekst, "żadnych znaków sterujących C1 z mojibake")
        self.assertIn("ISO-8859-2", blad.getvalue())
        self.assertIn("UTF-8", blad.getvalue())

    def test_ok_cp1250_z_deklaracja_utf8(self):
        dane = (self.SZABLON % "UTF-8").encode("cp1250")
        self.assertEqual(parse(dane).records[0]["n"], "zażółć gęślą")

    def test_ok_iso88592_zadeklarowane(self):
        self.assertEqual(
            parse((self.SZABLON % "ISO-8859-2").encode("iso-8859-2")).records[0]["n"],
            "zażółć gęślą")

    def test_ok_iso88592_bez_deklaracji_jest_rozpoznane(self):
        """Brak deklaracji: łańcuch awaryjny wybiera kodowanie po treści.

        Dla polskiego tekstu ``cp1250`` i ``ISO-8859-2`` różnią się właśnie na
        ``ś``/``ą``, więc ślepa kolejność zamieniałaby je na ``¶``/``±``.
        Wybór jest zgłaszany na stderr, żeby dało się go zweryfikować.
        """
        blad = io.StringIO()
        with contextlib.redirect_stderr(blad):
            doc = parse("<a><i><n>zażółć gęślą</n></i><i><n>b</n></i></a>"
                        .encode("iso-8859-2"))
        self.assertEqual(doc.records[0]["n"], "zażółć gęślą")
        self.assertIn("brak deklaracji kodowania", blad.getvalue())

    def test_ok_bom_utf8_przy_deklaracji_iso(self):
        dane = b"\xef\xbb\xbf" + (self.SZABLON % "ISO-8859-2").encode("iso-8859-2")
        self.assertEqual(parse(dane).records[0]["n"], "zażółć gęślą")

    def test_ok_utf16_z_bom(self):
        self.assertEqual(parse((self.SZABLON % "UTF-16").encode("utf-16")).records[0]["n"],
                         "zażółć gęślą")

    def test_ok_nieznana_nazwa_kodowania_daje_ostrzezenie(self):
        """Kodowanie nieznane Pythonowi nie może dać CICHYCH krzaków.

        ``x-mac-central-europe`` nie ma w Pythonie kodeka, więc dokładne
        odtworzenie treści jest niemożliwe — ale użytkownik musi się dowiedzieć,
        że nazwa z deklaracji nie została rozpoznana i że kodowanie zgadujemy.
        """
        dane = ("<?xml version='1.0' encoding='x-mac-central-europe'?>"
                "<a><i><n>zażółć</n></i><i><n>b</n></i></a>").encode("mac-latin2")
        blad = io.StringIO()
        with contextlib.redirect_stderr(blad):
            doc = parse(dane)
        self.assertEqual(len(doc.records), 2)
        komunikat = blad.getvalue()
        self.assertIn("x-mac-central-europe", komunikat)
        self.assertIn("nieznana nazwa kodowania", komunikat)

    def test_ok_polskie_znaki_docieraja_do_xlsx(self):
        self.zapisz("pl.xml", (self.SZABLON % "ISO-8859-2").encode("iso-8859-2"))
        kod, out, _ = self.uruchom()
        self.assertEqual(kod, 0, out)
        arkusz = self.arkusze()["Wszystkie aukcje"]
        self.assertIn("zażółć gęślą", [w[-1] for w in arkusz[1:]])


# --------------------------------------------------------------------------- #
# 5. Encje, XXE, bomby
# --------------------------------------------------------------------------- #

class TestEncji(BazaCLI):

    def test_ok_xxe_zablokowane(self):
        with self.assertRaises(xf.XmlParseError) as ctx:
            parse("<!DOCTYPE a [<!ENTITY xxe SYSTEM 'file:///etc/passwd'>]>"
                  "<a><i><n>&xxe;</n></i><i><n>b</n></i></a>")
        self.assertIn("XXE", str(ctx.exception))

    def test_ok_zewnetrzny_dtd_nie_jest_pobierany(self):
        doc = parse("<!DOCTYPE a SYSTEM 'http://192.0.2.1/evil.dtd'>"
                    "<a><i><n>1</n></i><i><n>2</n></i></a>")
        self.assertEqual([r["n"] for r in doc.records], ["1", "2"])

    def test_ok_bomba_encyjna_zablokowana(self):
        with self.assertRaises(xf.XmlParseError):
            parse("<!DOCTYPE a [<!ENTITY a 'aaaaaaaaaa'>"
                  "<!ENTITY b '&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;'>"
                  "<!ENTITY c '&b;&b;&b;&b;&b;&b;&b;&b;&b;&b;'>]>"
                  "<a><i><n>&c;</n></i><i><n>x</n></i></a>")

    def test_ok_zwykla_encja_z_amp_jest_rozwijana(self):
        """Legalna encja ``"Kowalski &amp; Syn"`` rozwija się, plik zostaje."""
        doc = parse("<!DOCTYPE a [<!ENTITY firma \"Kowalski &amp; Syn\">]>"
                    "<a><i><n>&firma;</n></i><i><n>b</n></i></a>")
        self.assertEqual(doc.records[0]["n"], "Kowalski & Syn")

    def test_ok_dluga_encja_nie_odrzuca_pliku(self):
        """Encja o treści 1100 znaków to jeszcze nie bomba — ma się rozwinąć."""
        doc = parse("<!DOCTYPE a [<!ENTITY op \"%s\">]>"
                    "<a><i><n>&op;</n></i><i><n>b</n></i></a>" % ("x" * 1100))
        self.assertEqual(doc.records[0]["n"], "x" * 1100)

    def test_ok_kilkadziesiat_encji_nie_odrzuca_pliku(self):
        """70 deklaracji encji mieści się w limicie — plik wczytuje się normalnie."""
        deklaracje = "".join('<!ENTITY e%d "v%d">' % (i, i) for i in range(70))
        doc = parse("<!DOCTYPE a [%s]><a><i><n>&e0;</n></i><i><n>b</n></i></a>"
                    % deklaracje)
        self.assertEqual(doc.records[0]["n"], "v0")

    def test_ok_nieznana_encja_zostaje_doslownie_z_ostrzezeniem(self):
        """``&nieznana;`` NIE znika po cichu — zostaje w treści, a stderr ostrzega."""
        blad = io.StringIO()
        with contextlib.redirect_stderr(blad):
            doc = parse("<a><i><n>AT&amp;T &nieznana; koniec</n></i>"
                        "<i><n>b</n></i></a>")
        self.assertEqual(doc.records[0]["n"], "AT&T &nieznana; koniec")
        self.assertIn("encje", blad.getvalue())

    def test_ok_encje_html_sa_ratowane(self):
        doc = parse("<a><i><n>5&nbsp;kg &oacute;</n></i><i><n>b</n></i></a>")
        self.assertEqual(doc.records[0]["n"], "5 kg ó")


# --------------------------------------------------------------------------- #
# 6. Uszkodzone i dziwne pliki w potoku CLI
# --------------------------------------------------------------------------- #

class TestUszkodzonychPlikow(BazaCLI):

    def test_ok_smieci_nie_przerywaja_budowania(self):
        self.zapisz("00_dobry.xml",
                    "<aukcja id='OK-1'><item><nazwa>Laptop</nazwa><cena>1 234,56</cena></item>"
                    "<item><nazwa>Monitor</nazwa><cena>99</cena></item></aukcja>")
        self.zapisz("01_pusty.xml", b"")
        self.zapisz("02_biale.xml", b"   \n\t ")
        self.zapisz("03_uciety.xml", "<aukcja><item><nazwa>Lap")
        self.zapisz("04_binarny.xml", bytes(range(256)) * 40)
        self.zapisz("05_dwa_korzenie.xml", "<a/><b/>")
        self.zapisz("06_zip.xml", b"PK\x03\x04" + os.urandom(200))
        kod, out, err = self.uruchom()
        self.assertEqual(kod, 3, out + err)
        arkusze = self.arkusze()
        wszystkie = arkusze["Wszystkie aukcje"]
        self.assertEqual(len(wszystkie) - 1, 2)
        self.assertEqual(wszystkie[1][4], "Laptop")
        self.assertEqual(wszystkie[1][5], 1234.56)
        bledy = [w for w in arkusze["Podsumowanie"] if str(w[5]).startswith("BŁĄD")]
        self.assertEqual(len(bledy), 6)

    def test_ok_uszkodzony_xml_to_XmlParseError(self):
        for zly in ("<a><b></a>", "<a", "<?xml version='1.0'?>", "<a>&#x0;</a>",
                    "<a b='1></a>", "<a><b/><a>"):
            with self.subTest(zly=zly):
                with self.assertRaises(xf.XmlParseError):
                    parse(zly)
                self.assertTrue(issubclass(xf.XmlParseError, ValueError))

    def test_ok_zagniezdzenie_ponad_limit_obcina_tylko_nadmiar(self):
        """Głębokość ponad ``_MAX_DEPTH`` obcina nadmiar, nie cały plik.

        Odrzucenie CAŁEGO dokumentu oznaczałoby utratę wszystkich pozycji przez
        jedno patologicznie zagnieżdżone pole.
        """
        def dokument(glebokosc):
            srodek = ("".join("<n%d>" % i for i in range(glebokosc)) + "X"
                      + "".join("</n%d>" % i for i in reversed(range(glebokosc))))
            return "<root><item>%s</item><item>%s</item></root>" % (srodek, srodek)

        self.assertEqual(len(parse(dokument(150)).records), 2)
        blad = io.StringIO()
        with contextlib.redirect_stderr(blad):
            doc = parse(dokument(500))
        self.assertEqual(len(doc.records), 2, "pozycje muszą przetrwać obcięcie")
        self.assertIn("obcięto", blad.getvalue())

    def test_ok_cdata_mieszana_tresc_i_puste_elementy(self):
        doc = parse("<a><i><n><![CDATA[<b>1 & 2</b>]]></n><pusty/><z x=''/></i>"
                    "<i><n>b</n></i></a>")
        self.assertEqual(doc.records[0]["n"], "<b>1 & 2</b>")

    def test_ok_tekst_rodzica_rekordow_zostaje_w_kontekscie(self):
        """Opis wpisany wprost w elemencie-rodzicu rekordów nie przepada."""
        doc = parse("<aukcja id='A-1'>Sprzet uzywany, odbior osobisty, brak zwrotow."
                    "<pozycja><nazwa>Laptop</nazwa></pozycja>"
                    "<pozycja><nazwa>Monitor</nazwa></pozycja></aukcja>")
        self.assertEqual(doc.context.get("aukcja@id"), "A-1")
        self.assertEqual(doc.context.get("aukcja"),
                         "Sprzet uzywany, odbior osobisty, brak zwrotow.")
        self.assertIn("aukcja", doc.columns)
        self.assertEqual([r["nazwa"] for r in doc.records], ["Laptop", "Monitor"])

# --------------------------------------------------------------------------- #
# 7. Kolizje z kolumnami technicznymi i eksplozja kolumn
# --------------------------------------------------------------------------- #

class TestKolumn(BazaCLI):

    def test_ok_tagi_o_nazwach_kolumn_technicznych(self):
        self.zapisz("kolizja.xml",
                    "<paczka nr='P-1'>"
                    "<item><Aukcja>W1</Aukcja><Plik>skan.pdf</Plik>"
                    "<spec name='Nr pozycji'>77</spec><nazwa>A</nazwa></item>"
                    "<item><Aukcja>W2</Aukcja><Plik>skan2.pdf</Plik>"
                    "<spec name='Nr pozycji'>78</spec><nazwa>B</nazwa></item></paczka>")
        kod, out, _ = self.uruchom()
        self.assertEqual(kod, 0, out)
        arkusz = self.arkusze()["Wszystkie aukcje"]
        self.assertEqual(arkusz[0][:3], ["Aukcja", "Plik", "Nr pozycji"])
        for nazwa in ("Aukcja (XML)", "Plik (XML)", "Nr pozycji (XML)"):
            self.assertIn(nazwa, arkusz[0])
        self.assertEqual(arkusz[1][arkusz[0].index("Aukcja (XML)")], "W1")

    def test_ok_eksplozja_kolumn_dzieli_arkusz_bez_utraty_danych(self):
        liczba = xlsxwrite.MAX_COLS_PER_SHEET + 400

        def pozycja(nr):
            return "<item>" + "".join(
                "<t%05d>v%d_%d</t%05d>" % (k, nr, k, k) for k in range(liczba)) + "</item>"

        self.zapisz("eksplozja.xml",
                    "<aukcja id='A9'>%s%s</aukcja>" % (pozycja(1), pozycja(2)))
        kod, out, err = self.uruchom()
        self.assertEqual(kod, 0, out)
        self.assertIn("limit Excela", err)
        arkusze = self.arkusze()
        self.assertIn("Wszystkie aukcje (2)", arkusze)
        pierwszy, drugi = arkusze["Wszystkie aukcje"], arkusze["Wszystkie aukcje (2)"]
        self.assertEqual(len(pierwszy[0]), xlsxwrite.MAX_COLS_PER_SHEET)
        self.assertEqual(drugi[0][-1], "t%05d" % (liczba - 1))
        self.assertEqual(drugi[1][-1], "v1_%d" % (liczba - 1))
        # Arkusz-kontynuacja powtarza kolumny techniczne, żeby dało się
        # połączyć wiersze z arkuszem pierwszym.
        self.assertEqual(drugi[0][:3], ["Aukcja", "Plik", "Nr pozycji"])
        self.assertEqual(drugi[1][:3], pierwszy[1][:3])

    def test_ok_podzial_wierszy_nie_gubi_danych(self):
        stare = xlsxwrite.MAX_ROWS_PER_SHEET
        xlsxwrite.MAX_ROWS_PER_SHEET = 5
        self.addCleanup(setattr, xlsxwrite, "MAX_ROWS_PER_SHEET", stare)
        wiersze = [[i, "a%d" % i] for i in range(1, 11)]
        arkusz = xlsxwrite.Sheet(name="Dane", columns=["k1", "k2"], rows=iter(wiersze))
        cel = os.path.join(self.katalog, "podzial.xlsx")
        with contextlib.redirect_stderr(io.StringIO()):
            xlsxwrite.write_workbook(cel, [arkusz])
        arkusze = self.arkusze(cel)
        zebrane = []
        for nazwa in arkusze:
            zebrane.extend(w for w in arkusze[nazwa][1:])
        self.assertEqual(zebrane, wiersze)


# --------------------------------------------------------------------------- #
# 8. Ogromne wartości i nagłówki
# --------------------------------------------------------------------------- #

class TestOgromnychWartosci(BazaCLI):

    XML_WIELKI_ATRYBUT = (
        "<aukcja id='A1'>"
        "<item><spec name='%s'>1</spec><sn>x1</sn></item>"
        "<item><spec name='%s'>2</spec><sn>x2</sn></item></aukcja>")

    def test_ok_naglowek_jest_przycinany_do_limitu_excela(self):
        """Nagłówek przechodzi przez ``sanitize_cell`` — także w zapisie stdlib."""
        nazwa = "A" * 40000
        self.zapisz("wielki.xml", self.XML_WIELKI_ATRYBUT % (nazwa, nazwa))
        cel = os.path.join(self.katalog, "stdlib.xlsx")
        kod, out, _ = self.uruchom(backend="stdlib", out=cel)
        self.assertEqual(kod, 0, out)
        with zipfile.ZipFile(cel) as paczka:
            tresc = paczka.read("xl/worksheets/sheet1.xml").decode("utf-8")
        self.assertNotIn("A" * (values.MAX_CELL_CHARS + 1), tresc)
        naglowek = self.arkusze(cel)["Wszystkie aukcje"][0][4]
        self.assertEqual(len(naglowek), values.MAX_CELL_CHARS)

    def test_ok_dwa_backendy_daja_ten_sam_naglowek(self):
        """Obie ścieżki zapisu przycinają nagłówek do tej samej długości."""
        nazwa = "A" * 40000
        self.zapisz("wielki.xml", self.XML_WIELKI_ATRYBUT % (nazwa, nazwa))
        a = os.path.join(self.katalog, "a.xlsx")
        b = os.path.join(self.katalog, "b.xlsx")
        self.uruchom(backend="openpyxl", out=a)
        self.uruchom(backend="stdlib", out=b)
        naglowek_a = self.arkusze(a)["Wszystkie aukcje"][0][4]
        naglowek_b = self.arkusze(b)["Wszystkie aukcje"][0][4]
        self.assertEqual(len(naglowek_a), values.MAX_CELL_CHARS)
        self.assertEqual(naglowek_a, naglowek_b)

    def test_ok_wielka_wartosc_komorki_jest_przycinana(self):
        wartosc = "B" * 50000
        self.zapisz("dlugi.xml",
                    "<aukcja id='A1'><item><opis>%s</opis></item>"
                    "<item><opis>krotki</opis></item></aukcja>" % wartosc)
        for backend in ("openpyxl", "stdlib"):
            with self.subTest(backend=backend):
                cel = os.path.join(self.katalog, "%s.xlsx" % backend)
                kod, out, _ = self.uruchom(backend=backend, out=cel)
                self.assertEqual(kod, 0, out)
                komorka = self.arkusze(cel)["Wszystkie aukcje"][1][4]
                self.assertEqual(len(komorka), values.MAX_CELL_CHARS)

    def test_ok_wstrzykniecie_formuly_zapisane_jako_tekst(self):
        self.zapisz("formula.xml",
                    "<aukcja id='A1'>"
                    "<item><uwaga>=SUM(A1:A9)</uwaga><b>-2+3</b><c>@SUM(1)</c></item>"
                    "<item><uwaga>+1</uwaga><b>ok</b><c>ok</c></item></aukcja>")
        for backend in ("openpyxl", "stdlib"):
            with self.subTest(backend=backend):
                cel = os.path.join(self.katalog, "f_%s.xlsx" % backend)
                kod, out, _ = self.uruchom(backend=backend, out=cel)
                self.assertEqual(kod, 0, out)
                skoroszyt = openpyxl.load_workbook(cel)
                arkusz = skoroszyt["Wszystkie aukcje"]
                for kolumna in (5, 6, 7):
                    komorka = arkusz.cell(row=2, column=kolumna)
                    self.assertEqual(komorka.data_type, "s", komorka.value)
                self.assertEqual(arkusz.cell(row=2, column=5).value, "=SUM(A1:A9)")


# --------------------------------------------------------------------------- #
# 9. Powtórzone rodzeństwo
# --------------------------------------------------------------------------- #

class TestPowtorzonegoRodzenstwa(BazaCLI):

    XML = ("<b><i><sn>1</sn><cecha>A</cecha><cecha></cecha><cecha>C</cecha></i>"
           "<i><sn>2</sn><cecha>A</cecha><cecha>B</cecha><cecha>C</cecha></i></b>")

    def test_ok_tryb_join_sklei_wartosci_bez_pozycji(self):
        """Świadomy kompromis trybu ``join``: 'A | C' vs 'A | B | C'.

        Sklejanie nie pokazuje, KTÓRA wartość była pusta — kto tego potrzebuje,
        używa ``repeat="index"`` (patrz test poniżej).
        """
        doc = parse(self.XML)
        self.assertEqual(doc.records[0]["cecha"], "A | C")
        self.assertEqual(doc.records[1]["cecha"], "A | B | C")

    def test_ok_tryb_index_zachowuje_pozycje(self):
        doc = parse(self.XML, repeat="index")
        self.assertEqual(doc.records[0]["cecha[2]"], None)
        self.assertEqual(doc.records[1]["cecha[2]"], "B")


# --------------------------------------------------------------------------- #
# 10. Odporność: fuzzing i duże dane
# --------------------------------------------------------------------------- #

class TestOdpornosci(BazaCLI):

    WZORZEC = (
        "<?xml version='1.0' encoding='UTF-8'?>\n<aukcja id='A-1' nr='7'>\n"
        " <meta><sprzedawca>Firma &amp; Syn</sprzedawca><data>2026-01-02</data></meta>\n"
        " <pozycje>\n"
        "  <pozycja id='1'><nazwa>Laptop</nazwa><spec name='RAM'>16 GB</spec>"
        "<cena waluta='PLN'>1 234,56</cena><opis><![CDATA[a<b]]></opis></pozycja>\n"
        "  <pozycja id='2'><nazwa>Monitor</nazwa><spec name='RAM'>8 GB</spec>"
        "<cena waluta='PLN'>99</cena></pozycja>\n"
        " </pozycje>\n</aukcja>\n").encode("utf-8")

    def test_ok_fuzzing_bajtowy_konczy_sie_zawsze_XmlParseError(self):
        """Żadna mutacja nie może wywołać wyjątku spoza kontraktu."""
        rnd = random.Random(20260908)
        for _ in range(1200):
            dane = bytearray(self.WZORZEC)
            for _ in range(rnd.randint(1, 6)):
                if not dane:
                    break
                pozycja = rnd.randrange(len(dane))
                los = rnd.random()
                if los < 0.4:
                    dane[pozycja] = rnd.randrange(256)
                elif los < 0.7:
                    del dane[pozycja:pozycja + rnd.randint(1, 12)]
                else:
                    dane[pozycja:pozycja] = bytes(
                        rnd.randrange(256) for _ in range(rnd.randint(1, 6)))
            try:
                xf.parse_bytes(bytes(dane), "fuzz.xml")
            except xf.XmlParseError:
                pass
            except Exception as exc:  # noqa: BLE001 - to jest właśnie badane
                self.fail("nieoczekiwany %s: %s\nDANE: %r"
                          % (type(exc).__name__, exc, bytes(dane)))

    def test_ok_fuzzing_strukturalny_nie_gubi_zadnej_wartosci(self):
        """Żadna wartość (tekst ani atrybut) nie może zniknąć z wyniku.

        300 losowych dokumentów o różnych kształtach: każda wartość musi dać się
        odnaleźć w kontekście albo w rekordach (jako wartość lub nazwa kolumny).
        """
        rnd = random.Random(7)
        tagi = ["pozycja", "spec", "opcja", "cecha", "name", "value", "nazwa", "lot", "x"]
        atrybuty = ["id", "name", "nazwa", "value", "wartosc", "typ"]

        def buduj(glebokosc, licznik):
            tag = rnd.choice(tagi)
            attrs = {rnd.choice(atrybuty): "V%05d" % next(licznik)
                     for _ in range(rnd.randint(0, 2))}
            dzieci, tekst = [], ""
            if glebokosc > 0 and rnd.random() < 0.7:
                dzieci = [buduj(glebokosc - 1, licznik) for _ in range(rnd.randint(1, 3))]
                if rnd.random() < 0.2:
                    tekst = "V%05d" % next(licznik)
            else:
                tekst = "V%05d" % next(licznik)
            return (tag, attrs, tekst, dzieci)

        def rysuj(wezel):
            tag, attrs, tekst, dzieci = wezel
            opis = "".join(' %s="%s"' % (k, v) for k, v in attrs.items())
            srodek = tekst + "".join(rysuj(k) for k in dzieci)
            if not srodek:
                return "<%s%s/>" % (tag, opis)
            return "<%s%s>%s</%s>" % (tag, opis, srodek, tag)

        def wartosci(wezel, zbior):
            tag, attrs, tekst, dzieci = wezel
            zbior.update(attrs.values())
            if tekst:
                zbior.add(tekst)
            for k in dzieci:
                wartosci(k, zbior)
            return zbior

        zgubione = 0
        proby = 300
        for numer in range(proby):
            licznik = itertools.count(1)
            korzen = ("aukcja", {"id": "A%04d" % numer}, "",
                      [buduj(3, licznik) for _ in range(rnd.randint(2, 5))])
            doc = parse("<?xml version='1.0' encoding='UTF-8'?>" + rysuj(korzen))
            blob = "\n".join(
                [str(v) for v in doc.context.values()]
                + [str(v) for r in doc.records for v in r.values()]
                + list(doc.context.keys())
                + [k for r in doc.records for k in r.keys()])
            if any(v not in blob for v in wartosci(korzen, set())):
                zgubione += 1
        self.assertEqual(zgubione, 0,
                         "%d z %d dokumentów zgubiło co najmniej jedną wartość"
                         % (zgubione, proby))

    def test_ok_duzy_plik_przechodzi_przez_potok(self):
        rekord = ("<pozycja><nr>%d</nr><nazwa>Serwer Dell R730 %d</nazwa>"
                  "<sn>SN%08d</sn><stan>uzywany</stan><cena>1 234,56</cena>"
                  "<opis>Opis pozycji %d, sprzet powystawowy</opis></pozycja>")
        sciezka = os.path.join(self.wejscie, "duzy.xml")
        with open(sciezka, "w", encoding="utf-8") as fh:
            fh.write("<?xml version='1.0' encoding='UTF-8'?><aukcja id='BIG-1'><pozycje>")
            fh.write("".join(rekord % (j, j, j, j) for j in range(8000)))
            fh.write("</pozycje></aukcja>")
        self.assertGreater(os.path.getsize(sciezka), 1_000_000)
        kod, out, _ = self.uruchom()
        self.assertEqual(kod, 0, out)
        arkusz = self.arkusze()["Wszystkie aukcje"]
        self.assertEqual(len(arkusz) - 1, 8000)
        self.assertEqual(arkusz[1][arkusz[0].index("cena")], 1234.56)


if __name__ == "__main__":
    unittest.main(verbosity=2)
