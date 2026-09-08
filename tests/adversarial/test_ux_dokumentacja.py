# -*- coding: utf-8 -*-
"""Testy ADWERSARYJNE: zgodność README.md z rzeczywistością (wymiar UX).

Punkt widzenia: **nietechniczny użytkownik z Polski**, Windows/„podstawy obsługi
komputera”, który kopiuje polecenia z ``README.md`` dosłownie, jedno po drugim,
w czystym katalogu.  Testy sprawdzają:

* czy polecenia z README dają obiecane efekty i **obiecane kody wyjścia**,
* czy komunikaty błędów są po polsku i mówią, co zrobić dalej,
* czy nazwy kolumn w wyniku da się zrozumieć bez znajomości XML-a,
* czy narzędzie działa bez ``openpyxl`` (symulacja ``ImportError``),
* czy część pobierająca zachowuje się tak, jak opisuje README
  (atrapa portalu na ``127.0.0.1`` — bez dostępu do sieci).

Wszystkie testy to **regresja README** — pilnują, żeby udokumentowane
zachowania nie zniknęły.  Historycznie moduł zawierał testy ``test_USTERKA_*``
utrwalające znalezione wady; po ich naprawieniu każdy został przepisany na
asercję stanu poprawnego.

Uruchamianie::

    python3 -m unittest tests.adversarial.test_ux_dokumentacja
    python3 tests/adversarial/test_ux_dokumentacja.py

UWAGA: ``python3 -m unittest discover -s tests`` (polecenie „cały zestaw”
z README) NIE zagląda do katalogu ``tests/adversarial`` — brakuje w nim
``__init__.py``, więc discovery go pomija bez słowa ostrzeżenia.
"""

import os
import re
import shlex
import shutil
import signal
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

_KORZEN = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if _KORZEN not in sys.path:
    sys.path.insert(0, _KORZEN)

import openpyxl  # noqa: E402  (dozwolone w testach — tylko do weryfikacji wyniku)

from flexit2xlsx import values  # noqa: E402

README = os.path.join(_KORZEN, "README.md")


# --------------------------------------------------------------------------- #
# Narzędzia: uruchamianie CLI dokładnie tak, jak robi to użytkownik
# --------------------------------------------------------------------------- #


def uruchom(argumenty, katalog, srodowisko=None, limit_czasu=180):
    """Uruchamia ``python3 -m flexit2xlsx ...`` jak z wiersza poleceń.

    Zwraca obiekt ``CompletedProcess`` z ``stdout``/``stderr`` w UTF-8.
    ``katalog`` to katalog roboczy (użytkownik robi ``cd`` do projektu).
    """
    env = dict(os.environ)
    env["PYTHONIOENCODING"] = "utf-8"
    # atrapa portalu stoi na 127.0.0.1 — proxy nie może się wtrącać
    env["no_proxy"] = "*"
    env["NO_PROXY"] = "*"
    if srodowisko:
        env.update(srodowisko)
    return subprocess.run(
        [sys.executable, "-m", "flexit2xlsx"] + list(argumenty),
        cwd=katalog,
        env=env,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        timeout=limit_czasu,
    )


def tekst(strumien):
    """Bajty z podprocesu -> tekst (UTF-8, bez wywracania się na krzakach)."""
    return (strumien or b"").decode("utf-8", "replace")


def czytaj_arkusz(sciezka, nazwa="Wszystkie aukcje"):
    """Zwraca listę krotek: [nagłówek, wiersz1, wiersz2, ...]."""
    skoroszyt = openpyxl.load_workbook(sciezka)
    return list(skoroszyt[nazwa].iter_rows(values_only=True))


def zapisz_xml(sciezka, tresc, kodowanie="utf-8"):
    """Zapisuje plik XML w podanym kodowaniu (użytkownik pobiera je z portalu)."""
    os.makedirs(os.path.dirname(sciezka), exist_ok=True)
    with open(sciezka, "wb") as uchwyt:
        uchwyt.write(tresc.encode(kodowanie))


#: Realistyczna „zawartość pakietu” — trzy sztuki sprzętu w jednym locie.
PAKIET_TRZY = """<?xml version="1.0" encoding="UTF-8"?>
<batch lotHash="e0c8f" auction="flexit-auctions-26-02-2026-1087">
  <lot>
    <title>12x Lenovo 8th-10th Gen Laptop Mix</title>
    <items>
      <item nr="1"><model>ThinkPad T480</model><ram>8 GB</ram><serial>PF1A2B3C</serial><grade>B</grade></item>
      <item nr="2"><model>ThinkPad T490</model><ram>16 GB</ram><serial>PF4D5E6F</serial><grade>A</grade></item>
      <item nr="3"><model>ThinkPad L390</model><ram>4 GB</ram><serial>PF7G8H9I</serial><grade>C</grade></item>
    </items>
  </lot>
</batch>
"""

#: Pakiet z JEDNĄ sztuką — pułapka na wykrywanie elementu powtarzalnego.
PAKIET_JEDEN = """<?xml version="1.0" encoding="UTF-8"?>
<batch lotHash="53439" auction="flexit-auctions-26-02-2026-1087">
  <lot>
    <title>10x HP EliteDesk Mix</title>
    <items>
      <item nr="1"><model>EliteDesk 800 G4</model><ram>8 GB</ram><serial>HP0001</serial><grade>B</grade></item>
    </items>
  </lot>
</batch>
"""


class BazaUX(unittest.TestCase):
    """Wspólny „czysty katalog projektu” — dokładnie to, co dostaje użytkownik."""

    @classmethod
    def setUpClass(cls):
        cls._tymczasowy = tempfile.mkdtemp(prefix="flexit-ux-")
        # użytkownik rozpakowuje projekt na pulpit: kopiujemy TYLKO to,
        # co jest potrzebne do uruchomienia (pakiet + README)
        cls.projekt = os.path.join(cls._tymczasowy, "Pulpit", "triwolak")
        os.makedirs(cls.projekt)
        shutil.copytree(
            os.path.join(_KORZEN, "flexit2xlsx"),
            os.path.join(cls.projekt, "flexit2xlsx"),
            ignore=shutil.ignore_patterns("__pycache__"),
        )
        shutil.copy2(README, os.path.join(cls.projekt, "README.md"))

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls._tymczasowy, ignore_errors=True)

    def cli(self, argumenty, katalog=None, srodowisko=None, limit_czasu=180):
        """Uruchamia CLI z ``PYTHONPATH`` wskazującym skopiowany projekt.

        Użytkownik robi ``cd`` do katalogu projektu; w testach katalog roboczy
        bywa podkatalogiem z danymi, więc pakiet wskazujemy przez ``PYTHONPATH``.
        """
        srodowisko = dict(srodowisko or {})
        sciezki = []
        if "PYTHONPATH" in srodowisko:
            sciezki.append(srodowisko.pop("PYTHONPATH"))
        sciezki.append(self.projekt)
        srodowisko["PYTHONPATH"] = os.pathsep.join(sciezki)
        return uruchom(argumenty, katalog or self.projekt, srodowisko, limit_czasu)

    def katalog_roboczy(self, nazwa):
        """Osobny podkatalog na dane jednego testu (żeby testy się nie mieszały)."""
        sciezka = os.path.join(self.projekt, nazwa)
        if os.path.isdir(sciezka):
            shutil.rmtree(sciezka)
        os.makedirs(sciezka)
        return sciezka


# --------------------------------------------------------------------------- #
# 1. Regresja README — polecenia „krok po kroku” muszą działać dosłownie
# --------------------------------------------------------------------------- #


class TestPoleceniaZReadme(BazaUX):
    """Polecenia skopiowane z README dosłownie, w podanej kolejności."""

    def test_help_dziala_i_jest_zrozumialy(self):
        """README, Windows krok 5: ``py -m flexit2xlsx --help``."""
        wynik = self.cli(["--help"], self.projekt)
        self.assertEqual(wynik.returncode, 0)
        pomoc = tekst(wynik.stdout)
        for podkomenda in ("download", "build", "all"):
            self.assertIn(podkomenda, pomoc)
        # opisy podkomend są po polsku
        self.assertIn("zbuduj jeden plik .xlsx", pomoc)

    def test_version(self):
        """README, tabela opcji wspólnych: ``--version``."""
        wynik = self.cli(["--version"], self.projekt)
        self.assertEqual(wynik.returncode, 0)
        self.assertRegex(tekst(wynik.stdout), r"flexit2xlsx \d+\.\d+")

    def test_build_z_katalogu_konczy_sie_kodem_zero(self):
        """README, „Przykłady poleceń” nr 1 — najprostszy scenariusz."""
        praca = self.katalog_roboczy("t_build")
        zapisz_xml(os.path.join(praca, "xml_flexit", "a.xml"), PAKIET_TRZY)
        zapisz_xml(os.path.join(praca, "xml_flexit", "b.xml"), PAKIET_JEDEN)

        wynik = self.cli(
            ["build", "--in", "xml_flexit", "--out", "aukcje.xlsx"], praca
        )
        self.assertEqual(wynik.returncode, 0, tekst(wynik.stderr))
        self.assertTrue(os.path.isfile(os.path.join(praca, "aukcje.xlsx")))

        wiersze = czytaj_arkusz(os.path.join(praca, "aukcje.xlsx"))
        # 3 + 1 pozycje, nic nie ginie
        self.assertEqual(len(wiersze) - 1, 4)
        # kolumny techniczne z README, w podanej kolejności
        self.assertEqual(wiersze[0][:3], ("Aukcja", "Plik", "Nr pozycji"))

    def test_kolumny_ujednolicone_dla_pliku_z_jedna_pozycja(self):
        """README, sekcja „O ``--no-unify``”: plik z jedną pozycją ma te same kolumny."""
        praca = self.katalog_roboczy("t_unify")
        zapisz_xml(os.path.join(praca, "xml", "duzy.xml"), PAKIET_TRZY)
        zapisz_xml(os.path.join(praca, "xml", "maly.xml"), PAKIET_JEDEN)

        self.cli(["build", "--in", "xml", "--out", "a.xlsx"], praca)
        wiersze = czytaj_arkusz(os.path.join(praca, "a.xlsx"))
        naglowek = list(wiersze[0])
        self.assertIn("model", naglowek)
        # przed poprawką pojawiłaby się druga kolumna z pełną ścieżką
        self.assertNotIn("lot/items/item/model", naglowek)
        # a z --no-unify README obiecuje zachowanie SPRZED poprawki:
        # plik z jedną pozycją dostaje kolumny z PEŁNĄ ścieżką, czyli te same
        # dane trafiają do dwóch różnych kolumn arkusza
        self.cli(["build", "--in", "xml", "--out", "b.xlsx", "--no-unify"], praca)
        naglowek_bez = list(czytaj_arkusz(os.path.join(praca, "b.xlsx"))[0])
        self.assertIn("model", naglowek_bez)
        self.assertIn("lot/items/item/model", naglowek_bez)

    def test_pojedyncze_pliki_zamiast_katalogu(self):
        """README, „Przykłady poleceń” nr 2: ``--in plik1.xml plik2.xml``."""
        praca = self.katalog_roboczy("t_pliki")
        zapisz_xml(os.path.join(praca, "pakiet1.xml"), PAKIET_TRZY)
        zapisz_xml(os.path.join(praca, "pakiet2.xml"), PAKIET_JEDEN)
        wynik = self.cli(
            ["build", "--in", "pakiet1.xml", "pakiet2.xml", "--out", "aukcje.xlsx"],
            praca,
        )
        self.assertEqual(wynik.returncode, 0, tekst(wynik.stderr))
        self.assertEqual(len(czytaj_arkusz(os.path.join(praca, "aukcje.xlsx"))) - 1, 4)

    def test_dry_run_niczego_nie_zapisuje(self):
        """README, „Przykłady poleceń” nr 5: ``--dry-run``."""
        praca = self.katalog_roboczy("t_dry")
        zapisz_xml(os.path.join(praca, "xml", "a.xml"), PAKIET_TRZY)
        wynik = self.cli(
            ["build", "--in", "xml", "--out", "podglad.xlsx", "--dry-run"], praca
        )
        self.assertEqual(wynik.returncode, 0)
        self.assertFalse(os.path.exists(os.path.join(praca, "podglad.xlsx")))
        self.assertIn("dry-run", tekst(wynik.stdout))

    def test_per_auction_nie_gubi_wierszy(self):
        """README, „Arkusze per aukcja”: osobna zakładka na aukcję."""
        praca = self.katalog_roboczy("t_per")
        zapisz_xml(os.path.join(praca, "xml", "a.xml"), PAKIET_TRZY)
        zapisz_xml(
            os.path.join(praca, "xml", "b.xml"),
            PAKIET_JEDEN.replace("26-02-2026-1087", "18-06-2026-1103"),
        )
        wynik = self.cli(
            ["build", "--in", "xml", "--out", "a.xlsx", "--per-auction"], praca
        )
        self.assertEqual(wynik.returncode, 0, tekst(wynik.stderr))
        skoroszyt = openpyxl.load_workbook(os.path.join(praca, "a.xlsx"))
        zbiorczy = skoroszyt["Wszystkie aukcje"].max_row - 1
        per = sum(
            skoroszyt[n].max_row - 1
            for n in skoroszyt.sheetnames
            if n not in ("Wszystkie aukcje", "Podsumowanie")
        )
        self.assertEqual(zbiorczy, per)
        self.assertEqual(zbiorczy, 4)

    def test_repeat_index_tworzy_kolumny_z_nawiasami(self):
        """README, „Przykłady poleceń” nr 7: ``--repeat index`` -> ``cecha[1]``."""
        praca = self.katalog_roboczy("t_repeat")
        zapisz_xml(
            os.path.join(praca, "xml", "a.xml"),
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            "<batch><items>"
            "<item><model>T480</model><cecha>8GB</cecha><cecha>SSD</cecha></item>"
            "<item><model>T490</model><cecha>16GB</cecha></item>"
            "</items></batch>\n",
        )
        self.cli(["build", "--in", "xml", "--out", "j.xlsx"], praca)
        self.cli(["build", "--in", "xml", "--out", "i.xlsx", "--repeat", "index"], praca)
        self.assertIn("cecha", czytaj_arkusz(os.path.join(praca, "j.xlsx"))[0])
        naglowek = czytaj_arkusz(os.path.join(praca, "i.xlsx"))[0]
        self.assertIn("cecha[1]", naglowek)
        self.assertIn("cecha[2]", naglowek)

    def test_polskie_znaki_i_sciezki_ze_spacjami(self):
        """Pulpit „Jan Kowalski” + katalog „xml zażółć” — typowe na Windows."""
        praca = self.katalog_roboczy("t_pl")
        katalog = os.path.join(praca, "Jan Kowalski", "xml zażółć")
        zapisz_xml(
            os.path.join(katalog, "pakiet.xml"),
            '<?xml version="1.0" encoding="ISO-8859-2"?>\n'
            "<batch auction=\"AUK-PL\"><pozycje>"
            "<pozycja><model>Zażółć gęślą jaźń</model></pozycja>"
            "<pozycja><model>Łódź Ćma</model></pozycja>"
            "</pozycje></batch>\n",
            kodowanie="iso-8859-2",
        )
        wyjscie = os.path.join(praca, "Jan Kowalski", "moje aukcje.xlsx")
        wynik = self.cli(
            ["build", "--in", katalog, "--out", wyjscie], praca
        )
        self.assertEqual(wynik.returncode, 0, tekst(wynik.stderr))
        wiersze = czytaj_arkusz(wyjscie)
        modele = [w[wiersze[0].index("model")] for w in wiersze[1:]]
        self.assertEqual(modele, ["Zażółć gęślą jaźń", "Łódź Ćma"])

    def test_nazwy_plikow_z_przegladarki(self):
        """README krok 4: „Nazwy plików mogą być dowolne”."""
        praca = self.katalog_roboczy("t_nazwy")
        katalog = os.path.join(praca, "xml")
        zapisz_xml(os.path.join(katalog, "Download Batch Details.xml"), PAKIET_TRZY)
        zapisz_xml(os.path.join(katalog, "Download Batch Details (1).xml"), PAKIET_JEDEN)
        zapisz_xml(os.path.join(katalog, "BATCH_DETAILS.XML"), PAKIET_JEDEN)
        wynik = self.cli(["build", "--in", "xml", "--out", "a.xlsx"], praca)
        self.assertEqual(wynik.returncode, 0, tekst(wynik.stderr))
        self.assertEqual(len(czytaj_arkusz(os.path.join(praca, "a.xlsx"))) - 1, 5)

    def test_snippet_uzycie_jako_biblioteki_z_readme(self):
        """README, „Dla programistów / Użycie jako biblioteki” — dosłownie."""
        praca = self.katalog_roboczy("t_lib")
        zapisz_xml(os.path.join(praca, "xml_flexit", "a.xml"), PAKIET_TRZY)
        skrypt = (
            "from flexit2xlsx import cli, xmlflatten, xlsxwrite\n"
            "\n"
            "docs, errors = cli.load_documents(cli.collect_xml_files(['xml_flexit']))\n"
            "xlsxwrite.write_workbook('aukcje.xlsx', cli.build_sheets(docs, errors))\n"
            "print('DOCS', len(docs), 'ERRORS', len(errors))\n"
        )
        env = dict(os.environ)
        env["PYTHONPATH"] = self.projekt
        wynik = subprocess.run(
            [sys.executable, "-c", skrypt],
            cwd=praca,
            env=env,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=120,
        )
        self.assertEqual(wynik.returncode, 0, tekst(wynik.stderr))
        self.assertIn("DOCS 1 ERRORS 0", tekst(wynik.stdout))
        self.assertTrue(os.path.isfile(os.path.join(praca, "aukcje.xlsx")))


# --------------------------------------------------------------------------- #
# 2. Kody wyjścia z README
# --------------------------------------------------------------------------- #


class TestKodyWyjscia(BazaUX):
    """Tabela „Kody wyjścia” z README — każdy kod musi dać się wywołać."""

    def test_kod_1_plik_juz_istnieje(self):
        praca = self.katalog_roboczy("k1")
        zapisz_xml(os.path.join(praca, "xml", "a.xml"), PAKIET_TRZY)
        self.cli(["build", "--in", "xml", "--out", "aukcje.xlsx"], praca)
        wynik = self.cli(["build", "--in", "xml", "--out", "aukcje.xlsx"], praca)
        self.assertEqual(wynik.returncode, 1)
        komunikat = tekst(wynik.stderr)
        self.assertIn("już istnieje", komunikat)
        self.assertIn("--overwrite", komunikat)

    def test_kod_2_bledna_skladnia(self):
        praca = self.katalog_roboczy("k2")
        wynik = self.cli(["build", "--outt", "a.xlsx"], praca)
        self.assertEqual(wynik.returncode, 2)

    def test_kod_3_czesciowy_sukces(self):
        """Jeden plik uszkodzony nie może przerwać całości (README: kod 3)."""
        praca = self.katalog_roboczy("k3")
        zapisz_xml(os.path.join(praca, "xml", "dobry.xml"), PAKIET_TRZY)
        zapisz_xml(
            os.path.join(praca, "xml", "logowanie.xml"),
            "<!DOCTYPE html><html><head><title>Sign in</title></head>"
            "<body><form><input name='email'></form></body></html>",
        )
        zapisz_xml(os.path.join(praca, "xml", "przerwany.xml"), "")
        wynik = self.cli(["build", "--in", "xml", "--out", "a.xlsx"], praca)
        self.assertEqual(wynik.returncode, 3, tekst(wynik.stderr))
        self.assertTrue(os.path.isfile(os.path.join(praca, "a.xlsx")))
        # README: „Pliki, których nie udało się wczytać, też tu są — ze statusem BŁĄD”
        podsumowanie = czytaj_arkusz(os.path.join(praca, "a.xlsx"), "Podsumowanie")
        statusy = [w[-1] for w in podsumowanie[1:]]
        self.assertEqual(sum(1 for s in statusy if str(s).startswith("BŁĄD")), 2)
        self.assertTrue(any(str(w[0]) == "RAZEM" for w in podsumowanie))

    def test_kod_4_brak_plikow_xml(self):
        praca = self.katalog_roboczy("k4")
        os.makedirs(os.path.join(praca, "xml_flexit"))
        wynik = self.cli(["build", "--in", "xml_flexit", "--out", "a.xlsx"], praca)
        self.assertEqual(wynik.returncode, 4)
        # komunikat z README, tabela „Najczęstsze problemy”
        self.assertIn("Nie znalazłem żadnego pliku .xml", tekst(wynik.stderr))

    def test_kod_4_literowka_w_sciezce(self):
        """Najczęstszy błąd użytkownika: literówka w nazwie katalogu."""
        praca = self.katalog_roboczy("k4b")
        wynik = self.cli(["build", "--in", "xml_flexitt", "--out", "a.xlsx"], praca)
        self.assertEqual(wynik.returncode, 4)
        komunikat = tekst(wynik.stderr)
        self.assertIn("nie ma takiego pliku ani katalogu", komunikat)
        self.assertIn("Sprawdź ścieżkę", komunikat)

    def test_kod_5_portal_nieosiagalny(self):
        """README: kod 5 = błąd komunikacji z portalem."""
        praca = self.katalog_roboczy("k5")
        # port, na którym na pewno nikt nie słucha
        wynik = self.cli(
            [
                "download", "--out", "xml", "--base-url", "http://127.0.0.1:9/",
                "--retries", "0", "--timeout", "1", "--delay", "0",
            ],
            praca,
        )
        self.assertEqual(wynik.returncode, 5, tekst(wynik.stderr))
        komunikat = tekst(wynik.stderr)
        self.assertIn("BŁĄD", komunikat)
        self.assertIn("Podpowiedzi", komunikat)

    def test_kod_130_przerwanie_ctrl_c(self):
        """README: kod 130 = przerwane przez użytkownika (Ctrl+C).

        Portal udaje bardzo wolny serwer, żeby program na pewno jeszcze
        pracował w chwili wysłania sygnału.
        """
        praca = self.katalog_roboczy("k130")

        class WolnyHandler(BaseHTTPRequestHandler):
            def do_GET(self):  # noqa: N802
                time.sleep(30)

            def log_message(self, *a):
                pass

        serwer = ThreadingHTTPServer(("127.0.0.1", 0), WolnyHandler)
        serwer.daemon_threads = True
        threading.Thread(target=serwer.serve_forever, daemon=True).start()
        baza = "http://127.0.0.1:%d/" % serwer.server_address[1]

        env = dict(os.environ)
        env["PYTHONIOENCODING"] = "utf-8"
        env["no_proxy"] = "*"
        env["NO_PROXY"] = "*"
        env["PYTHONPATH"] = self.projekt
        proces = subprocess.Popen(
            [
                sys.executable, "-m", "flexit2xlsx", "download",
                "--out", "xml", "--base-url", baza,
                "--timeout", "60", "--retries", "0", "--delay", "0",
            ],
            cwd=praca,
            env=env,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        try:
            time.sleep(2.0)
            self.assertIsNone(proces.poll(), "program powinien jeszcze pracować")
            proces.send_signal(signal.SIGINT)  # to samo co Ctrl+C
            _, blad = proces.communicate(timeout=30)
        finally:
            if proces.poll() is None:  # pragma: no cover - zabezpieczenie
                proces.kill()
                proces.communicate()
            serwer.shutdown()
            serwer.server_close()
        self.assertEqual(proces.returncode, 130)
        self.assertIn("Przerwano przez użytkownika", tekst(blad))


# --------------------------------------------------------------------------- #
# 3. Praca bez openpyxl (README: „Biblioteki zewnętrzne — żadne”)
# --------------------------------------------------------------------------- #


class TestBezOpenpyxl(BazaUX):
    """README obiecuje, że bez ``openpyxl`` program nadal zapisuje .xlsx."""

    def _atrapa_bez_openpyxl(self):
        """Katalog na ``PYTHONPATH`` z modułem ``openpyxl`` rzucającym ImportError."""
        katalog = os.path.join(self.projekt, "_bez_openpyxl")
        os.makedirs(katalog, exist_ok=True)
        with open(os.path.join(katalog, "openpyxl.py"), "w", encoding="utf-8") as uchwyt:
            uchwyt.write('raise ImportError("symulacja: brak openpyxl")\n')
        return katalog

    def test_zapis_bez_zainstalowanego_openpyxl(self):
        praca = self.katalog_roboczy("bez")
        zapisz_xml(os.path.join(praca, "xml", "a.xml"), PAKIET_TRZY)
        wynik = self.cli(
            ["build", "--in", "xml", "--out", "a.xlsx"],
            praca,
            srodowisko={"PYTHONPATH": self._atrapa_bez_openpyxl()},
        )
        self.assertEqual(wynik.returncode, 0, tekst(wynik.stderr))
        self.assertIn("stdlib", tekst(wynik.stdout))
        # plik czytelny NIEZALEŻNYM narzędziem (openpyxl w tym procesie działa)
        wiersze = czytaj_arkusz(os.path.join(praca, "a.xlsx"))
        self.assertEqual(len(wiersze) - 1, 3)

    def test_backend_stdlib_daje_ten_sam_wynik_co_openpyxl(self):
        """README, „Wymuszenie sposobu zapisu XLSX”: oba backendy = te same dane."""
        praca = self.katalog_roboczy("dwa_backendy")
        zapisz_xml(os.path.join(praca, "xml", "a.xml"), PAKIET_TRZY)
        zapisz_xml(os.path.join(praca, "xml", "b.xml"), PAKIET_JEDEN)
        self.cli(
            ["build", "--in", "xml", "--out", "op.xlsx"],
            praca,
            srodowisko={"FLEXIT_XLSX_BACKEND": "openpyxl"},
        )
        self.cli(
            ["build", "--in", "xml", "--out", "st.xlsx"],
            praca,
            srodowisko={"FLEXIT_XLSX_BACKEND": "stdlib"},
        )
        self.assertEqual(
            czytaj_arkusz(os.path.join(praca, "op.xlsx")),
            czytaj_arkusz(os.path.join(praca, "st.xlsx")),
        )

    def test_wymuszenie_openpyxl_gdy_go_brak_nie_wywala_programu(self):
        praca = self.katalog_roboczy("wymus")
        zapisz_xml(os.path.join(praca, "xml", "a.xml"), PAKIET_TRZY)
        wynik = self.cli(
            ["build", "--in", "xml", "--out", "a.xlsx"],
            praca,
            srodowisko={
                "PYTHONPATH": self._atrapa_bez_openpyxl(),
                "FLEXIT_XLSX_BACKEND": "openpyxl",
            },
        )
        self.assertEqual(wynik.returncode, 0, tekst(wynik.stderr))
        self.assertIn("nie jest zainstalowana", tekst(wynik.stderr))


# --------------------------------------------------------------------------- #
# 4. Atrapa portalu — sprawdzenie części `download` opisanej w README
# --------------------------------------------------------------------------- #


class _AtrapaHandler(BaseHTTPRequestHandler):
    """Serwuje statyczne strony z ``self.server.strony``."""

    def do_GET(self):  # noqa: N802 - nazwa narzucona przez BaseHTTPRequestHandler
        wpis = self.server.strony.get(self.path)
        if wpis is None:
            self.send_error(404)
            return
        dane = wpis[0].encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", wpis[1])
        self.send_header("Content-Length", str(len(dane)))
        self.end_headers()
        self.wfile.write(dane)

    def log_message(self, *argumenty):  # cisza w wynikach testów
        pass


class AtrapaPortalu:
    """Mały serwer HTTP na losowym porcie — zastępuje niedostępny portal."""

    def __init__(self, strony):
        self.serwer = ThreadingHTTPServer(("127.0.0.1", 0), _AtrapaHandler)
        self.serwer.strony = strony
        self.serwer.daemon_threads = True
        self.watek = threading.Thread(target=self.serwer.serve_forever, daemon=True)
        self.watek.start()

    @property
    def base(self):
        return "http://127.0.0.1:%d/" % self.serwer.server_address[1]

    def stop(self):
        self.serwer.shutdown()
        self.serwer.server_close()
        self.watek.join(timeout=5)


HTML = "text/html; charset=utf-8"
XML = "application/xml"


def strony_portalu():
    """Portal w układzie ze SITE_NOTES: lista -> aukcja -> loty -> XML."""
    return {
        "/": (
            "<html><head><title>Online Auctions</title></head><body>"
            "<a href='/auction/flexit-auctions-18-06-2026-1103'>Aukcja 1103</a>"
            "</body></html>",
            HTML,
        ),
        # strona aukcji ma WŁASNY zbiorczy XML *i* listę lotów
        "/auction/flexit-auctions-18-06-2026-1103": (
            "<html><head><title>Aukcja 1103</title></head><body>"
            "<a href='/media/batch-1103-all.xml'>Download Batch Details (whole auction)</a>"
            "<a href='/lot/thinkpad-t480-mix-bcd12'>15x ThinkPad T480 Mix</a>"
            "</body></html>",
            HTML,
        ),
        "/lot/thinkpad-t480-mix-bcd12": (
            "<html><head><title>Lot bcd12</title></head><body>"
            "<a href='/media/batch-1103-bcd12.xml'>Download Batch Details</a>"
            "</body></html>",
            HTML,
        ),
        "/media/batch-1103-all.xml": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<batch auction="flexit-auctions-18-06-2026-1103"><lot><items>'
            "<item nr='1'><model>Whole auction export</model></item>"
            "</items></lot></batch>",
            XML,
        ),
        "/media/batch-1103-bcd12.xml": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<batch lotHash="bcd12" auction="flexit-auctions-18-06-2026-1103"><lot><items>'
            "<item nr='1'><model>ThinkPad T480</model><serial>TP0001</serial></item>"
            "<item nr='2'><model>ThinkPad T480</model><serial>TP0002</serial></item>"
            "</items></lot></batch>",
            XML,
        ),
    }


class TestPobieranieZAtrapy(BazaUX):
    """Część ``download``/``all`` na atrapie portalu (bez dostępu do sieci)."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.portal = AtrapaPortalu(strony_portalu())

    @classmethod
    def tearDownClass(cls):
        cls.portal.stop()
        super().tearDownClass()

    def test_all_pobiera_i_buduje(self):
        """README, Windows krok 6: jedno polecenie ``all``."""
        praca = self.katalog_roboczy("p_all")
        wynik = self.cli(
            [
                "all", "--in", "xml_flexit", "--out", "aukcje.xlsx",
                "--base-url", self.portal.base, "--delay", "0",
            ],
            praca,
        )
        self.assertEqual(wynik.returncode, 0, tekst(wynik.stderr))
        self.assertTrue(os.path.isfile(os.path.join(praca, "aukcje.xlsx")))

    def test_diagnose_wypisuje_co_widac_na_stronie(self):
        """README, „Gdy skrypt nie znajduje aukcji lub XML-i”, krok 1."""
        praca = self.katalog_roboczy("p_diag")
        portal = AtrapaPortalu(
            {"/": ("<html><head><title>Pusto</title></head><body></body></html>", HTML)}
        )
        try:
            wynik = self.cli(
                [
                    "download", "--out", "xml", "--dry-run", "--diagnose",
                    "--base-url", portal.base, "--delay", "0",
                ],
                praca,
            )
        finally:
            portal.stop()
        wyjscie = tekst(wynik.stdout) + tekst(wynik.stderr)
        self.assertIn("Tytuł strony", wyjscie)
        self.assertIn("Linków razem", wyjscie)

    def test_strona_logowania_zamiast_xml_daje_wskazowke_o_cookie(self):
        """README, „Logowanie i ciasteczka”: objaw wygasłej sesji."""
        praca = self.katalog_roboczy("p_login")
        strony = strony_portalu()
        strony["/media/batch-1103-all.xml"] = (
            "<!DOCTYPE html><html><head><title>Sign in</title></head>"
            "<body><form method='post'><input name='password'></form></body></html>",
            HTML,
        )
        strony["/media/batch-1103-bcd12.xml"] = strony["/media/batch-1103-all.xml"]
        portal = AtrapaPortalu(strony)
        try:
            wynik = self.cli(
                [
                    "download", "--out", "xml", "--base-url", portal.base,
                    "--delay", "0", "--descend", "always",
                ],
                praca,
            )
        finally:
            portal.stop()
        komunikat = tekst(wynik.stderr)
        self.assertIn("nie jest XML-em", komunikat)
        self.assertIn("--cookie", komunikat)

    def test_ciasteczko_nie_wycieka_przy_przekierowaniu_na_obca_domene(self):
        """README, „Uwagi bezpieczeństwa”: przekierowanie poza witrynę jest przerywane.

        Sprawdzamy realnie: obcy serwer zapisuje każdy otrzymany nagłówek.
        """
        praca = self.katalog_roboczy("p_cookie")
        otrzymane = []

        class ObcyHandler(BaseHTTPRequestHandler):
            def do_GET(self):  # noqa: N802
                otrzymane.append(dict(self.headers))
                dane = b"<batch><items><item><a>1</a></item></items></batch>"
                self.send_response(200)
                self.send_header("Content-Type", XML)
                self.send_header("Content-Length", str(len(dane)))
                self.end_headers()
                self.wfile.write(dane)

            def log_message(self, *a):
                pass

        obcy = ThreadingHTTPServer(("127.0.0.1", 0), ObcyHandler)
        threading.Thread(target=obcy.serve_forever, daemon=True).start()
        obcy_url = "http://127.0.0.1:%d/kradziez.xml" % obcy.server_address[1]

        class PrzekierowanieHandler(_AtrapaHandler):
            def do_GET(self):  # noqa: N802
                if self.path == "/redirect-out":
                    self.send_response(302)
                    self.send_header("Location", obcy_url)
                    self.end_headers()
                    return
                _AtrapaHandler.do_GET(self)

        serwer = ThreadingHTTPServer(("localhost", 0), PrzekierowanieHandler)
        serwer.strony = {
            "/": (
                "<html><head><title>Online Auctions</title></head><body>"
                "<a href='/auction/flexit-auctions-01-01-2026-9001'>A</a></body></html>",
                HTML,
            ),
            "/auction/flexit-auctions-01-01-2026-9001": (
                "<html><head><title>A</title></head><body>"
                "<a href='/lot/testowy-lot-abcde'>Lot</a></body></html>",
                HTML,
            ),
            "/lot/testowy-lot-abcde": (
                "<html><head><title>L</title></head><body>"
                "<a href='/redirect-out'>Download Batch Details</a></body></html>",
                HTML,
            ),
        }
        serwer.daemon_threads = True
        watek = threading.Thread(target=serwer.serve_forever, daemon=True)
        watek.start()
        baza = "http://localhost:%d/" % serwer.server_address[1]
        try:
            self.cli(
                [
                    "download", "--out", "xml", "--base-url", baza, "--delay", "0",
                    "--descend", "always",
                    "--cookie", "sessionid=TAJNE_HASLO_SESJI",
                ],
                praca,
            )
        finally:
            serwer.shutdown()
            serwer.server_close()
            obcy.shutdown()
            obcy.server_close()

        self.assertEqual(
            otrzymane, [], "obcy serwer NIE MOŻE dostać żadnego żądania (ani ciasteczka)"
        )

    # ------------------------------------------------------------------ #
    # USTERKI
    # ------------------------------------------------------------------ #

    def test_descend_auto_nie_pomija_lotow(self):
        """``--descend auto`` schodzi na strony lotów także wtedy, gdy XML JEST.

        Portal, który obok zbiorczego „whole auction export'' trzyma właściwą
        zawartość pakietów na stronach pojedynczych lotów (dokładnie taki układ
        opisuje ``SITE_NOTES.md``), nie może po cichu oddać tylko tego jednego
        zbiorczego pliku — reszta lotów zniknęłaby z arkusza bez ostrzeżenia.
        """
        praca = self.katalog_roboczy("u_descend")
        wynik_auto = self.cli(
            [
                "download", "--out", "auto", "--dry-run",
                "--base-url", self.portal.base, "--delay", "0",
            ],
            praca,
        )
        wynik_always = self.cli(
            [
                "download", "--out", "always", "--dry-run", "--descend", "always",
                "--base-url", self.portal.base, "--delay", "0",
            ],
            praca,
        )
        auto = tekst(wynik_auto.stdout)
        always = tekst(wynik_always.stdout)

        # domyślny tryb widzi ZARÓWNO plik zbiorczy, JAK I plik lotu
        self.assertIn("batch-1103-all.xml", auto)
        self.assertIn("batch-1103-bcd12.xml", auto)
        self.assertIn("batch-1103-bcd12.xml", always)
        # bez duplikatów: każdy plik wymieniony raz
        self.assertEqual(auto.count("batch-1103-bcd12.xml  ->"), 1, auto)

    def test_all_sprawdza_plik_wyjsciowy_zanim_zacznie_pobierac(self):
        """``all`` odmawia nadpisania .xlsx PRZED pierwszym żądaniem HTTP.

        Warunek („plik wynikowy istnieje, a nie podano ``--overwrite``'') jest
        znany od początku.  Użytkownik, który po prostu powtórzył polecenie
        z README, nie może stracić kilkunastu minut pobierania „w błoto'' ani
        niepotrzebnie obciążyć portalu.
        """
        praca = self.katalog_roboczy("u_all")
        with open(os.path.join(praca, "aukcje.xlsx"), "w", encoding="utf-8") as uchwyt:
            uchwyt.write("stary plik")
        wynik = self.cli(
            [
                "all", "--in", "xml_flexit", "--out", "aukcje.xlsx",
                "--base-url", self.portal.base, "--delay", "0",
            ],
            praca,
        )
        self.assertEqual(wynik.returncode, 1)
        self.assertIn("już istnieje", tekst(wynik.stderr))
        # żadnego pobierania: katalog na XML-e nawet nie powstał
        self.assertFalse(os.path.exists(os.path.join(praca, "xml_flexit")))
        self.assertNotIn("Szukam aukcji", tekst(wynik.stdout))
        self.assertNotIn("pobrano", tekst(wynik.stdout))
        # stary plik został nietknięty
        with open(os.path.join(praca, "aukcje.xlsx"), encoding="utf-8") as uchwyt:
            self.assertEqual(uchwyt.read(), "stary plik")

# --------------------------------------------------------------------------- #
# 5. USTERKI w wyniku i w dokumentacji
# --------------------------------------------------------------------------- #


class TestUsterkiWynikuIDokumentacji(BazaUX):
    """Rozbieżności README <-> rzeczywistość oraz pułapki w danych."""

    def test_niejednoznaczny_separator_tysiecy_zostaje_tekstem(self):
        """``1,234`` z angielskiego portalu NIE może stać się ``1.234``.

        ``SITE_NOTES.md`` mówi wprost: „Waluta EUR, treść po angielsku'', więc
        przecinek w polu liczbowym bywa separatorem TYSIĘCY: ``1,234`` = tysiąc
        dwieście trzydzieści cztery.  Odczytany po polsku dałby wartość
        mniejszą **1000 razy** — i nic by tego nie sygnalizowało.

        Separator (``,`` albo ``.``) z dokładnie TRZEMA cyframi po nim i bez
        drugiego separatora jest nierozstrzygalny, więc taka wartość zostaje
        TEKSTEM.  Zapisy jednoznaczne dalej są liczbami.
        """
        # 1) na poziomie funkcji z kontraktu — niejednoznaczne zostaje tekstem
        for zapis in ("1,234", "1,500", "12,345", "1.234", "2,500", "3,750"):
            with self.subTest(zapis=zapis):
                self.assertEqual(values.coerce_value(zapis), zapis)
        # ...a jednoznaczne dalej są liczbami
        self.assertEqual(values.coerce_value("950"), 950)
        self.assertEqual(values.coerce_value("1 234,56"), 1234.56)
        self.assertEqual(values.coerce_value("1.234,56"), 1234.56)
        self.assertEqual(values.coerce_value("1,234.56"), 1234.56)

        # 2) i to samo w gotowym arkuszu — tak zobaczy to użytkownik
        praca = self.katalog_roboczy("u_ceny")
        zapisz_xml(
            os.path.join(praca, "xml", "eur.xml"),
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<batch auction="flexit-auctions-18-06-2026-1103"><lot>'
            "<currency>EUR</currency><items>"
            "<item><model>T480</model><reserve>1,234</reserve><hammer>2,500</hammer></item>"
            "<item><model>T490</model><reserve>950</reserve><hammer>3,750</hammer></item>"
            "</items></lot></batch>\n",
        )
        self.cli(["build", "--in", "xml", "--out", "ceny.xlsx"], praca)
        wiersze = czytaj_arkusz(os.path.join(praca, "ceny.xlsx"))
        naglowek = list(wiersze[0])
        rezerwa = [w[naglowek.index("reserve")] for w in wiersze[1:]]
        mlotek = [w[naglowek.index("hammer")] for w in wiersze[1:]]
        self.assertEqual(rezerwa, ["1,234", 950])
        self.assertEqual(mlotek, ["2,500", "3,750"])

    def test_naglowek_atrybutu_rekordu_ma_nazwe_elementu(self):
        """Atrybut elementu-rekordu daje nagłówek ``item@nr``, nie samo ``@nr``.

        README, „Co znajdziesz w pliku wynikowym'', obiecuje:
        „atrybut → ``element@atrybut``'' — także dla atrybutów samego elementu
        powtarzalnego (``<item nr="1">``).
        """
        praca = self.katalog_roboczy("u_atrybut")
        zapisz_xml(os.path.join(praca, "xml", "a.xml"), PAKIET_TRZY)
        self.cli(["build", "--in", "xml", "--out", "a.xlsx"], praca)
        naglowek = list(czytaj_arkusz(os.path.join(praca, "a.xlsx"))[0])
        self.assertIn("item@nr", naglowek)
        self.assertNotIn("@nr", naglowek)

    def test_dlugi_tekst_obcinany_z_ostrzezeniem(self):
        """Tekst dłuższy niż 32767 znaków jest ucinany, ale NIE po cichu.

        Limit pochodzi z Excela, jednak użytkownik musi się dowiedzieć, że część
        opisu pakietu została odrzucona — i gdzie szukać pełnej treści.
        """
        praca = self.katalog_roboczy("u_dlugi")
        opis = "A" * 40000
        zapisz_xml(
            os.path.join(praca, "xml", "a.xml"),
            '<?xml version="1.0" encoding="UTF-8"?>\n<batch><items>'
            "<item><model>X</model><opis>%s</opis></item>"
            "<item><model>Y</model><opis>krotki</opis></item>"
            "</items></batch>\n" % opis,
        )
        wynik = self.cli(["build", "--in", "xml", "--out", "a.xlsx"], praca)
        self.assertEqual(wynik.returncode, 0)
        wiersze = czytaj_arkusz(os.path.join(praca, "a.xlsx"))
        wartosc = wiersze[1][list(wiersze[0]).index("opis")]
        self.assertEqual(len(wartosc), 32767)  # ucięte z 40000
        pelne = tekst(wynik.stdout) + tekst(wynik.stderr)
        self.assertIn("skrócono", pelne.lower())
        self.assertIn("32767", pelne)

    def test_version_dziala_takze_w_podkomendzie(self):
        """README wymienia ``--version`` wśród opcji WSPÓLNYCH — i tak jest.

        Opcja musi działać w każdej podkomendzie, nie tylko przed nią.
        """
        praca = self.katalog_roboczy("u_wersja")
        wzorzec = tekst(self.cli(["--version"], praca).stdout).strip()
        self.assertTrue(wzorzec.startswith("flexit2xlsx"), wzorzec)
        for podkomenda in ("build", "download", "all"):
            with self.subTest(podkomenda=podkomenda):
                wynik = self.cli([podkomenda, "--version"], praca)
                self.assertEqual(wynik.returncode, 0, tekst(wynik.stderr))
                self.assertEqual(tekst(wynik.stdout).strip(), wzorzec)

    def test_komunikaty_argparse_sa_po_polsku(self):
        """README (Windows, krok 5) obiecuje „pomoc po polsku'' — i tak jest.

        Po polsku muszą być nie tylko opisy opcji, ale też nagłówki sekcji
        i komunikaty o błędach składni: to właśnie one pojawiają się
        w chwili pomyłki adresata README („znam tylko podstawy'').
        """
        praca = self.katalog_roboczy("u_ang")
        pomoc = tekst(self.cli(["--help"], praca).stdout)
        self.assertNotIn("positional arguments", pomoc)
        self.assertNotIn("options:", pomoc)
        self.assertNotIn("show this help message", pomoc)
        self.assertIn("argumenty pozycyjne", pomoc)
        self.assertIn("opcje:", pomoc)
        self.assertIn("użycie:", pomoc)

        zly = self.cli(["build", "--nie-ma-takiej-opcji"], praca)
        blad = tekst(zly.stderr)
        self.assertNotEqual(zly.returncode, 0)
        self.assertNotIn("unrecognized arguments", blad)
        self.assertIn("błąd:", blad)
        self.assertIn("nierozpoznany argument", blad)

    def test_readme_wyjasnia_dwa_znaczenia_limit(self):
        """``--limit`` znaczy co innego w ``build`` i w ``all`` — README to mówi.

        W tabeli ``build`` to „weź najwyżej N plików XML'', w ``download``
        — „pobierz najwyżej N aukcji''.  Sekcja ``all`` musi rozstrzygnąć,
        które z dwóch znaczeń obowiązuje.
        """
        with open(README, encoding="utf-8") as uchwyt:
            tresc = uchwyt.read()
        sekcja = tresc.split("### `all` — pobierz i zbuduj", 1)[1].split("---", 1)[0]
        self.assertIn("--limit", sekcja)
        self.assertIn("AUKCJI", sekcja)

    def test_zapisana_strona_logowania_podpowiada_cookie(self):
        """W ``build`` strona logowania to czytelna diagnoza, nie „Uszkodzony XML''.

        README (krok 4) namawia, żeby pobrać pliki ręcznie z przeglądarki —
        „to działa zawsze''.  Gdy użytkownik nieświadomie zapisze pod nazwą
        ``.xml`` stronę logowania (bo sesja wygasła), musi się dowiedzieć,
        co naprawdę ma na dysku i co z tym zrobić.
        """
        praca = self.katalog_roboczy("u_html")
        zapisz_xml(os.path.join(praca, "xml", "dobry.xml"), PAKIET_TRZY)
        zapisz_xml(
            os.path.join(praca, "xml", "pakiet.xml"),
            "<!DOCTYPE html>\n<html lang='en'>\n"
            "<head><meta charset='utf-8'><title>Sign in</title></head>\n"
            "<body><form method='post'><input name='password'></form></body></html>\n",
        )
        wynik = self.cli(["build", "--in", "xml", "--out", "a.xlsx"], praca)
        self.assertEqual(wynik.returncode, 3)
        komunikat = tekst(wynik.stderr)
        self.assertNotIn("mismatched tag", komunikat)
        self.assertIn("to strona HTML, a nie XML", komunikat)
        self.assertIn("--cookie", komunikat)
        self.assertIn("pakiet.xml", komunikat)
        # dobry plik i tak trafia do arkusza
        wiersze = czytaj_arkusz(os.path.join(praca, "a.xlsx"))
        self.assertEqual(len(wiersze) - 1, 3)


if __name__ == "__main__":
    unittest.main(verbosity=2)
