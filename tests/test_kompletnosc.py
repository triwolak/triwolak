# -*- coding: utf-8 -*-
"""Testy KOMPLETNOŚCI: czy użytkownik naprawdę dostaje wszystkie aukcje w jednym pliku.

Ten moduł nie sprawdza pojedynczych funkcji "czy liczą dobrze" — od tego są
``test_cli.py``, ``test_xmlflatten.py`` i reszta.  Pyta o coś innego:

* czy dane, które użytkownik MA, dają się w ogóle wprowadzić do programu
  (katalog, pojedynczy plik, paczka ``.zip``, plik ``.xml.gz``),
* czy nic nie ginie i nic nie dubluje się PO CICHU (powtórzone pliki, wiersze
  zgubione w chwili zapisu, komórki obcięte w eksporcie),
* czy przy 300 plikach widać, że program pracuje,
* czy nieudany zapis ``.xlsx`` nie kasuje całej pracy (ratunkowy CSV).

Testy są w stylu ``unittest`` i nie potrzebują żadnej biblioteki zewnętrznej;
``openpyxl`` jest używany tylko tam, gdzie służy do NIEZALEŻNEGO sprawdzenia
wyniku, i jest pomijany, gdy go nie ma.
"""

from __future__ import annotations

import contextlib
import csv
import gzip
import io
import os
import shutil
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from flexit2xlsx import cli, values, xlsxwrite  # noqa: E402

try:  # tylko do niezależnej weryfikacji wyniku
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


XML_PAKIET = """<?xml version="1.0" encoding="UTF-8"?>
<batch auction="{auction}">
  <lot id="{lot}">
    <items>
      <item><model>Lenovo T480</model><serial>0071</serial><grade>A</grade></item>
      <item><model>Dell 5490</model><serial>0072</serial><grade>B</grade></item>
    </items>
  </lot>
</batch>
"""


def pakiet(auction: str = "1103", lot: str = "A1") -> str:
    return XML_PAKIET.format(auction=auction, lot=lot)


class Baza(unittest.TestCase):
    """Katalog roboczy + uruchamianie CLI z przechwyconymi strumieniami."""

    def setUp(self) -> None:
        self.tmp = tempfile.mkdtemp(prefix="flexit-kompletnosc-")
        self.addCleanup(shutil.rmtree, self.tmp, True)
        # katalogi tymczasowe archiwów są wspólne dla procesu — czyścimy je,
        # żeby testy nie zależały od kolejności uruchomienia
        self.addCleanup(cli.cleanup_archives)
        self.dane = os.path.join(self.tmp, "xml")
        os.makedirs(self.dane)
        self.wynik = os.path.join(self.tmp, "aukcje.xlsx")

    # -- narzędzia ---------------------------------------------------------

    def zapisz(self, nazwa: str, tresc) -> str:
        """Zapisuje plik w katalogu z danymi; zwraca ścieżkę."""
        path = os.path.join(self.dane, nazwa)
        tryb = "wb" if isinstance(tresc, (bytes, bytearray)) else "w"
        with open(path, tryb, **({} if "b" in tryb else {"encoding": "utf-8"})) as handle:
            handle.write(tresc)
        return path

    def uruchom(self, *argv):
        """Uruchamia CLI; zwraca ``(kod, stdout, stderr)``."""
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            kod = cli.main(list(argv))
        return kod, out.getvalue(), err.getvalue()

    def arkusz(self, path=None, nazwa=cli.SHEET_ALL):
        """Zawartość arkusza jako lista wierszy (wymaga openpyxl)."""
        book = openpyxl.load_workbook(path or self.wynik, read_only=True, data_only=True)
        try:
            return [list(row) for row in book[nazwa].iter_rows(values_only=True)]
        finally:
            book.close()

    def czytaj_csv(self, path, sep=cli.DEFAULT_CSV_SEP):
        with open(path, "r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.reader(handle, delimiter=sep))


# --------------------------------------------------------------------------- #
# 1. Wejście: archiwa ZIP
# --------------------------------------------------------------------------- #


class TestArchiwaZip(Baza):
    """Paczka .zip z XML-ami to normalne wejście, a nie ślepy zaułek."""

    def zip_z(self, nazwa: str, wpisy: dict) -> str:
        path = os.path.join(self.dane, nazwa)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for member, tresc in wpisy.items():
                archive.writestr(member, tresc)
        return path

    def test_zip_w_katalogu_jest_rozpakowany(self):
        self.zip_z("paczka.zip", {"loty/a.xml": pakiet(lot="A1"),
                                  "loty/b.xml": pakiet(lot="A2")})
        pliki = cli.collect_xml_files([self.dane])
        self.assertEqual(len(pliki), 2)
        for path in pliki:
            self.assertTrue(os.path.isfile(path))
            self.assertTrue(path.lower().endswith(".xml"))

    def test_zip_wskazany_wprost_jako_wejscie(self):
        path = self.zip_z("paczka.zip", {"a.xml": pakiet()})
        self.assertEqual(len(cli.collect_xml_files([path])), 1)

    def test_nazwa_pliku_niesie_slad_pochodzenia(self):
        self.zip_z("paczka_2026.zip", {"loty/lot_e0c8f.xml": pakiet()})
        nazwa = os.path.basename(cli.collect_xml_files([self.dane])[0])
        self.assertIn("paczka_2026", nazwa)
        self.assertIn("lot_e0c8f", nazwa)

    def test_zip_nie_wypuszcza_plikow_poza_katalog_tymczasowy(self):
        """Wpis ``../../.bashrc`` nie może wylądować poza katalogiem roboczym."""
        self.zip_z("zlosliwa.zip", {"../../ucieczka.xml": pakiet(),
                                    "/absolutna.xml": pakiet()})
        pliki = cli.collect_xml_files([self.dane])
        self.assertEqual(len(pliki), 2)
        scratch = os.path.realpath(os.path.dirname(pliki[0]))
        for path in pliki:
            self.assertEqual(os.path.realpath(os.path.dirname(path)), scratch)
            self.assertNotIn("..", os.path.basename(path).split(os.sep))

    def test_dwa_wywolania_daja_te_same_sciezki(self):
        """``all`` ogląda katalog dwa razy — wiersze nie mogą się zdublować."""
        self.zip_z("paczka.zip", {"a.xml": pakiet()})
        pierwsze = cli.collect_xml_files([self.dane])
        drugie = cli.collect_xml_files([self.dane])
        self.assertEqual(pierwsze, drugie)

    def test_paczka_paczek_jest_rozpakowana_rekurencyjnie(self):
        """"Pobierz wszystko" bywa jednym ZIP-em z ZIP-ami — po jednym na lot."""
        wewnetrzny = io.BytesIO()
        with zipfile.ZipFile(wewnetrzny, "w") as archive:
            archive.writestr("lot.xml", pakiet(lot="WEWN"))
        self.zip_z("wszystko.zip", {"loty/lot_1.zip": wewnetrzny.getvalue(),
                                    "luzem.xml": pakiet(lot="LUZ")})
        pliki = cli.collect_xml_files([self.dane])
        self.assertEqual(len(pliki), 2)
        self.assertTrue(all(p.lower().endswith(".xml") for p in pliki), pliki)

    def test_zip_bez_xml_ostrzega_i_nie_wywraca(self):
        self.zip_z("zdjecia.zip", {"foto.jpg": b"\xff\xd8\xff", "opis.txt": "nic"})
        bufor = io.StringIO()
        reporter = cli.Reporter(err=bufor)
        self.assertEqual(cli.collect_xml_files([self.dane], reporter), [])
        self.assertIn("nie zawiera plików .xml", bufor.getvalue())

    def test_uszkodzone_archiwum_nie_przerywa_pracy(self):
        self.zapisz("zepsuta.zip", b"to nie jest archiwum")
        self.zapisz("dobry.xml", pakiet())
        bufor = io.StringIO()
        reporter = cli.Reporter(err=bufor)
        pliki = cli.collect_xml_files([self.dane], reporter)
        self.assertEqual([os.path.basename(p) for p in pliki], ["dobry.xml"])
        self.assertIn("Pomijam archiwum", bufor.getvalue())

    @unittest.skipIf(openpyxl is None, "openpyxl potrzebny do weryfikacji wyniku")
    def test_end_to_end_zip_plus_xml_w_jednym_arkuszu(self):
        self.zapisz("luzem.xml", pakiet(lot="LUZ"))
        self.zip_z("paczka.zip", {"a.xml": pakiet(lot="ZIP1"),
                                  "b.xml": pakiet(lot="ZIP2")})
        kod, out, err = self.uruchom("build", "--in", self.dane, "--out", self.wynik)
        self.assertEqual(kod, cli.EXIT_OK, out + err)
        wiersze = self.arkusz()
        self.assertEqual(len(wiersze) - 1, 6)          # 3 pliki x 2 pozycje
        loty = {row[wiersze[0].index("batch/lot@id")] for row in wiersze[1:]}
        self.assertEqual(loty, {"LUZ", "ZIP1", "ZIP2"})


# --------------------------------------------------------------------------- #
# 2. Wejście: gzip i pojedynczy plik
# --------------------------------------------------------------------------- #


class TestPojedynczePlikiIGzip(Baza):

    def test_gz_jest_rozpakowany(self):
        path = os.path.join(self.dane, "batch.xml.gz")
        with gzip.open(path, "wb") as handle:
            handle.write(pakiet().encode("utf-8"))
        pliki = cli.collect_xml_files([self.dane])
        self.assertEqual(len(pliki), 1)
        with open(pliki[0], "rb") as handle:
            self.assertTrue(handle.read().lstrip().startswith(b"<?xml"))

    @unittest.skipIf(openpyxl is None, "openpyxl potrzebny do weryfikacji wyniku")
    def test_pojedynczy_plik_zamiast_katalogu(self):
        path = self.zapisz("jeden.xml", pakiet())
        kod, out, err = self.uruchom("build", "--in", path, "--out", self.wynik)
        self.assertEqual(kod, cli.EXIT_OK, out + err)
        self.assertEqual(len(self.arkusz()) - 1, 2)

    @unittest.skipIf(openpyxl is None, "openpyxl potrzebny do weryfikacji wyniku")
    def test_pojedynczy_plik_bez_opcji_in(self):
        """Ścieżka podana bez ``--in`` też działa (tak ludzie piszą naprawdę)."""
        path = self.zapisz("jeden.xml", pakiet())
        kod, out, err = self.uruchom("build", path, "--out", self.wynik)
        self.assertEqual(kod, cli.EXIT_OK, out + err)
        self.assertEqual(len(self.arkusz()) - 1, 2)

    def test_archiwum_pod_nazwa_xml_podpowiada_co_zrobic(self):
        """Paczka zapisana jako ``batch.xml`` to nie jest 'uszkodzony XML'."""
        bufor = io.BytesIO()
        with zipfile.ZipFile(bufor, "w") as archive:
            archive.writestr("a.xml", pakiet())
        self.zapisz("batch.xml", bufor.getvalue())
        kod, out, err = self.uruchom("build", "--in", self.dane, "--out", self.wynik)
        self.assertEqual(kod, cli.EXIT_NO_DATA)
        self.assertIn("archiwum ZIP", err)
        self.assertIn(".zip", err)


# --------------------------------------------------------------------------- #
# 3. Ciche dublowanie danych
# --------------------------------------------------------------------------- #


class TestPowtorzonychPlikow(Baza):
    """Ten sam pakiet pobrany dwa razy nie może po cichu podwoić stanu magazynu."""

    def test_find_duplicate_files_wskazuje_kopie(self):
        a = self.zapisz("lot.xml", pakiet())
        b = self.zapisz("lot (1).xml", pakiet())
        c = self.zapisz("inny.xml", pakiet(lot="INNY"))
        duplikaty = cli.find_duplicate_files([a, b, c])
        self.assertEqual(duplikaty, {b: a})

    def test_rozne_pliki_o_tym_samym_rozmiarze_to_nie_duplikaty(self):
        a = self.zapisz("a.xml", "<r><i><n>AAA</n></i><i><n>BBB</n></i></r>")
        b = self.zapisz("b.xml", "<r><i><n>CCC</n></i><i><n>DDD</n></i></r>")
        self.assertEqual(os.path.getsize(a), os.path.getsize(b))
        self.assertEqual(cli.find_duplicate_files([a, b]), {})

    def test_domyslnie_ostrzezenie_a_dane_zostaja(self):
        self.zapisz("lot.xml", pakiet())
        self.zapisz("lot-kopia.xml", pakiet())
        kod, out, err = self.uruchom("build", "--in", self.dane, "--out", self.wynik)
        self.assertEqual(kod, cli.EXIT_OK, out + err)
        self.assertIn("identyczną", err)
        self.assertIn("--skip-duplicates", err)
        self.assertIn("4 wiersze", out)          # nic nie usuwamy bez pytania

    def test_skip_duplicates_usuwa_kopie(self):
        self.zapisz("lot.xml", pakiet())
        self.zapisz("lot-kopia.xml", pakiet())
        kod, out, err = self.uruchom("build", "--in", self.dane, "--out", self.wynik,
                                     "--skip-duplicates")
        self.assertEqual(kod, cli.EXIT_OK, out + err)
        self.assertIn("2 wiersze", out)


# --------------------------------------------------------------------------- #
# 4. Eksport CSV — zwykły i ratunkowy
# --------------------------------------------------------------------------- #


class TestEksportuCsv(Baza):

    def test_csv_obok_xlsx_ma_te_same_wiersze(self):
        self.zapisz("lot.xml", pakiet())
        cel = os.path.join(self.tmp, "aukcje.csv")
        kod, out, err = self.uruchom("build", "--in", self.dane, "--out", self.wynik,
                                     "--csv", cel)
        self.assertEqual(kod, cli.EXIT_OK, out + err)
        wiersze = self.czytaj_csv(cel)
        self.assertEqual(len(wiersze) - 1, 2)
        self.assertEqual(wiersze[0][:3], cli.TECH_COLUMNS)
        self.assertIn("Lenovo T480", wiersze[1])

    def test_csv_jest_czytelny_dla_polskiego_excela(self):
        """UTF-8 z BOM + średnik: plik otwiera się dwuklikiem, z polskimi znakami."""
        self.zapisz("lot.xml", "<r><i><n>Żółć ąęś</n></i><i><n>Drugi</n></i></r>")
        cel = os.path.join(self.tmp, "a.csv")
        self.uruchom("build", "--in", self.dane, "--out", self.wynik, "--csv", cel)
        with open(cel, "rb") as handle:
            surowe = handle.read()
        self.assertTrue(surowe.startswith(b"\xef\xbb\xbf"))
        self.assertIn("Żółć ąęś".encode("utf-8"), surowe)
        self.assertIn(b";", surowe.splitlines()[0])

    def test_csv_nie_wykonuje_formul(self):
        """Wartość ``=cmd|...`` z pliku z sieci musi zostać tekstem."""
        cel = os.path.join(self.tmp, "a.csv")
        cli.write_csv(cel, ["A"], [["=cmd|' /C calc'!A0"], ["@SUM(1)"], ["-2+3"]])
        wiersze = self.czytaj_csv(cel)
        for wiersz in wiersze[1:]:
            self.assertTrue(wiersz[0].startswith("'"), wiersz)

    def test_csv_nie_obcina_dlugich_komorek(self):
        """W CSV nie obowiązuje limit 32767 znaków — obcinanie byłoby stratą."""
        dlugi = "x" * (values.MAX_CELL_CHARS + 500)
        cel = os.path.join(self.tmp, "a.csv")
        cli.write_csv(cel, ["A"], [[dlugi]])
        self.assertEqual(len(self.czytaj_csv(cel)[1][0]), len(dlugi))

    def test_csv_zapisuje_liczby_daty_i_bool_jednoznacznie(self):
        import datetime

        cel = os.path.join(self.tmp, "a.csv")
        cli.write_csv(cel, ["liczba", "data", "flaga", "puste"],
                      [[1234.5, datetime.date(2026, 6, 18), True, None]])
        wiersz = self.czytaj_csv(cel)[1]
        self.assertEqual(wiersz[0], "1234,5")        # przecinek — polski Excel
        self.assertEqual(wiersz[1], "2026-06-18")
        self.assertEqual(wiersz[2], "PRAWDA")
        self.assertEqual(wiersz[3], "")

    def test_csv_z_separatorem_przecinkowym_ma_kropke_dziesietna(self):
        cel = os.path.join(self.tmp, "a.csv")
        cli.write_csv(cel, ["liczba"], [[1234.5]], sep=",")
        self.assertEqual(self.czytaj_csv(cel, sep=",")[1][0], "1234.5")

    def test_nieudany_zapis_xlsx_ratuje_dane_do_csv(self):
        """Padnięty zapis .xlsx nie może skasować godzin pracy."""
        self.zapisz("lot.xml", pakiet())
        oryginal = xlsxwrite.write_workbook

        def wybuch(path, sheets, **kw):
            raise OSError("brak miejsca na dysku")

        xlsxwrite.write_workbook = wybuch
        try:
            kod, out, err = self.uruchom("build", "--in", self.dane, "--out", self.wynik)
        finally:
            xlsxwrite.write_workbook = oryginal
        self.assertEqual(kod, cli.EXIT_ERROR)
        ratunek = os.path.join(self.tmp, "aukcje.csv")
        self.assertTrue(os.path.exists(ratunek), err)
        self.assertIn("uratowane", err)
        self.assertEqual(len(self.czytaj_csv(ratunek)) - 1, 2)

    def test_ratunkowy_csv_nie_nadpisuje_cudzego_pliku(self):
        self.zapisz("lot.xml", pakiet())
        istniejacy = os.path.join(self.tmp, "aukcje.csv")
        with open(istniejacy, "w", encoding="utf-8") as handle:
            handle.write("cudze dane")
        oryginal = xlsxwrite.write_workbook
        xlsxwrite.write_workbook = lambda *a, **k: (_ for _ in ()).throw(OSError("nie"))
        try:
            self.uruchom("build", "--in", self.dane, "--out", self.wynik)
        finally:
            xlsxwrite.write_workbook = oryginal
        with open(istniejacy, "r", encoding="utf-8") as handle:
            self.assertEqual(handle.read(), "cudze dane")
        self.assertTrue(os.path.exists(os.path.join(self.tmp, "aukcje-2.csv")))

    def test_istniejacy_csv_wymaga_overwrite(self):
        self.zapisz("lot.xml", pakiet())
        cel = os.path.join(self.tmp, "a.csv")
        with open(cel, "w", encoding="utf-8") as handle:
            handle.write("stare")
        kod, out, err = self.uruchom("build", "--in", self.dane, "--out", self.wynik,
                                     "--csv", cel)
        self.assertEqual(kod, cli.EXIT_ERROR)
        self.assertIn("--overwrite", err)
        self.assertFalse(os.path.exists(self.wynik))   # nie zaczynamy zapisu


# --------------------------------------------------------------------------- #
# 5. Informacja zwrotna przy dużej liczbie plików
# --------------------------------------------------------------------------- #


class TestPostepu(Baza):

    def test_licznik_pokazuje_ile_zostalo(self):
        for number in range(12):
            self.zapisz("lot_%02d.xml" % number, pakiet(lot="L%02d" % number))
        kod, out, err = self.uruchom("build", "--in", self.dane, "--out", self.wynik)
        self.assertEqual(kod, cli.EXIT_OK, out + err)
        self.assertIn("[1/12]", out)
        self.assertIn("[12/12]", out)

    def test_widac_ze_zapis_sie_zaczal(self):
        self.zapisz("lot.xml", pakiet())
        kod, out, err = self.uruchom("build", "--in", self.dane, "--out", self.wynik)
        self.assertEqual(kod, cli.EXIT_OK, out + err)
        self.assertIn("Zapisuję", out)
        self.assertLess(out.index("Zapisuję"), out.index("Zapisano"))

    def test_postep_zapisu_co_kilkadziesiat_tysiecy_wierszy(self):
        """Przy dużym korpusie milczenie przez minuty wygląda jak zawieszenie."""
        bufor = io.StringIO()
        reporter = cli.Reporter(out=bufor)

        class Uchwyt:
            source = "duzy.xml"
            auction = "A"

            def load(self):
                import types

                doc = types.SimpleNamespace(
                    context={},
                    records=[{"a": i} for i in range(cli.PROGRESS_ROWS + 5)],
                )
                return doc

        rows = cli._iter_rows([Uchwyt()], ["a"], with_auction=True,
                              use_typing=False, reporter=reporter)
        self.assertEqual(len(list(rows)), cli.PROGRESS_ROWS + 5)
        self.assertIn("... zapisano", bufor.getvalue())

    def test_tryb_cichy_naprawde_milczy(self):
        for number in range(3):
            self.zapisz("lot_%d.xml" % number, pakiet(lot="L%d" % number))
        kod, out, err = self.uruchom("build", "--in", self.dane, "--out", self.wynik,
                                     "--quiet")
        self.assertEqual(kod, cli.EXIT_OK, err)
        self.assertEqual(out, "")


# --------------------------------------------------------------------------- #
# 6. Cicha utrata wierszy przy zapisie
# --------------------------------------------------------------------------- #


class TestUtratyWierszy(Baza):
    """Plik zniknięty w trakcie pracy nie może zostawić arkusza z cichą dziurą."""

    @unittest.skipIf(openpyxl is None, "openpyxl potrzebny do weryfikacji wyniku")
    def test_znikniety_plik_konczy_kodem_czesciowym(self):
        self.zapisz("zostaje.xml", pakiet(lot="ZOSTAJE"))
        znika = self.zapisz("znika.xml", pakiet(lot="ZNIKA"))

        # wymuszamy tryb strumieniowy: dokumenty są parsowane ponownie w chwili
        # zapisu, więc usunięcie pliku po wczytaniu jest realnym scenariuszem
        oryginal = cli.MEMORY_ROW_BUDGET
        cli.MEMORY_ROW_BUDGET = 0
        try:
            files = cli.collect_xml_files([self.dane])
            docs, errors = cli.load_documents(files)
            os.remove(znika)
            bufor = io.StringIO()
            reporter = cli.Reporter(err=bufor)
            lost = []
            sheets = cli.build_sheets(docs, errors, reporter=reporter, failures=lost)
            xlsxwrite.write_workbook(self.wynik, sheets)
        finally:
            cli.MEMORY_ROW_BUDGET = oryginal

        self.assertEqual([os.path.basename(p) for p in lost], ["znika.xml"])
        self.assertIn("pomijam jego wiersze", bufor.getvalue())
        self.assertEqual(len(self.arkusz()) - 1, 2)      # został tylko jeden plik

    def test_cli_zglasza_utrate_kodem_3(self):
        self.zapisz("lot.xml", pakiet())
        oryginal = cli._iter_rows

        def kaleki(docs, columns, **kw):
            failures = kw.get("failures")
            if failures is not None:
                failures.append("lot.xml")
            return iter(())

        cli._iter_rows = kaleki
        try:
            kod, out, err = self.uruchom("build", "--in", self.dane, "--out", self.wynik)
        finally:
            cli._iter_rows = oryginal
        self.assertEqual(kod, cli.EXIT_PARTIAL)
        self.assertIn("NIE trafiły do arkusza", err)


# --------------------------------------------------------------------------- #
# 7. Scenariusz "wszystko naraz"
# --------------------------------------------------------------------------- #


class TestScenariuszaUzytkownika(Baza):

    @unittest.skipIf(openpyxl is None, "openpyxl potrzebny do weryfikacji wyniku")
    def test_katalog_z_mieszanka_wejsc_daje_jeden_komplet(self):
        """Katalog jak u człowieka: luźne XML-e, paczka, kopia, śmieci."""
        self.zapisz("lot_1.xml", pakiet(lot="L1"))
        self.zapisz("lot_1_kopia.xml", pakiet(lot="L1"))
        self.zapisz("notatka.txt", "nic")
        with zipfile.ZipFile(os.path.join(self.dane, "paczka.zip"), "w") as archive:
            archive.writestr("lot_2.xml", pakiet(lot="L2"))
            archive.writestr("lot_3.xml", pakiet(lot="L3"))
        with gzip.open(os.path.join(self.dane, "lot_4.xml.gz"), "wb") as handle:
            handle.write(pakiet(lot="L4").encode("utf-8"))
        csv_cel = os.path.join(self.tmp, "aukcje.csv")

        kod, out, err = self.uruchom(
            "build", "--in", self.dane, "--out", self.wynik,
            "--csv", csv_cel, "--skip-duplicates")

        self.assertEqual(kod, cli.EXIT_OK, out + err)
        wiersze = self.arkusz()
        self.assertEqual(len(wiersze) - 1, 8)            # 4 loty x 2 pozycje
        indeks = wiersze[0].index("batch/lot@id")
        self.assertEqual({row[indeks] for row in wiersze[1:]},
                         {"L1", "L2", "L3", "L4"})
        self.assertEqual(len(self.czytaj_csv(csv_cel)) - 1, 8)


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
