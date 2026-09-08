# -*- coding: utf-8 -*-
"""Testy modułu :mod:`flexit2xlsx.values` — normalizacja wartości komórek.

Testy są tabelaryczne (``subTest``), żeby jeden nieudany przypadek nie ukrywał
pozostałych.  Uruchamiane przez ``python3 -m unittest`` oraz ``python3 -m pytest``.
"""

import datetime as dt
import math
import os
import sys
import unittest
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from flexit2xlsx.values import (  # noqa: E402
    MAX_CELL_CHARS,
    MAX_INT_DIGITS,
    MAX_SCALAR_LEN,
    coerce_value,
    looks_like_formula,
    normalize_ws,
    sanitize_cell,
)

NBSP = " "        # twarda spacja
NNBSP = " "       # wąska twarda spacja
THIN = " "        # cienka spacja
MINUS = "−"       # matematyczny minus
POLSKIE = "Zażółć gęślą jaźń ŁÓDŹ"


class TestNormalizeWs(unittest.TestCase):
    """Zwijanie białych znaków."""

    CASES = [
        ("Ala ma kota", "Ala ma kota"),
        ("  Ala   ma\tkota  ", "Ala ma kota"),
        ("\n\n Ala \n ma \r\n kota \n", "Ala ma kota"),
        ("Ala" + NBSP + "ma" + NBSP + "kota", "Ala ma kota"),
        ("Ala" + NNBSP + THIN + "ma", "Ala ma"),
        ("", ""),
        ("   ", ""),
        ("\t\r\n", ""),
        (NBSP, ""),
        (POLSKIE, POLSKIE),
        ("  Zażółć   gęślą  ", "Zażółć gęślą"),
        ("emoji 🐍  test", "emoji 🐍 test"),
        ("jeden", "jeden"),
        ("a" * 5000, "a" * 5000),
    ]

    def test_table(self):
        for source, expected in self.CASES:
            with self.subTest(source=source):
                self.assertEqual(normalize_ws(source), expected)

    def test_none_i_typy_nietekstowe(self):
        self.assertEqual(normalize_ws(None), "")
        self.assertEqual(normalize_ws(5), "5")
        self.assertEqual(normalize_ws(1.5), "1.5")

    def test_jest_idempotentna(self):
        for source, _ in self.CASES:
            with self.subTest(source=source):
                once = normalize_ws(source)
                self.assertEqual(normalize_ws(once), once)


class TestCoercePuste(unittest.TestCase):
    """Puste wejście -> ``None``."""

    CASES = [None, "", " ", "    ", "\t", "\n", "\r\n", " \t\r\n ", NBSP, NBSP * 3,
             NNBSP + THIN, "\n\n\n"]

    def test_table(self):
        for source in self.CASES:
            with self.subTest(source=repr(source)):
                self.assertIsNone(coerce_value(source))


class TestCoerceInt(unittest.TestCase):
    """Liczby całkowite."""

    CASES = [
        ("0", 0),
        ("1", 1),
        ("-1", -1),
        ("42", 42),
        ("2024", 2024),
        ("  7  ", 7),
        ("1000", 1000),
        ("1 000", 1000),
        ("1" + NBSP + "000", 1000),
        ("1" + NNBSP + "234", 1234),
        ("12 345 678", 12345678),
        ("1.234.567", 1234567),
        ("1,234,567", 1234567),
        ("-1 000", -1000),
        ("-12", -12),
        (MINUS + "12", -12),
        ("999", 999),
        ("100 000 000", 100000000),
        ("9" * MAX_INT_DIGITS, int("9" * MAX_INT_DIGITS)),
    ]

    def test_table(self):
        for source, expected in self.CASES:
            with self.subTest(source=source):
                value = coerce_value(source)
                self.assertIsInstance(value, int)
                self.assertNotIsInstance(value, bool)
                self.assertEqual(value, expected)


class TestCoerceFloat(unittest.TestCase):
    """Liczby dziesiętne — wszystkie warianty separatorów."""

    CASES = [
        ("1.5", 1.5),
        ("1,5", 1.5),
        ("-12,5", -12.5),
        ("-12.5", -12.5),
        (MINUS + "5,5", -5.5),
        ("0.5", 0.5),
        ("0,5", 0.5),
        ("-0,25", -0.25),
        ("1234.56", 1234.56),
        ("1234,56", 1234.56),
        ("1 234,56", 1234.56),
        ("1" + NBSP + "234,56", 1234.56),
        ("1" + NNBSP + "234,56", 1234.56),
        ("1" + THIN + "234,56", 1234.56),
        ("1.234,56", 1234.56),
        ("1,234.56", 1234.56),
        ("1 234 567,89", 1234567.89),
        ("1.234.567,89", 1234567.89),
        ("1,234,567.89", 1234567.89),
        ("12 345,678", 12345.678),
        ("3.14159", 3.14159),
        ("2,71828", 2.71828),
        ("-1 000,01", -1000.01),
        ("999,99", 999.99),
        # Jednoznaczne mimo trzech cyfr po separatorze: wiodące zero i część
        # całkowita dłuższa niż trzy cyfry wykluczają odczyt "tysiące".
        ("0,500", 0.5),
        ("1234,567", 1234.567),
        ("0.500", 0.5),
    ]

    def test_table(self):
        for source, expected in self.CASES:
            with self.subTest(source=source):
                value = coerce_value(source)
                self.assertIsInstance(value, float)
                self.assertAlmostEqual(value, expected, places=9)

    def test_REGRESJA_niejednoznaczny_tysiac_zostaje_tekstem(self):
        """"1,234" może znaczyć 1234 (EN) albo 1.234 (PL) — zostaje tekstem.

        Wcześniej przecinek był ZAWSZE separatorem dziesiętnym, więc kwota
        "1,234" z anglojęzycznego portalu trafiała do arkusza jako 1.234,
        czyli TYSIĄC RAZY mniejsza — bez żadnego ostrzeżenia.
        """
        for source in ("1,234", "1,500", "12,345", "123,456", "2,500",
                       "1.234", "1.500", "12.345", "-1,234", "-1.234"):
            with self.subTest(source=source):
                self.assertEqual(coerce_value(source), source)
        # Zapisy jednoznaczne nadal są liczbami:
        self.assertEqual(coerce_value("1,234,567"), 1234567)
        self.assertEqual(coerce_value("1.234.567"), 1234567)
        self.assertAlmostEqual(coerce_value("1 234,56"), 1234.56, places=9)
        self.assertAlmostEqual(coerce_value("1,234.56"), 1234.56, places=9)
        self.assertAlmostEqual(coerce_value("1,50"), 1.5, places=9)
        self.assertEqual(coerce_value("1 234"), 1234)


class TestCoerceZostajeTekstem(unittest.TestCase):
    """Przypadki NEGATYWNE — wartość musi pozostać dokładnie tym samym tekstem."""

    CASES = [
        # numery katalogowe / wiodące zera
        "007", "0012", "00", "000", "01", "0007,5", "00.5", "0x1F", "0b101",
        # wersje i identyfikatory z kropkami
        "1.2.3", "1.2.3.4", "v2.0", "V1", "2.0.0-rc1", "ver 1.0",
        # adresy IP
        "192.168.0.1", "10.0.0.255", "127.0.0.1", "255.255.255.0",
        # jednostki i opisy sprzętu
        "16 GB", "500GB", "2x8GB", "16 GB RAM", "1 TB SSD", "2 x 500 GB",
        "1 234,56 zł", "12%", "10 szt.", "5 kg", "230V",
        # notacja wykładnicza — świadomie nieobsługiwana
        "1e3", "1E5", "2,5e3", "1.5E-3",
        # telefony
        "+48 123 456 789", "+48123456789", "(22) 123-45-67", "22-123-45-67",
        "123-456-789", "0048 123 456 789", "+1 (555) 010-9999",
        # bool po polsku i inne "prawie bool"
        "tak", "nie", "TAK", "Nie", "yes", "no", "prawda", "fałsz", "truex",
        "true false", "T", "N",
        # zepsute liczby
        "1,234,56", "1 23", "10 00", "1 0000", "1.", "1,", ".5", ",5", "-", "--5",
        "- 5", "+5", "+12", "1..2", "1,,2", "1 . 2", "1.2 3", "1-2", "1/2", "½",
        "1_000", "1'234", "١٢٣", "١٢٣٤٥",
        # daty niepoprawne / niepełne / nieobsługiwane formaty
        "2024-13-01", "2024-00-10", "2024-03-32", "2023-02-29", "31.02.2024",
        "99.99.9999", "2024-03", "03/15/2024", "15/03/2024", "10:30", "10:30:00",
        "2024-03-15T25:00:00", "2024-03-15T10:61:00", "2024-03-15T10:30:00+99:00",
        # zwykły tekst
        "abc", "Dell PowerEdge R720", "SN0012345", "PL-1234", "NaN", "nan",
        "inf", "Infinity", "-inf", "None", "null", POLSKIE, "🐍💾", "Zażółć",
        "Lot nr 5", "Pakiet 1/3",
        # zbyt długa liczba całkowita (utrata precyzji w Excelu)
        "1234567890123456", "12345678901234567890",
        # zbyt długi ciąg, żeby w ogóle próbować rozpoznawania
        "1" * (MAX_SCALAR_LEN + 1),
    ]

    def test_table(self):
        for source in self.CASES:
            with self.subTest(source=source):
                self.assertEqual(coerce_value(source), source)

    def test_biale_znaki_sa_przycinane(self):
        for source in ["  16 GB  ", "\t007\n", NBSP + "v2.0" + NBSP]:
            with self.subTest(source=repr(source)):
                self.assertEqual(coerce_value(source), source.strip())


class TestCoerceBool(unittest.TestCase):
    """Tylko dosłowne true/false."""

    PRAWDA = ["true", "True", "TRUE", "tRuE", "  true  ", "\ttrue\n"]
    FALSZ = ["false", "False", "FALSE", "fAlSe", " false "]

    def test_prawda(self):
        for source in self.PRAWDA:
            with self.subTest(source=repr(source)):
                self.assertIs(coerce_value(source), True)

    def test_falsz(self):
        for source in self.FALSZ:
            with self.subTest(source=repr(source)):
                self.assertIs(coerce_value(source), False)

    def test_nie_bool(self):
        for source in ["tak", "nie", "yes", "no", "1", "0", "T", "F", "on", "off",
                       "prawda", "fałsz"]:
            with self.subTest(source=source):
                self.assertNotIsInstance(coerce_value(source), bool)


class TestCoerceDaty(unittest.TestCase):
    """Daty i daty z czasem."""

    DATY = [
        ("2024-03-15", dt.date(2024, 3, 15)),
        ("2024-3-5", dt.date(2024, 3, 5)),
        ("1999-12-31", dt.date(1999, 12, 31)),
        ("2024-02-29", dt.date(2024, 2, 29)),
        ("15.03.2024", dt.date(2024, 3, 15)),
        ("01.02.2024", dt.date(2024, 2, 1)),
        ("1.2.2024", dt.date(2024, 2, 1)),
        ("15-03-2024", dt.date(2024, 3, 15)),
        ("31-12-2023", dt.date(2023, 12, 31)),
        ("  2024-03-15  ", dt.date(2024, 3, 15)),
    ]

    CZASY = [
        ("2024-03-15T10:30:00", dt.datetime(2024, 3, 15, 10, 30, 0)),
        ("2024-03-15 10:30:00", dt.datetime(2024, 3, 15, 10, 30, 0)),
        ("2024-03-15T10:30", dt.datetime(2024, 3, 15, 10, 30)),
        ("2024-03-15T10:30:00.123456", dt.datetime(2024, 3, 15, 10, 30, 0, 123456)),
        ("2024-03-15T10:30:00,500", dt.datetime(2024, 3, 15, 10, 30, 0, 500000)),
        ("2024-03-15T00:00:00", dt.datetime(2024, 3, 15, 0, 0, 0)),
    ]

    CZASY_ZE_STREFA = [
        ("2024-03-15T10:30:00Z",
         dt.datetime(2024, 3, 15, 10, 30, tzinfo=dt.timezone.utc)),
        ("2024-03-15T10:30:00z",
         dt.datetime(2024, 3, 15, 10, 30, tzinfo=dt.timezone.utc)),
        ("2024-03-15T10:30:00+02:00",
         dt.datetime(2024, 3, 15, 10, 30,
                     tzinfo=dt.timezone(dt.timedelta(hours=2)))),
        ("2024-03-15T10:30:00-05:00",
         dt.datetime(2024, 3, 15, 10, 30,
                     tzinfo=dt.timezone(dt.timedelta(hours=-5)))),
        ("2024-03-15T10:30:00+0200",
         dt.datetime(2024, 3, 15, 10, 30,
                     tzinfo=dt.timezone(dt.timedelta(hours=2)))),
    ]

    def test_daty(self):
        for source, expected in self.DATY:
            with self.subTest(source=source):
                value = coerce_value(source)
                self.assertIsInstance(value, dt.date)
                self.assertNotIsInstance(value, dt.datetime)
                self.assertEqual(value, expected)

    def test_daty_z_czasem(self):
        for source, expected in self.CZASY:
            with self.subTest(source=source):
                value = coerce_value(source)
                self.assertIsInstance(value, dt.datetime)
                self.assertIsNone(value.tzinfo)
                self.assertEqual(value, expected)

    def test_daty_ze_strefa(self):
        for source, expected in self.CZASY_ZE_STREFA:
            with self.subTest(source=source):
                value = coerce_value(source)
                self.assertIsInstance(value, dt.datetime)
                self.assertIsNotNone(value.tzinfo)
                self.assertEqual(value, expected)

    def test_rok_nie_jest_data(self):
        self.assertEqual(coerce_value("2024"), 2024)


class TestCoerceTekst(unittest.TestCase):
    """Tekst wielolinijkowy, bardzo długi, ze znakami specjalnymi."""

    def test_wielolinijkowy_zachowuje_podzial(self):
        source = "Linia 1\nLinia 2\nLinia 3"
        self.assertEqual(coerce_value(source), source)

    def test_wielolinijkowy_normalizuje_konce_linii(self):
        self.assertEqual(coerce_value("a\r\nb\rc"), "a\nb\nc")

    def test_wielolinijkowy_przycina_wciecia_z_xml(self):
        source = "\n      Procesor: i7\n      RAM: 16 GB\n   "
        self.assertEqual(coerce_value(source), "Procesor: i7\nRAM: 16 GB")

    def test_wielolinijkowy_zwija_puste_linie(self):
        self.assertEqual(coerce_value("a\n\n\n\nb"), "a\n\nb")

    def test_wielolinijkowa_liczba_zostaje_tekstem(self):
        self.assertEqual(coerce_value("12\n34"), "12\n34")

    def test_bardzo_dlugi_tekst_bez_zmian(self):
        source = "x" * 100000
        self.assertEqual(coerce_value(source), source)

    def test_polskie_znaki_i_emoji(self):
        for source in [POLSKIE, "🐍 wąż", "ĄĆĘŁŃÓŚŹŻ", "µ Ω ½ – —"]:
            with self.subTest(source=source):
                self.assertEqual(coerce_value(source), source)

    def test_znaki_sterujace_nie_sa_usuwane_przez_coerce(self):
        """Za czyszczenie odpowiada sanitize_cell, nie coerce_value."""
        self.assertEqual(coerce_value("a\x00b"), "a\x00b")

    def test_typy_nietekstowe_przechodza_bez_zmian(self):
        marker = object()
        for value in [5, 1.5, True, dt.date(2024, 1, 1), marker]:
            with self.subTest(value=value):
                self.assertIs(coerce_value(value), value)

    def test_wynik_nigdy_nie_zawiera_crlf(self):
        for source in ["a\r\nb", "a\rb", "  a\r\n\r\n  b  "]:
            with self.subTest(source=repr(source)):
                self.assertNotIn("\r", coerce_value(source))


class TestSanitizeCell(unittest.TestCase):
    """Czyszczenie wartości przed zapisem do XLSX."""

    ZNAKI_DO_USUNIECIA = [
        ("a\x00b", "ab"),
        ("a\x01b", "ab"),
        ("a\x07b", "ab"),
        ("a\x08b", "ab"),
        ("a\x0bb", "ab"),
        ("a\x0cb", "ab"),
        ("a\x0eb", "ab"),
        ("a\x1fb", "ab"),
        ("\x00\x01\x02", ""),
        ("a￾b", "ab"),
        ("a￿b", "ab"),
        # Znaki sterujące C1 – w komórce nie niosą treści, a trafiają tam
        # ze ŹLE ZDEKODOWANYCH bajtów (mojibake typu "zaĹźĂłĹ\x82Ä\x87").
        ("a\x80b", "ab"),
        ("a\x82b", "ab"),
        ("a\x9fb", "ab"),
        ("zaĹźĂłĹ\x82Ä\x87", "zaĹźĂłĹÄ"),
    ]

    ZNAKI_DO_ZACHOWANIA = [
        "a\tb",
        "a\nb",
        "a\rb",
        "a\r\nb",
        POLSKIE,
        "🐍💾",
        "a\x7fb",       # DEL jest legalny w XML 1.0 i musi zostać
        "=SUMA(A1)",
    ]

    def test_usuwa_znaki_sterujace(self):
        for source, expected in self.ZNAKI_DO_USUNIECIA:
            with self.subTest(source=repr(source)):
                self.assertEqual(sanitize_cell(source), expected)

    def test_zachowuje_dozwolone_znaki(self):
        for source in self.ZNAKI_DO_ZACHOWANIA:
            with self.subTest(source=repr(source)):
                self.assertEqual(sanitize_cell(source), source)

    def test_usuwa_samotne_surogaty(self):
        source = "a\ud800b\udfffc"
        self.assertEqual(sanitize_cell(source), "abc")

    def test_przycina_do_limitu(self):
        value = sanitize_cell("x" * (MAX_CELL_CHARS + 5000))
        self.assertEqual(len(value), MAX_CELL_CHARS)
        self.assertEqual(value, "x" * MAX_CELL_CHARS)

    def test_czysci_przed_przycieciem(self):
        """Znaki sterujące znikają zanim policzymy limit długości."""
        source = "\x00" * 10 + "y" * MAX_CELL_CHARS
        value = sanitize_cell(source)
        self.assertEqual(len(value), MAX_CELL_CHARS)
        self.assertNotIn("\x00", value)

    def test_tekst_o_dokladnej_dlugosci_limitu(self):
        source = "z" * MAX_CELL_CHARS
        self.assertEqual(sanitize_cell(source), source)

    def test_nan_i_inf(self):
        for value in [float("nan"), float("inf"), float("-inf")]:
            with self.subTest(value=value):
                result = sanitize_cell(value)
                self.assertIsInstance(result, str)
        self.assertEqual(sanitize_cell(float("inf")), "inf")
        self.assertEqual(sanitize_cell(float("-inf")), "-inf")
        self.assertEqual(sanitize_cell(float("nan")), "nan")

    def test_liczby_przechodza_bez_zmian(self):
        for value in [0, 1, -1, 2 ** 40, 1.5, -0.25, 3.14159]:
            with self.subTest(value=value):
                self.assertEqual(sanitize_cell(value), value)

    def test_bool_pozostaje_boolem(self):
        self.assertIs(sanitize_cell(True), True)
        self.assertIs(sanitize_cell(False), False)

    def test_none(self):
        self.assertIsNone(sanitize_cell(None))

    def test_ogromny_int_staje_sie_tekstem(self):
        value = sanitize_cell(10 ** 400)
        self.assertIsInstance(value, str)
        self.assertTrue(value.startswith("1"))

    def test_daty(self):
        data = dt.date(2024, 3, 15)
        czas = dt.datetime(2024, 3, 15, 10, 30)
        self.assertEqual(sanitize_cell(data), data)
        self.assertEqual(sanitize_cell(czas), czas)

    def test_data_ze_strefa_jest_sprowadzana_do_utc(self):
        aware = dt.datetime(2024, 3, 15, 12, 30,
                            tzinfo=dt.timezone(dt.timedelta(hours=2)))
        value = sanitize_cell(aware)
        self.assertIsNone(value.tzinfo)
        self.assertEqual(value, dt.datetime(2024, 3, 15, 10, 30))

    def test_czas_ze_strefa_traci_tzinfo(self):
        value = sanitize_cell(dt.time(10, 30, tzinfo=dt.timezone.utc))
        self.assertIsNone(value.tzinfo)

    def test_decimal(self):
        self.assertEqual(sanitize_cell(Decimal("1.50")), 1.5)
        self.assertIsInstance(sanitize_cell(Decimal("1.50")), float)
        self.assertIsInstance(sanitize_cell(Decimal("NaN")), str)

    def test_bytes(self):
        self.assertEqual(sanitize_cell(b"ab\x00c"), "abc")
        self.assertEqual(sanitize_cell("zażółć".encode("utf-8")), "zażółć")

    def test_inne_typy_na_tekst(self):
        self.assertEqual(sanitize_cell([1, 2]), "[1, 2]")
        self.assertEqual(sanitize_cell({"a": 1}), "{'a': 1}")

    def test_nie_dodaje_apostrofu(self):
        for source in ["=1+1", "+1", "-1", "@SUM(A1)"]:
            with self.subTest(source=source):
                self.assertEqual(sanitize_cell(source), source)

    def test_jest_idempotentna(self):
        for source, _ in self.ZNAKI_DO_USUNIECIA:
            with self.subTest(source=repr(source)):
                once = sanitize_cell(source)
                self.assertEqual(sanitize_cell(once), once)


class TestLooksLikeFormula(unittest.TestCase):
    """Wykrywanie ciągów, które Excel mógłby uznać za formułę."""

    FORMULY = ["=1+1", "=SUM(A1:A2)", "=", "+1", "+", "-1", "-abc", "@SUM", "@",
               "  =A1", "\t=A1", "\tabc", "\rabc", "\t", "\r",
               "=HYPERLINK(\"http://zly.pl\")", "-12,5 zł", "+48 123 456 789",
               "\n=A1", "\n\n  -1"]

    NIE_FORMULY = ["", "a=1", "1+1", "abc", "  ", "'=1", "Ala ma kota", "007",
                   "2024-03-15", POLSKIE, "\n", "  ", " ", "abc=1", "x-1"]

    NIE_TEKST = [None, 0, 1, -1, -1.5, True, False, dt.date(2024, 1, 1),
                 dt.datetime(2024, 1, 1), b"=1+1", []]

    def test_formuly(self):
        for source in self.FORMULY:
            with self.subTest(source=repr(source)):
                self.assertTrue(looks_like_formula(source))

    def test_nie_formuly(self):
        for source in self.NIE_FORMULY:
            with self.subTest(source=repr(source)):
                self.assertFalse(looks_like_formula(source))

    def test_wartosci_nietekstowe(self):
        for value in self.NIE_TEKST:
            with self.subTest(value=repr(value)):
                self.assertFalse(looks_like_formula(value))


class TestIntegracjaWartosci(unittest.TestCase):
    """Wspólne własności całego łańcucha coerce_value -> sanitize_cell."""

    PROBKI = [
        "16 GB", "007", "1 234,56", "2024-03-15", "true", "", "  ", POLSKIE,
        "a\x00b", "x" * 40000, "🐍", "1.2.3", "2024-03-15T10:30:00Z",
        "Linia 1\nLinia 2", "-12,5", "192.168.0.1", "+48 123 456 789",
    ]

    def test_lancuch_nie_rzuca_wyjatkow(self):
        for source in self.PROBKI:
            with self.subTest(source=repr(source[:30])):
                sanitize_cell(coerce_value(source))

    def test_tekst_po_lancuchu_miesci_sie_w_limicie(self):
        for source in self.PROBKI:
            with self.subTest(source=repr(source[:30])):
                value = sanitize_cell(coerce_value(source))
                if isinstance(value, str):
                    self.assertLessEqual(len(value), MAX_CELL_CHARS)

    def test_wynik_nie_zawiera_niedozwolonych_znakow(self):
        zabronione = set(range(0x00, 0x09)) | {0x0b, 0x0c} | set(range(0x0e, 0x20))
        for source in self.PROBKI:
            value = sanitize_cell(coerce_value(source))
            if isinstance(value, str):
                with self.subTest(source=repr(source[:30])):
                    self.assertFalse(
                        any(ord(char) in zabronione for char in value)
                    )


class TestOdpornoscNaSmieci(unittest.TestCase):
    """Deterministyczny (stały seed) przegląd losowych śmieci — nic nie może wybuchnąć."""

    ALFABET = "0123456789.,+- \u00a0\u2212eE:TZxGB/()v\u0105\u017b\U0001f40d\t\n\r\x00\x0b\ud800"
    ZABRONIONE = set(range(0x00, 0x09)) | {0x0b, 0x0c} | set(range(0x0e, 0x20))

    def test_losowe_ciagi(self):
        import random

        generator = random.Random(20240315)  # stały seed = powtarzalny wynik
        for _ in range(3000):
            source = "".join(
                generator.choice(self.ALFABET)
                for _ in range(generator.randint(0, 24))
            )
            value = sanitize_cell(coerce_value(source))
            normalize_ws(source)
            looks_like_formula(source)
            if isinstance(value, str):
                self.assertLessEqual(len(value), MAX_CELL_CHARS)
                self.assertFalse(
                    any(ord(char) in self.ZABRONIONE for char in value),
                    "niedozwolony znak sterujący w wyniku dla %r" % source,
                )
                self.assertFalse(
                    any(0xD800 <= ord(char) <= 0xDFFF for char in value),
                    "surogat w wyniku dla %r" % source,
                )


try:  # openpyxl służy WYŁĄCZNIE do weryfikacji wyników w testach
    import openpyxl  # noqa: F401
    from openpyxl import Workbook, load_workbook
    _HAS_OPENPYXL = True
except ImportError:  # pragma: no cover
    _HAS_OPENPYXL = False


@unittest.skipUnless(_HAS_OPENPYXL, "openpyxl niedostępny")
class TestZgodnoscZOpenpyxl(unittest.TestCase):
    """Sprawdza, że wynik sanitize_cell faktycznie da się zapisać do XLSX."""

    TRUDNE = [
        "a\x00b", "\x01\x02\x03", "a\ud800b", "x" * 40000, POLSKIE, "🐍💾",
        "=1+1", "-12", "a\tb\nc", "a￾b", "16 GB", "007",
    ]

    def test_zapis_i_odczyt(self):
        import io

        book = Workbook()
        sheet = book.active
        wartosci = [sanitize_cell(coerce_value(item)) for item in self.TRUDNE]
        wartosci += [
            sanitize_cell(float("nan")),
            sanitize_cell(dt.datetime(2024, 3, 15, 12, 30,
                                      tzinfo=dt.timezone.utc)),
            sanitize_cell(dt.date(2024, 3, 15)),
            sanitize_cell(True),
            sanitize_cell(10 ** 400),
        ]
        for index, value in enumerate(wartosci, start=1):
            sheet.cell(row=index, column=1, value=value)

        buffer = io.BytesIO()
        book.save(buffer)
        buffer.seek(0)
        wczytany = load_workbook(buffer).active

        for index, value in enumerate(wartosci, start=1):
            with self.subTest(index=index):
                odczyt = wczytany.cell(row=index, column=1).value
                if value == "":
                    self.assertIn(odczyt, ("", None))
                elif isinstance(value, dt.date) and not isinstance(value, dt.datetime):
                    # openpyxl zwraca daty zawsze jako datetime
                    self.assertEqual(
                        odczyt.date() if isinstance(odczyt, dt.datetime) else odczyt,
                        value,
                    )
                else:
                    self.assertEqual(odczyt, value)


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
